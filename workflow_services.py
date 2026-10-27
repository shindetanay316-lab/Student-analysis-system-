"""Roster/marks validation and short, idempotent confirmation transactions."""
from decimal import Decimal,InvalidOperation
from datetime import datetime,timezone
from hashlib import sha256
from io import BytesIO
import json
from openpyxl import Workbook,load_workbook
from flask import abort
from flask_login import current_user
from models import db,Student,Subject,Semester,Enrollment,TheoryMarks,InternalMarksSubmission
from teaching_models import StaffAccessAudit
from teaching_access import require_assignment
from workflow_models import WorkflowReview,AcceptedInternalMark

COMPONENTS={'CT1':('ct1',10),'CT2':('ct2',10),'ASSIGNMENT':('assignment',10),'MIDSEM':('midsem',20)}
XLSX='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def subject_context(values):
    from routes_teaching import normalize_academic_year
    from batch_utils import clean_filter
    try:sem=int(values.get('semester_id',''))
    except (TypeError,ValueError):raise ValueError('Select a semester.')
    year=normalize_academic_year(values.get('academic_year',''))
    code=str(values.get('subject_code','')).strip()
    subject=db.session.get(Subject,code)
    if not subject or not subject.is_active or subject.semester_id!=sem or (subject.is_elective and not subject.parent_subject_code):
        raise ValueError('Select an active subject or actual elective option in this semester.')
    if db.session.get(Semester,(sem,year)) is None:raise ValueError('Assign the subject for this academic year first.')
    return {'subject_code':code,'semester_id':sem,'academic_year':year,
        'division':clean_filter(values.get('division','')),'batch':clean_filter(values.get('batch',''))}


def roster(context):
    from batch_utils import apply_student_batch_filters
    query=Student.query.join(Enrollment,Enrollment.prn==Student.prn).filter(
        Enrollment.subject_code==context['subject_code'],Enrollment.semester_id==context['semester_id'],Enrollment.academic_year==context['academic_year'])
    return apply_student_batch_filters(query,Student,context.get('division',''),context.get('batch','')).order_by(Student.prn).all()


def matching_marks(context):
    return TheoryMarks.query.filter_by(**{k:context[k] for k in ('subject_code','semester_id','academic_year')})


def fingerprint(context,kind,rows=None):
    if kind in {'LAB_MARKS','EXTERNAL_MARKS'}:
        from coursework_services import fingerprint as coursework_fingerprint
        return coursework_fingerprint(context)
    keys={k:context[k] for k in ('subject_code','semester_id','academic_year')}
    if kind=='MARKS':
        data={'roster':[(s.prn,s.name,s.division,s.batch) for s in roster(context)],
            'marks':[(m.prn,str(m.ct1),str(m.ct2),str(m.assignment),str(m.midsem)) for m in matching_marks(context).order_by(TheoryMarks.prn)],
            'locks':[(r.exam_type,r.division,r.batch) for r in InternalMarksSubmission.query.filter_by(**keys).order_by(InternalMarksSubmission.id)],
            'accepted':[(r.prn,r.component,r.state,str(r.value)) for r in AcceptedInternalMark.query.filter_by(**keys).order_by(AcceptedInternalMark.id)]}
    else:
        prns=[r['prn'] for r in rows]
        data={'students':[(s.prn,s.name,s.year_of_admission,s.division,s.batch) for s in Student.query.filter(Student.prn.in_(prns)).order_by(Student.prn)],
            'enrollments':[(e.prn,e.subject_code,e.semester_id,e.academic_year) for e in Enrollment.query.filter(
                Enrollment.prn.in_(prns),Enrollment.semester_id==context['semester_id'],Enrollment.academic_year==context['academic_year']).order_by(Enrollment.id)]}
    if kind=='MARKS':
        from coursework_services import approved_revisions
        revisions=approved_revisions(context)
        if revisions:data['revisions']=[r.token for r in revisions]
    return sha256(json.dumps(data,sort_keys=True,default=str).encode()).hexdigest()


def new_review(kind,context,rows,file=None):
    review=WorkflowReview(owner_id=current_user.user_id,kind=kind,context=context,rows=rows,snapshot=fingerprint(context,kind,rows))
    if file:
        from werkzeug.utils import secure_filename
        file.stream.seek(0);data=file.read();file.stream.seek(0)
        review.source_name=secure_filename(file.filename or 'source.xlsx')
        review.source_hash=sha256(data).hexdigest();review.source_workbook=data
    db.session.add(review);db.session.commit();return review


def workbook(file):
    if not file or not file.filename.lower().endswith('.xlsx'):raise ValueError('Upload an .xlsx workbook.')
    try:book=load_workbook(file,data_only=False)
    except Exception:raise ValueError('Cannot open this workbook. Use the downloaded .xlsx template.')
    if any(sheet.max_row>10000 or sheet.max_column>100 for sheet in book):raise ValueError('Workbook exceeds the supported roster size.')
    return book


def parse_marks(file,context):
    book=workbook(file);sheet=book.active
    if '_meta' not in book.sheetnames:raise ValueError('Download a fresh marks template; its context metadata is required.')
    meta={str(row[0]):str(row[1] or '').strip() for row in book['_meta'].iter_rows(min_row=1,max_row=8,max_col=2,values_only=True) if row[0]}
    for key in ('subject_code','semester_id','academic_year','exam_type'):
        if meta.get(key,'')!=str(context[key]):raise ValueError('Workbook belongs to a different subject, academic year, semester or assessment.')
    for key in ('division','batch'):
        if (meta.get(key) or 'ALL')!=(context.get(key) or 'ALL'):raise ValueError('Workbook belongs to a different division or practical batch.')
    header=next((row for row in range(1,min(30,sheet.max_row)+1) if 'PRN' in str(sheet.cell(row,3).value).upper() and 'ROLL' in str(sheet.cell(row,2).value).upper()),None)
    if header is None:raise ValueError('PRN header was not found. Use the official marks template.')
    students={s.prn:s for s in roster(context)}
    if not students:raise ValueError('No students are enrolled in this subject/context. Ask the administrator to configure its roster.')
    seen=set();rows=[];errors=[]
    for number in range(header+1,sheet.max_row+1):
        cell=sheet.cell(number,3);raw=cell.value
        if raw is None:continue
        prn=str(raw).strip()
        if prn.lower().startswith('note'):continue
        if cell.data_type=='f' or not isinstance(raw,str):errors.append(f'Row {number}: PRN must remain text, exactly as downloaded.');continue
        if prn in seen:errors.append(f'Row {number}: duplicate PRN {prn}.');continue
        seen.add(prn)
        if prn not in students:errors.append(f'Row {number}: {prn} is not in this subject roster.');continue
        name=' '.join(str(sheet.cell(number,4).value or '').split()).casefold()
        if name!=' '.join(students[prn].name.split()).casefold():errors.append(f'Row {number}: student name and PRN do not match the roster. Download a fresh template.');continue
        mark=sheet.cell(number,5)
        if mark.data_type=='f':errors.append(f'Row {number}: enter a mark, not an Excel formula.');continue
        raw=mark.value;text=str(raw or '').strip().upper()
        if raw is None or (isinstance(raw,str) and not raw.strip()):errors.append(f'Row {number}: mark is missing. Enter a number or A for absent.');continue
        if text in {'A','AB','ABSENT'}:value=None;state='ABSENT'
        else:
            try:value=Decimal(str(raw));valid=value.is_finite() and Decimal('0')<=value<=COMPONENTS[context['exam_type']][1] and value.as_tuple().exponent>=-2
            except (InvalidOperation,ValueError):valid=False
            if not valid:errors.append(f'Row {number}: mark must be 0–{COMPONENTS[context["exam_type"]][1]} with at most two decimal places, or A.');continue
            value=str(value);state='PRESENT'
        rows.append({'prn':prn,'name':students[prn].name,'value':value,'state':state})
    missing=set(students)-seen
    if missing:errors.append(f'{len(missing)} enrolled students are missing. Download a fresh template for this filter.')
    if errors:raise ValueError('\n'.join(errors[:25]))
    ensure_unlocked(context,rows)
    return rows


def ensure_unlocked(context,rows):
    field=COMPONENTS[context['exam_type']][0];students={s.prn:s for s in roster(context)}
    marks={r.prn:r for r in matching_marks(context)}
    locks=InternalMarksSubmission.query.filter_by(subject_code=context['subject_code'],semester_id=context['semester_id'],academic_year=context['academic_year'],exam_type=context['exam_type']).all()
    accepted={r.prn for r in AcceptedInternalMark.query.filter_by(subject_code=context['subject_code'],semester_id=context['semester_id'],academic_year=context['academic_year'],component=context['exam_type'])}
    for row in rows:
        student=students.get(row['prn'])
        if student is None:raise ValueError('The roster changed. Download a fresh template.')
        if row['prn'] in accepted or (marks.get(row['prn']) and getattr(marks[row['prn']],field) is not None) or any(
            (r.division=='ALL' or r.division==(student.division or '')) and (r.batch=='ALL' or r.batch==(student.batch or '')) for r in locks):
            raise ValueError(f'{student.prn}: this assessment is already submitted/locked. A saved component cannot be overwritten through another filter.')


def roster_template(context,source_id=None):
    book=Workbook();sheet=book.active;sheet.title='Roster'
    sheet.append(['PRN','Student Name','Admission Year','Division','Practical Batch'])
    if source_id:
        from result_data.models import ImportJob
        job=db.get_or_404(ImportJob,source_id)
        if not job.report or not job.confirmed_at or job.status not in {'READY','EXPORTED'}:raise ValueError('Confirm the university extraction before using its names for a roster template.')
        if str(job.report.semester)!=str(context['semester_id']):raise ValueError('Choose a university PDF for this same semester.')
        source=[(s.seat_number,s.name) for s in sorted(job.report.students,key=lambda s:s.seat_number)]
    else:source=[(s.prn,s.name) for s in roster(context)]
    for prn,name in source:
        existing=db.session.get(Student,prn)
        sheet.append([prn,name,existing.year_of_admission if existing else None,existing.division if existing else '',existing.batch if existing else ''])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value,str):cell.data_type='s'
    for row in range(2,max(102,len(source)+2)):
        sheet.cell(row,1).number_format='@'
    for column,width in [('A',24),('B',38),('C',20),('D',16),('E',20)]:sheet.column_dimensions[column].width=width
    sheet.freeze_panes='C2'
    meta=book.create_sheet('_meta');meta.sheet_state='hidden'
    for key in ('subject_code','semester_id','academic_year'):meta.append([key,context[key]])
    result=BytesIO();book.save(result);result.seek(0);return result


def validate_roster_rows(context,rows):
    seen=set();subject=db.session.get(Subject,context['subject_code'])
    for row in rows:
        prn=row['prn'];existing=db.session.get(Student,prn)
        if not prn or len(prn)>20 or prn in seen:raise ValueError(f'Duplicate, missing or invalid PRN: {prn}.')
        seen.add(prn)
        if not row['name'] or len(row['name'])>100 or row['name'].startswith(('=','+','-','@')):raise ValueError(f'{prn}: enter a valid student name.')
        if existing:
            if ' '.join(existing.name.split()).casefold()!=' '.join(row['name'].split()).casefold():raise ValueError(f'{prn}: the name conflicts with an existing student. Resolve the identity before enrollment.')
            if row['admission_year'] is not None and row['admission_year']!=existing.year_of_admission:raise ValueError(f'{prn}: admission year conflicts with the existing profile.')
            for key in ('division','batch'):
                if row.get(key) and row[key]!=(getattr(existing,key) or ''):raise ValueError(f'{prn}: {key} conflicts with the existing profile. Use Divisions and batches for placement changes.')
        elif row['admission_year'] is None:raise ValueError(f'{prn}: enter the actual admission year. It cannot be inferred from PRN or university results.')
        if row['admission_year'] is not None and not 1900<=row['admission_year']<=int(context['academic_year'][:4])+1:raise ValueError(f'{prn}: invalid admission year.')
        if len(row.get('division',''))>10 or len(row.get('batch',''))>10:raise ValueError(f'{prn}: division and batch must be at most 10 characters.')
        if subject.is_elective:
            others=Enrollment.query.join(Subject).filter(Enrollment.prn==prn,Enrollment.semester_id==context['semester_id'],Enrollment.academic_year==context['academic_year'],
                Subject.elective_group==subject.elective_group,Subject.subject_code!=subject.subject_code,Subject.is_elective.is_(True)).first()
            if others:raise ValueError(f'{prn}: another option in this elective group is already enrolled. Use Elective assignment to change options.')
    if not rows:raise ValueError('The roster contains no students.')


def parse_roster(file,context):
    book=workbook(file);sheet=book.active
    if [sheet.cell(1,i).value for i in range(1,6)]!=['PRN','Student Name','Admission Year','Division','Practical Batch']:raise ValueError('Use the downloaded roster template with its five named columns.')
    if '_meta' not in book.sheetnames:raise ValueError('Roster context metadata is missing. Download a fresh template.')
    meta={str(r[0]):str(r[1]) for r in book['_meta'].iter_rows(max_col=2,values_only=True) if r[0]}
    if any(meta.get(k)!=str(context[k]) for k in ('subject_code','semester_id','academic_year')):raise ValueError('Roster workbook belongs to another subject or academic term.')
    rows=[]
    for index in range(2,sheet.max_row+1):
        cells=[sheet.cell(index,c) for c in range(1,6)]
        if not any(c.value is not None for c in cells):continue
        if any(c.data_type=='f' for c in cells):raise ValueError(f'Row {index}: formulas are not allowed in a roster.')
        if not isinstance(cells[0].value,str):raise ValueError(f'Row {index}: PRN must be text to preserve all digits and leading zeros.')
        raw=cells[2].value
        try:year=int(str(raw)) if raw is not None and str(raw).strip() else None
        except (ValueError,TypeError):raise ValueError(f'Row {index}: admission year must be a whole year.')
        rows.append({'prn':cells[0].value.strip(),'name':str(cells[1].value or '').strip(),'admission_year':year,
            'division':str(cells[3].value or '').strip(),'batch':str(cells[4].value or '').strip()})
    validate_roster_rows(context,rows);return rows


def confirm(review):
    if review.kind in {'LAB_MARKS','EXTERNAL_MARKS'}:
        from coursework_services import confirm as coursework_confirm
        return coursework_confirm(review)
    context=review.context
    subject_context(context)
    if review.owner_id!=current_user.user_id and current_user.role!='ADMIN':abort(403)
    if review.kind=='ROSTER' and current_user.role!='ADMIN':abort(403)
    require_assignment(context['subject_code'],context['semester_id'],context['academic_year'])
    if review.confirmed_at:return False
    expiry=review.expires_at.replace(tzinfo=timezone.utc) if review.expires_at.tzinfo is None else review.expires_at
    if datetime.now(timezone.utc)>expiry:raise ValueError('This preview expired. Upload again to create a fresh preview.')
    # Serialize saves per subject, then claim the preview. The conditional write
    # also obtains SQLite's write lock before reading the final snapshot.
    db.session.execute(db.select(Subject).where(Subject.subject_code==context['subject_code']).with_for_update()).scalar_one()
    claimed=db.session.execute(db.update(WorkflowReview).where(WorkflowReview.token==review.token,WorkflowReview.confirmed_at.is_(None),WorkflowReview.claimed.is_(False)).values(claimed=True)).rowcount
    if not claimed:raise ValueError('This preview is already being confirmed. Refresh to see its status.')
    db.session.expire_all()
    if fingerprint(context,review.kind,review.rows)!=review.snapshot:raise ValueError('Students, marks or enrollment changed after this preview. Upload again and review the updated data.')
    if review.kind=='ROSTER':
        validate_roster_rows(context,review.rows)
        for row in review.rows:
            student=db.session.get(Student,row['prn'])
            if student is None:
                db.session.add(Student(prn=row['prn'],name=row['name'],year_of_admission=row['admission_year'],division=row['division'] or None,batch=row['batch'] or None))
                db.session.flush()
            keys={k:context[k] for k in ('subject_code','semester_id','academic_year')}
            if not Enrollment.query.filter_by(prn=row['prn'],**keys).first():db.session.add(Enrollment(prn=row['prn'],**keys))
    else:
        ensure_unlocked(context,review.rows)
        keys={k:context[k] for k in ('subject_code','semester_id','academic_year')}
        for row in review.rows:
            db.session.add(AcceptedInternalMark(prn=row['prn'],**keys,component=context['exam_type'],value=Decimal(row['value']) if row['value'] is not None else None,
                state=row['state'],review_token=review.token,actor_id=current_user.user_id))
        db.session.flush()
        from calculations import update_internal_totals
        for row in review.rows:
            mark=matching_marks(context).filter_by(prn=row['prn']).first()
            if mark is None:mark=TheoryMarks(prn=row['prn'],**keys);db.session.add(mark)
            setattr(mark,COMPONENTS[context['exam_type']][0],Decimal(row['value']) if row['value'] is not None else None)
            update_internal_totals(mark)
        db.session.add(InternalMarksSubmission(**keys,exam_type=context['exam_type'],division=context.get('division') or 'ALL',batch=context.get('batch') or 'ALL'))
    if review.kind=='MARKS':
        db.session.flush()
        from coursework_services import recalculate
        recalculate(context,[r['prn'] for r in review.rows])
    review.confirmed_at=datetime.now(timezone.utc)
    db.session.add(StaffAccessAudit(actor_id=current_user.user_id,user_id=current_user.user_id,action=review.kind+'_CONFIRMED',details={'review':review.token,'context':context,'students':len(review.rows)}))
    db.session.commit();return True
