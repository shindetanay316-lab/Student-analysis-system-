"""Calendar-aware attendance plans and reviewed, atomic attendance changes."""
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
import json
from types import SimpleNamespace
from zoneinfo import ZoneInfo
from flask_login import current_user
from openpyxl import load_workbook
from sqlalchemy import update
from models import db, Subject, Semester, Student, TimetableSlot, AttendanceSession, AttendanceEntry, Attendance
from attendance_models import AttendanceCalendar, AttendanceHoliday, AttendanceChange
from workflow_models import WorkflowReview, now
from teaching_access import require_assignment
from routes_teaching import normalize_academic_year, audit

COUNTED = {'COMPLETED', 'EXTRA'}


def local_now():
    return datetime.now(ZoneInfo('Asia/Kolkata'))


def context(values):
    from attendance_utils import clean_filter_value
    try:
        semester = int(values.get('semester_id'))
    except (ValueError, TypeError):
        raise ValueError('Choose a valid semester.')
    year = normalize_academic_year(values.get('academic_year'))
    code = str(values.get('subject_code', '')).strip().upper()
    subject = db.session.get(Subject, code)
    if semester not in range(1, 9) or not db.session.get(Semester, (semester, year)) or not subject or subject.semester_id != semester:
        raise ValueError('Choose an existing subject and teaching term.')
    require_assignment(code, semester, year)
    division = clean_filter_value(values.get('division'))
    batch = clean_filter_value(values.get('batch'))
    if len(division) > 10 or len(batch) > 10:
        raise ValueError('Division and batch must be at most 10 characters.')
    return dict(semester_id=semester, academic_year=year, subject_code=code, division=division, batch=batch)


def term_values(c):
    return {key: c[key] for key in ('semester_id', 'academic_year')}


def course_values(c):
    return {key: c[key] for key in ('semester_id', 'academic_year', 'subject_code')}


def lock_term(c):
    # Calendars affect every subject; serialize calendar, session and upload
    # decisions on the same composite teaching-term row on PostgreSQL.
    return db.session.execute(db.select(Semester).where(Semester.semester_id == c['semester_id'],
        Semester.academic_year == c['academic_year']).with_for_update()).scalar_one()


def roster(c):
    from attendance_utils import students_for_attendance
    return students_for_attendance(c['subject_code'], c['semester_id'], c['academic_year'], c['division'], c['batch'])


def applies(division, batch, student):
    return (division == 'ALL' or division == (student.division or '').upper()) and (batch == 'ALL' or batch == (student.batch or '').upper())


def eligible(c, spec, student):
    return applies(spec['division'], spec['batch'], student) and not holiday_on(c, date.fromisoformat(spec['date']), (student.division or 'ALL').upper())


def holiday_on(c, on_date, division='ALL'):
    return AttendanceHoliday.query.filter_by(**term_values(c), holiday_date=on_date, is_active=True).filter(
        AttendanceHoliday.division.in_(['ALL', division])).first()


def validate_range(c, start, end):
    if not start or not end or start > end or (end - start).days > 366:
        raise ValueError('Choose an ordered date range of at most 367 days.')
    calendar = AttendanceCalendar.query.filter_by(**term_values(c)).first()
    if calendar and (start < calendar.start_date or end > calendar.end_date):
        raise ValueError(f'Dates must be inside the term calendar ({calendar.start_date} to {calendar.end_date}).')


def key(item):
    return (item['date'], item['start_time'], item['end_time'], item['division'], item['batch'])


def session_spec(session):
    return dict(date=session.lecture_date.isoformat(), start_time=session.start_time.strftime('%H:%M'),
        end_time=session.end_time.strftime('%H:%M'), division=session.division, batch=session.batch,
        session_type=session.session_type)


def find_session(c, spec):
    from attendance_utils import parse_date, parse_time_value
    return AttendanceSession.query.filter_by(**course_values(c), lecture_date=parse_date(spec['date']),
        start_time=parse_time_value(spec['start_time']), end_time=parse_time_value(spec['end_time']),
        division=spec['division'], batch=spec['batch']).first()


def planned(c, start, end):
    """One canonical session per real slot; template filters never change its key."""
    from attendance_utils import DAY_NAMES
    validate_range(c, start, end)
    query = TimetableSlot.query.filter_by(**course_values(c), is_active=True)
    if c['division'] != 'ALL':
        query = query.filter(TimetableSlot.division.in_(['ALL', c['division']]))
    if c['batch'] != 'ALL':
        query = query.filter(TimetableSlot.batch.in_(['ALL', c['batch']]))
    slots = query.all()
    holidays = {(h.holiday_date, h.division) for h in AttendanceHoliday.query.filter_by(**term_values(c), is_active=True)
        .filter(AttendanceHoliday.holiday_date.between(start, end))}
    saved_sessions = AttendanceSession.query.filter_by(**course_values(c)).filter(AttendanceSession.lecture_date.between(start, end)).all()
    saved_by_key = {key(session_spec(s)): s for s in saved_sessions}
    results = {}
    on_date = start
    while on_date <= end:
        for slot in slots:
            division = c['division'] if slot.division == 'ALL' else slot.division
            if slot.day_of_week != DAY_NAMES[on_date.weekday()] or (on_date, 'ALL') in holidays or (on_date, division) in holidays:
                continue
            spec = dict(date=on_date.isoformat(), start_time=slot.start_time.strftime('%H:%M'),
                end_time=slot.end_time.strftime('%H:%M'), division=slot.division, batch=slot.batch, session_type=slot.session_type)
            saved = saved_by_key.get(key(spec))
            if saved and saved.status == 'CANCELLED':
                continue
            spec['saved_status'] = saved.status if saved else 'SCHEDULED'
            results[key(spec)] = spec
        on_date += timedelta(days=1)
    extras = [s for s in saved_sessions if s.status == 'EXTRA']
    for session in extras:
        if c['division'] != 'ALL' and session.division not in {'ALL', c['division']}:
            continue
        if c['batch'] != 'ALL' and session.batch not in {'ALL', c['batch']}:
            continue
        spec = session_spec(session)
        spec['saved_status'] = 'EXTRA'
        results[key(spec)] = spec
    return sorted(results.values(), key=key)


def generated_items(c, start, end):
    from attendance_utils import parse_time_value
    items = []
    for spec in planned(c, start, end):
        slot = SimpleNamespace(start_time=parse_time_value(spec['start_time']), end_time=parse_time_value(spec['end_time']),
            division=spec['division'], batch=spec['batch'], session_type=spec['session_type'])
        label = f"{date.fromisoformat(spec['date']).strftime('%d-%b')}\n{spec['start_time']}-{spec['end_time']}\n{spec['division']}/{spec['batch']}"
        items.append(dict(date=date.fromisoformat(spec['date']), slot=slot, label=label))
    return items


def snapshot(c):
    """Includes roster, timetable, calendar and saved entries, so stale previews fail."""
    sessions = AttendanceSession.query.filter_by(**course_values(c)).order_by(AttendanceSession.id).all()
    ids = [s.id for s in sessions]
    data = dict(roster=[(s.prn, s.name, s.division, s.batch) for s in roster(c)],
        sessions=[(s.id, session_spec(s), s.status) for s in sessions],
        entries=[(e.session_id, e.prn, e.status) for e in AttendanceEntry.query.filter(AttendanceEntry.session_id.in_(ids)).order_by(AttendanceEntry.session_id, AttendanceEntry.prn)],
        slots=[s.to_dict() for s in TimetableSlot.query.filter_by(**course_values(c)).order_by(TimetableSlot.id)],
        calendars=[(v.start_date.isoformat(), v.end_date.isoformat()) for v in AttendanceCalendar.query.filter_by(**term_values(c))],
        holidays=[(h.holiday_date.isoformat(), h.division, h.is_active) for h in AttendanceHoliday.query.filter_by(**term_values(c)).order_by(AttendanceHoliday.id)])
    return sha256(json.dumps(data, sort_keys=True, default=str).encode()).hexdigest()


def parse_workbook(content):
    from attendance_utils import parse_date, parse_time_value, PRESENT_VALUES, ABSENT_VALUES, EXCUSED_VALUES
    if len(content) > 12 * 1024 * 1024:
        raise ValueError('Attendance workbook must be at most 12 MB.')
    try:
        wb = load_workbook(BytesIO(content), data_only=False)
    except Exception as exc:
        raise ValueError('Cannot read this XLSX workbook.') from exc
    if not {'Attendance', '_meta'}.issubset(wb.sheetnames):
        raise ValueError('Use the generated attendance template with its metadata sheet.')
    ws, meta = wb['Attendance'], wb['_meta']
    if ws.max_row > 5005 or ws.max_column > 510 or meta.max_row > 520:
        raise ValueError('Workbook exceeds 5,000 students or 500 sessions.')
    values = {str(meta.cell(r, 1).value): meta.cell(r, 2).value for r in range(1, 8) if meta.cell(r, 1).value}
    c = context(values)
    students = {s.prn: s for s in roster(c)}
    if not students:
        raise ValueError('No enrolled students for the selected attendance scope. Ask admin to set up the roster.')
    specs = []
    for row in range(11, meta.max_row + 1):
        if not meta.cell(row, 1).value:
            continue
        lecture = parse_date(meta.cell(row, 1).value)
        start, end = parse_time_value(meta.cell(row, 2).value), parse_time_value(meta.cell(row, 3).value)
        if not lecture or not start or not end or start >= end:
            raise ValueError(f'Invalid session metadata at row {row}. Download a fresh template.')
        spec = dict(date=lecture.isoformat(), start_time=start.strftime('%H:%M'), end_time=end.strftime('%H:%M'),
            session_type=str(meta.cell(row, 4).value or 'THEORY'),
            division=str(meta.cell(row, 5).value or c['division']).upper(), batch=str(meta.cell(row, 6).value or c['batch']).upper())
        specs.append(spec)
    if not specs or len(specs) > 500 or len({key(s) for s in specs}) != len(specs):
        raise ValueError('No sessions, duplicate sessions, or too many sessions. Download a fresh template.')
    start, end = min(s['date'] for s in specs), max(s['date'] for s in specs)
    allowed = planned(c, date.fromisoformat(start), date.fromisoformat(end))
    valid = {key(s): s for s in allowed}
    # Older templates used the filter as the session scope. Resolve them to the
    # real weekly slot, rather than storing a duplicate class for another batch.
    legacy = not meta.cell(11, 5).value
    for index, spec in enumerate(specs):
        if legacy:
            matches = [s for s in allowed if all(s[k] == spec[k] for k in ('date', 'start_time', 'end_time'))]
            if len(matches) != 1:
                raise ValueError('Old attendance template is ambiguous or no longer scheduled. Download a fresh template.')
            specs[index] = dict(matches[0])
        elif key(spec) not in valid or spec['session_type'] != valid[key(spec)]['session_type']:
            raise ValueError('Session is cancelled, on a holiday, outside the calendar, or differs from the current timetable. Download a fresh template.')
    holidays = {(h.holiday_date.isoformat(), h.division) for h in AttendanceHoliday.query.filter_by(**term_values(c), is_active=True)}
    def eligible_student(spec, student):
        return applies(spec['division'], spec['batch'], student) and (spec['date'], 'ALL') not in holidays and (spec['date'], (student.division or 'ALL').upper()) not in holidays
    rows, seen = [], set()
    for line in range(5, ws.max_row + 1):
        value = ws.cell(line, 2).value
        if value in (None, ''):
            continue
        if not isinstance(value, str) or ws.cell(line, 2).data_type == 'f':
            raise ValueError(f'Row {line}: PRN must be text.')
        prn = value.strip()
        if prn in seen or prn not in students:
            raise ValueError(f'Row {line}: duplicate or out-of-scope PRN {prn}.')
        if str(ws.cell(line, 3).value or '').strip() != students[prn].name.strip():
            raise ValueError(f'Row {line}: name does not match PRN {prn}.')
        seen.add(prn)
        statuses = []
        for index, spec in enumerate(specs):
            cell = ws.cell(line, index + 4)
            raw = str(cell.value or '').strip().upper() if cell.value != 0 else '0'
            if cell.data_type == 'f':
                raise ValueError(f'Row {line}: attendance cells cannot contain formulas.')
            applicable = eligible_student(spec, students[prn])
            if not applicable:
                if raw not in {'', 'N/A', '-'}:
                    raise ValueError(f'{prn}: attendance supplied for another division/batch.')
                statuses.append(None)
                continue
            if raw in {'', '-'}:
                statuses.append(None)
            elif raw in PRESENT_VALUES:
                statuses.append('P')
            elif raw in ABSENT_VALUES:
                statuses.append('A')
            elif raw in EXCUSED_VALUES:
                statuses.append('O' if raw in {'O', 'OD', 'OFFICIAL'} else 'L')
            else:
                raise ValueError(f'Row {line}: use P, A, L or O. Leave future/unconducted columns blank.')
        rows.append(dict(prn=prn, name=students[prn].name, statuses=statuses))
    if seen != set(students):
        raise ValueError('The student roster changed or rows are missing. Download a fresh template.')
    conducted = []
    for index, spec in enumerate(specs):
        relevant = [r for r in rows if eligible_student(spec, students[r['prn']])]
        filled = [r for r in relevant if r['statuses'][index] is not None]
        if not filled:
            continue  # Entirely unfilled is pending, not an absent class.
        if len(filled) != len(relevant):
            raise ValueError(f"{spec['date']} {spec['start_time']}: complete every applicable student for a conducted class; leave the whole column blank if pending.")
        ending = datetime.combine(date.fromisoformat(spec['date']), datetime.strptime(spec['end_time'], '%H:%M').time(), tzinfo=ZoneInfo('Asia/Kolkata'))
        if ending > local_now():
            raise ValueError('A future or still-running class cannot be marked conducted. Leave its column blank.')
        conducted.append(index)
    if not conducted:
        raise ValueError('No conducted class is filled. Blank columns remain pending and do not affect attendance.')
    c['sessions'] = [specs[i] for i in conducted]
    for row in rows:
        row['statuses'] = [row['statuses'][i] for i in conducted]
    return c, rows


def new_review(file):
    if not file or not file.filename.lower().endswith('.xlsx'):
        raise ValueError('Select a generated, filled XLSX attendance workbook.')
    content = file.read(12 * 1024 * 1024 + 1)
    c, rows = parse_workbook(content)
    for index, spec in enumerate(c['sessions']):
        saved = find_session(c, spec)
        previous = {e.prn: e.status for e in saved.entries} if saved else {}
        for row in rows:
            row.setdefault('previous', []).append(previous.get(row['prn']))
    review = WorkflowReview(owner_id=current_user.user_id, kind='ATTENDANCE', context=c, rows=rows,
        snapshot=snapshot(c), source_name=file.filename.replace('\\', '/').rsplit('/', 1)[-1][:255],
        source_hash=sha256(content).hexdigest(), source_workbook=content)
    db.session.add(review)
    db.session.commit()
    return review


def confirm(review):
    from attendance_utils import parse_date, parse_time_value
    if review.confirmed_at:
        return False
    c = context(review.context)
    c['sessions'] = review.context['sessions']
    expires = review.expires_at.replace(tzinfo=timezone.utc) if review.expires_at.tzinfo is None else review.expires_at
    if expires <= now():
        raise ValueError('Preview expired. Upload again to review current attendance.')
    lock_term(c)
    db.session.execute(db.select(Subject).where(Subject.subject_code == c['subject_code']).with_for_update())
    if snapshot(c) != review.snapshot:
        raise ValueError('The roster, calendar, timetable or attendance changed after preview. Upload again; no attendance was changed.')
    if db.session.execute(update(WorkflowReview).where(WorkflowReview.token == review.token,
        WorkflowReview.claimed.is_(False), WorkflowReview.confirmed_at.is_(None)).values(claimed=True)).rowcount != 1:
        raise ValueError('This preview was already processed. Reload the page.')
    for index, spec in enumerate(c['sessions']):
        session = find_session(c, spec)
        before = {}
        if session:
            if session.status == 'CANCELLED':
                raise ValueError('A class was cancelled after preview. Upload a fresh template.')
            before = {e.prn: e.status for e in session.entries}
        else:
            session = AttendanceSession(**course_values(c), lecture_date=parse_date(spec['date']),
                start_time=parse_time_value(spec['start_time']), end_time=parse_time_value(spec['end_time']),
                division=spec['division'], batch=spec['batch'], session_type=spec['session_type'], status='COMPLETED')
            db.session.add(session)
            db.session.flush()
        if session.status != 'EXTRA':
            session.status = 'COMPLETED'
        changed = {}
        for row in review.rows:
            status = row['statuses'][index]
            if status is None:
                continue
            entry = AttendanceEntry.query.filter_by(session_id=session.id, prn=row['prn']).first()
            if not entry:
                entry = AttendanceEntry(session_id=session.id, prn=row['prn'])
                db.session.add(entry)
            entry.status = status
            changed[row['prn']] = status
        db.session.add(AttendanceChange(session_id=session.id, actor_id=current_user.user_id,
            action='UPLOAD_CONFIRMED', reason='Reviewed attendance workbook confirmed', before={p: before.get(p) for p in changed},
            after=changed, review_token=review.token))
    db.session.flush()
    refresh_summary(c)
    review.confirmed_at = now()
    db.session.commit()
    return True


def refresh_summary(c):
    # Include PRNs from cancelled classes so their old totals are reset to zero.
    buckets = {}
    pairs = db.session.query(AttendanceEntry, AttendanceSession).join(AttendanceSession).filter(
        AttendanceSession.semester_id == c['semester_id'], AttendanceSession.academic_year == c['academic_year'],
        AttendanceSession.subject_code == c['subject_code']).all()
    for entry, _session in pairs:
        buckets.setdefault(entry.prn, [0, 0])
    for item in effective_records(c):
        counts = buckets.setdefault(item['prn'], [0, 0])
        if item['status'] in {'P', 'A'}:
            counts[0] += 1
        if item['status'] == 'P':
            counts[1] += 1
    for prn, (total, present) in buckets.items():
        summary = Attendance.query.filter_by(prn=prn, **course_values(c)).first()
        if summary is None:
            summary = Attendance(prn=prn, **course_values(c))
            db.session.add(summary)
        summary.total_lectures, summary.attended_lectures = total, present
        summary.percentage = Decimal(str(round(100 * present / total, 2))) if total else Decimal(0)
    return len(buckets)


def reason_from(values):
    reason = str(values.get('reason', '')).strip()
    if not 5 <= len(reason) <= 500:
        raise ValueError('Enter a reason between 5 and 500 characters.')
    return reason


def save_calendar(values):
    from attendance_utils import parse_date
    semester = int(values.get('semester_id', 0))
    year = normalize_academic_year(values.get('academic_year'))
    if not db.session.get(Semester, (semester, year)):
        raise ValueError('Create/select an existing teaching term first.')
    lock_term(dict(semester_id=semester, academic_year=year))
    start, end = parse_date(values.get('start_date')), parse_date(values.get('end_date'))
    if not start or not end or start > end or (end - start).days > 366:
        raise ValueError('Enter valid term dates spanning at most 367 days.')
    existing_sessions = AttendanceSession.query.filter_by(semester_id=semester, academic_year=year).filter(
        AttendanceSession.status != 'CANCELLED', db.or_(AttendanceSession.lecture_date < start, AttendanceSession.lecture_date > end)).first()
    if existing_sessions:
        raise ValueError('These dates would exclude saved classes. Preserve the historical term range.')
    calendar = AttendanceCalendar.query.filter_by(semester_id=semester, academic_year=year).first()
    if not calendar:
        calendar = AttendanceCalendar(semester_id=semester, academic_year=year)
        db.session.add(calendar)
    calendar.start_date, calendar.end_date, calendar.actor_id = start, end, current_user.user_id
    audit(current_user, 'attendance_calendar_saved', {'semester_id': semester, 'academic_year': year, 'start': str(start), 'end': str(end)})
    db.session.commit()


def add_holiday(values):
    from attendance_utils import parse_date, clean_filter_value
    semester = int(values.get('semester_id', 0))
    year = normalize_academic_year(values.get('academic_year'))
    day = parse_date(values.get('holiday_date'))
    division = clean_filter_value(values.get('division'))
    reason = reason_from(values)
    if not day or len(division) > 10 or not db.session.get(Semester, (semester, year)):
        raise ValueError('Choose a valid term, date and division.')
    c = dict(semester_id=semester, academic_year=year)
    lock_term(c)
    validate_range(c, day, day)
    query = AttendanceSession.query.filter_by(**c, lecture_date=day).filter(AttendanceSession.status.in_(COUNTED))
    if division != 'ALL':
        query = query.filter(AttendanceSession.division.in_(['ALL', division]))
    if query.first():
        raise ValueError('Saved conducted/makeup classes exist on this date. Cancel them explicitly before adding a holiday.')
    holiday = AttendanceHoliday.query.filter_by(**c, holiday_date=day, division=division).first()
    if not holiday:
        holiday = AttendanceHoliday(**c, holiday_date=day, division=division)
        db.session.add(holiday)
    holiday.reason, holiday.is_active, holiday.actor_id = reason[:300], True, current_user.user_id
    audit(current_user, 'attendance_holiday_added', {'semester_id': semester, 'academic_year': year, 'date': str(day), 'division': division, 'reason': reason})
    db.session.commit()


def cancel_session(c, values):
    from attendance_utils import parse_date
    reason = reason_from(values)
    specs = planned(c, parse_date(values.get('date')), parse_date(values.get('date')))
    wanted = [s for s in specs if s['start_time'] == values.get('start_time') and s['end_time'] == values.get('end_time')
        and s['division'] == values.get('slot_division', c['division']) and s['batch'] == values.get('slot_batch', c['batch'])]
    if len(wanted) != 1:
        raise ValueError('Choose an existing scheduled class to cancel.')
    spec = wanted[0]
    session = find_session(c, spec)
    if not session:
        from attendance_utils import parse_time_value
        session = AttendanceSession(**course_values(c), lecture_date=parse_date(spec['date']),
            start_time=parse_time_value(spec['start_time']), end_time=parse_time_value(spec['end_time']),
            division=spec['division'], batch=spec['batch'], session_type=spec['session_type'], status='SCHEDULED')
        db.session.add(session)
        db.session.flush()
    before = session.status
    session.status = 'CANCELLED'
    db.session.add(AttendanceChange(session_id=session.id, actor_id=current_user.user_id, action='CANCELLED',
        reason=reason, before={'status': before}, after={'status': 'CANCELLED'}))
    db.session.flush()
    refresh_summary(c)
    db.session.commit()


def add_makeup(c, values):
    from attendance_utils import parse_date, parse_time_value
    reason = reason_from(values)
    day = parse_date(values.get('date'))
    start, end = parse_time_value(values.get('start_time')), parse_time_value(values.get('end_time'))
    validate_range(c, day, day)
    if not start or not end or start >= end:
        raise ValueError('Enter a valid makeup class time.')
    if not roster(c):
        raise ValueError('Set up the student roster before adding a makeup class.')
    if holiday_on(c, day, c['division']):
        raise ValueError('Remove the holiday first if a makeup class is officially scheduled on this date.')
    spec = dict(date=day.isoformat(), start_time=start.strftime('%H:%M'), end_time=end.strftime('%H:%M'), division=c['division'], batch=c['batch'])
    if find_session(c, spec):
        raise ValueError('This session already exists (including cancelled sessions). Its history was preserved.')
    # Avoid double counting a weekly class as a makeup class with another scope.
    overlap = AttendanceSession.query.filter_by(**term_values(c), lecture_date=day).filter(
        AttendanceSession.status != 'CANCELLED', AttendanceSession.start_time < end, AttendanceSession.end_time > start,
        db.or_(AttendanceSession.division == 'ALL', c['division'] == 'ALL', AttendanceSession.division == c['division']),
        db.or_(AttendanceSession.batch == 'ALL', c['batch'] == 'ALL', AttendanceSession.batch == c['batch'])).first()
    from attendance_utils import DAY_NAMES
    weekly = TimetableSlot.query.filter_by(**term_values(c), day_of_week=DAY_NAMES[day.weekday()], is_active=True).filter(
        TimetableSlot.start_time < end, TimetableSlot.end_time > start,
        db.or_(TimetableSlot.division == 'ALL', c['division'] == 'ALL', TimetableSlot.division == c['division']),
        db.or_(TimetableSlot.batch == 'ALL', c['batch'] == 'ALL', TimetableSlot.batch == c['batch'])).all()
    for slot in weekly:
        saved = AttendanceSession.query.filter_by(**term_values(c), lecture_date=day, subject_code=slot.subject_code,
            start_time=slot.start_time, end_time=slot.end_time, division=slot.division, batch=slot.batch).first()
        if not saved or saved.status != 'CANCELLED':
            raise ValueError('Makeup class overlaps the weekly timetable. Cancel/move the conflicting class first.')
    if overlap:
        raise ValueError('Makeup class overlaps another saved class for the same division/batch.')
    subject = db.session.get(Subject, c['subject_code'])
    session = AttendanceSession(**course_values(c), lecture_date=day, start_time=start, end_time=end,
        division=c['division'], batch=c['batch'], session_type=subject.subject_type, status='EXTRA')
    db.session.add(session)
    db.session.flush()
    db.session.add(AttendanceChange(session_id=session.id, actor_id=current_user.user_id, action='MAKEUP_SCHEDULED',
        reason=reason, before={}, after=session_spec(session)))
    db.session.commit()


def effective_records(c, start=None, end=None):
    """Deduplicate old filter-scoped copies without deleting their evidence.

    Conflicting legacy copies stay pending until a reviewed upload supplies an
    authoritative value. Same-date classes at different times remain distinct.
    """
    clock = local_now()
    query = db.session.query(AttendanceEntry, AttendanceSession).join(AttendanceSession).filter(
        AttendanceSession.semester_id == c['semester_id'], AttendanceSession.academic_year == c['academic_year'],
        AttendanceSession.subject_code == c['subject_code'], AttendanceSession.status.in_(COUNTED),
        db.or_(AttendanceSession.lecture_date < clock.date(), db.and_(
            AttendanceSession.lecture_date == clock.date(), AttendanceSession.end_time <= clock.time().replace(tzinfo=None))))
    if start:
        query = query.filter(AttendanceSession.lecture_date >= start)
    if end:
        query = query.filter(AttendanceSession.lecture_date <= end)
    pairs = query.all()
    ids = {s.id for _, s in pairs}
    reviewed = {}
    for change in AttendanceChange.query.filter(AttendanceChange.session_id.in_(ids), AttendanceChange.action == 'UPLOAD_CONFIRMED').order_by(AttendanceChange.id):
        for prn, status in change.after.items():
            reviewed[(change.session_id, prn)] = (change.id, status)
    cancelled = AttendanceSession.query.filter_by(**course_values(c), status='CANCELLED').all()
    from term_services import students_in_term
    students = {s.prn: s for s in students_in_term(Student.query.filter(Student.prn.in_([e.prn for e, _ in pairs])).all(),c['semester_id'],c['academic_year'])}
    groups = {}
    for entry, session in pairs:
        # Explicit cancellations also suppress a legacy copy of that real class.
        if any(v.lecture_date == session.lecture_date and v.start_time == session.start_time and v.end_time == session.end_time
               and applies(v.division, v.batch, students[entry.prn]) for v in cancelled):
            continue
        identifier = (entry.prn, session.lecture_date, session.start_time, session.end_time)
        group = groups.setdefault(identifier, dict(prn=entry.prn, values=set(), decisions=[]))
        group['values'].add(entry.status)
        if (session.id, entry.prn) in reviewed:
            group['decisions'].append(reviewed[(session.id, entry.prn)])
    result = []
    for group in groups.values():
        status = max(group['decisions'])[1] if group['decisions'] else next(iter(group['values'])) if len(group['values']) == 1 else 'CONFLICT'
        result.append(dict(prn=group['prn'], status=status))
    return result


def report_rows(c, start=None, end=None):
    """Compute without writing. Pending and leave entries never become absences."""
    buckets = {}
    for item in effective_records(c, start, end):
        data = buckets.setdefault(item['prn'], dict(total=0, present=0, leave=0, official=0, conflicts=0))
        if item['status'] in {'P', 'A'}:
            data['total'] += 1
        if item['status'] == 'P':
            data['present'] += 1
        elif item['status'] == 'L':
            data['leave'] += 1
        elif item['status'] == 'O':
            data['official'] += 1
        elif item['status'] == 'CONFLICT':
            data['conflicts'] += 1
    result = []
    for student in roster(c):
        data = buckets.get(student.prn, dict(total=0, present=0, leave=0, official=0, conflicts=0))
        data['percentage'] = round(data['present'] / data['total'] * 100, 2) if data['total'] else None
        result.append(dict(student=student, **data))
    return result
