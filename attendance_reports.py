"""Attendance summary and source-entry registers using the application's exporters."""
from io import BytesIO
from xml.sax.saxutils import escape
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from models import db, AttendanceEntry, AttendanceSession
from attendance_services import course_values


HEADERS = ['Sr. No.', 'PRN', 'Student Name', 'Counted classes', 'Present', 'Absent', 'Leave', 'Official duty', 'Percentage', 'Status']


def values(rows):
    return [[i, r['student'].prn, r['student'].name, r['total'], r['present'], r['total'] - r['present'], r['leave'], r['official'],
        r['percentage'] if r['percentage'] is not None else 'Pending', 'Review conflicts' if r.get('conflicts') else 'Pending' if not r['total'] else 'Below 75%' if r['percentage'] < 75 else '75% or above'] for i, r in enumerate(rows, 1)]


def generate_report(c, rows, fmt):
    title = f"Attendance - {c['subject_name']} ({c['subject_code']})"
    subtitle = f"Semester {c['semester_id']} | {c['academic_year']} | {c['division']}/{c['batch']} | {c.get('start_date') or 'All dates'} to {c.get('end_date') or 'latest'}"
    note = 'P / (P + A) x 100. Leave/official duty excluded. Pending, cancelled and conflicting legacy entries do not count.'
    buffer = BytesIO()
    filename = f"Attendance_Summary_{c['subject_code']}_{c['academic_year']}"
    data = values(rows)
    if fmt == 'excel':
        wb = Workbook()
        summary = wb.active
        summary.title = 'Attendance Summary'
        for text in (title, subtitle, note):
            summary.append([text])
            summary.merge_cells(start_row=summary.max_row, start_column=1, end_row=summary.max_row, end_column=10)
        summary.append(HEADERS)
        for row in data:
            summary.append(row)
        register = wb.create_sheet('Session Register')
        register.append(['Sr. No.', 'PRN', 'Student Name', 'Date', 'Start', 'End', 'Division', 'Batch', 'Class status', 'Attendance'])
        prns = {r['student'].prn: r['student'].name for r in rows}
        query = db.session.query(AttendanceEntry, AttendanceSession).join(AttendanceSession).filter(
            AttendanceSession.semester_id == c['semester_id'], AttendanceSession.academic_year == c['academic_year'],
            AttendanceSession.subject_code == c['subject_code'], AttendanceEntry.prn.in_(list(prns)))
        if c.get('start_date'):
            from attendance_utils import parse_date
            query = query.filter(AttendanceSession.lecture_date >= parse_date(c['start_date']))
        if c.get('end_date'):
            from attendance_utils import parse_date
            query = query.filter(AttendanceSession.lecture_date <= parse_date(c['end_date']))
        for i, (entry, session) in enumerate(query.order_by(AttendanceSession.lecture_date, AttendanceSession.start_time, AttendanceEntry.prn), 1):
            register.append([i, entry.prn, prns[entry.prn], session.lecture_date.isoformat(), session.start_time.strftime('%H:%M'),
                session.end_time.strftime('%H:%M'), session.division, session.batch, session.status, entry.status])
        for sheet, header, widths in [(summary, 4, [9, 22, 34, 18, 12, 12, 12, 16, 16, 19]), (register, 1, [9, 22, 34, 15, 12, 12, 12, 12, 18, 16])]:
            sheet.freeze_panes = f'D{header + 1}'
            sheet.auto_filter.ref = f'A{header}:J{sheet.max_row}'
            for col, width in enumerate(widths, 1):
                sheet.column_dimensions[get_column_letter(col)].width = width
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
                    cell.font = Font(name='Calibri', size=11)
                    cell.alignment = Alignment(vertical='center', wrap_text=True)
            for cell in sheet[header]:
                cell.font = Font(name='Calibri', bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='14283F')
            sheet.row_dimensions[header].height = 32
        summary.row_dimensions[1].height = 30
        summary['A1'].font = Font(name='Calibri', size=14, bold=True)
        summary.row_dimensions[3].height = 32
        wb.save(buffer)
        mime, suffix = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', '.xlsx'
    else:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        styles = getSampleStyleSheet()
        small = ParagraphStyle('AttendanceCell', fontName='Helvetica', fontSize=8, leading=11)
        story = [Paragraph('CSMSS Chh. Shahu College of Engineering', styles['Title']),
            Paragraph(escape(title), styles['Heading2']), Paragraph(escape(subtitle), styles['Normal']),
            Spacer(1, 10), Paragraph(escape(note), styles['Normal']), Spacer(1, 12)]
        body = [[Paragraph(escape(str(v)), small) for v in row] for row in [HEADERS] + data]
        table = Table(body, colWidths=[35, 110, 167, 60, 44, 44, 42, 50, 68, 166 - 94], repeatRows=1)
        table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e5eaf0')),
            ('GRID', (0, 0), (-1, -1), .3, colors.HexColor('#cbd5e1')), ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
            ('LEFTPADDING', (0, 0), (-1, -1), 5), ('RIGHTPADDING', (0, 0), (-1, -1), 5)]))
        story.append(table)
        SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=20, rightMargin=20, topMargin=24, bottomMargin=24).build(story)
        mime, suffix = 'application/pdf', '.pdf'
    buffer.seek(0)
    return buffer, filename + suffix, mime
