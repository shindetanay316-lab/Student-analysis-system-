"""A department workbook using the application's existing Excel export stack."""
from io import BytesIO
from datetime import datetime
from zoneinfo import ZoneInfo
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from models import db, Student, Enrollment, Subject
from completion_models import StudentTerm, AssessmentPolicy, AssessmentScore, AssessmentScoreChange, ProgressionDecision, CourseOutcomeSelection
from academic_link_models import UniversityLinkedResult, StudentPortfolio
from coursework_models import MarkRevision
from curriculum_models import StudentCurriculumTerm
from batch_utils import apply_student_batch_filters
from term_services import students_in_term
from progression_services import local_result
from attendance_services import report_rows


def build_master(c, filters):
    query = Student.query.join(Enrollment,Enrollment.prn==Student.prn).filter(
        Enrollment.semester_id==c['semester_id'],Enrollment.academic_year==c['academic_year']).distinct()
    raw_cohort = str(filters.get('cohort') or '').strip()
    if raw_cohort:
        try:
            query=query.filter(Student.year_of_admission==int(raw_cohort))
        except ValueError:
            raise ValueError('Admission cohort must be a whole year.')
    query=apply_student_batch_filters(query,Student,filters.get('division'),filters.get('batch'),**c)
    students=students_in_term(query.order_by(Student.prn).limit(5001).all(),**c)
    if not students:
        raise ValueError('No enrolled students match the selected teaching term and filters.')
    if len(students)>5000:
        raise ValueError('Filter by admission cohort/division/batch to export at most 5,000 students.')
    prns={s.prn for s in students}
    names={s.prn:s.name for s in students}
    def identity(index,prn):
        return [index,prn,names[prn]]
    headers=['Sr. No.','PRN','Student Name']
    tables={}
    result_rows,profile_rows=[],[]
    from coursework_services import effective_states
    state_cache={}
    for i,student in enumerate(students,1):
        result=local_result(student.prn,c['semester_id'],c['academic_year'])
        profile_rows.append(identity(i,student.prn)+[student.year_of_admission,student.roll_number,student.division,student.batch,
            student.placement_origin,result['total_credits'],result['credits_earned'],result['sgpa'] if result['sgpa'] is not None else 'Pending',result['result']])
        for row in result['subjects']:
            code=row['subject_code']
            if code not in state_cache:
                state_cache[code]=effective_states(dict(c,subject_code=code,division='',batch=''))
            absent=[component for (prn,category,component),state in state_cache[code].items() if prn==student.prn and state=='ABSENT']
            result_rows.append(identity(len(result_rows)+1,student.prn)+[c['academic_year'],c['semester_id'],row['subject_name'],row['subject_code'],row['credits'],
                row['internal'] if row['internal'] is not None else 'Pending',row['external'] if row['external'] is not None else 'Absent' if any(k in absent for k in ('EXTERNAL','LAB_EXTERNAL')) else 'Pending',
                row['total'] if row['total'] is not None else 'Pending',row['grade'] if row['grade'] not in {None,'-',''} else 'Pending',
                row['grade_point'] if row['grade_point'] is not None else 'Pending',row['result'],', '.join(absent)])
    tables['Students']=(headers+['Actual admission year','Term roll number','Term division','Practical batch','Placement source','Registered credits','Earned local credits','Local SGPA','Local result'],profile_rows)
    tables['Local Results']=(headers+['Teaching year','Semester','Subject','Code','Credits','Internal','External','Total','Grade','Grade point','Local result','Absent components'],result_rows)
    official=[]
    latest=set()
    for link in UniversityLinkedResult.query.filter_by(**c).filter(UniversityLinkedResult.prn.in_(prns)).order_by(UniversityLinkedResult.id.desc()):
        earlier=link.attempt_key in latest
        latest.add(link.attempt_key)
        source=link.source
        official.append(identity(len(official)+1,link.prn)+[link.subject_code,source['subject_name'],source['subject_code'],
            link.publication.examination_session,link.publication.examination_type,
            float(source['total']) if source.get('total') is not None else source.get('total_raw') or 'Not printed',source.get('grade_display') or 'NOT PRINTED',
            float(source['grade_point']) if source.get('grade_point') is not None else 'Not printed',source['student_status'],
            float(source['sgpa']) if source.get('sgpa') is not None else 'Not printed','Earlier source revision' if earlier else 'Current source revision',
            link.publication.metadata_snapshot['filename'],link.publication.source_hash,json.dumps(source['components'],ensure_ascii=False)])
    tables['University Evidence']=(headers+['Local code','Printed subject','Printed code','Examination session','Examination type','Printed total','Printed grade','Printed grade point',
        'Source status','Printed SGPA','Source revision','Source PDF','Source SHA256','Printed components'],official)
    attendance=[]
    codes={e.subject_code for e in Enrollment.query.filter_by(**c).filter(Enrollment.prn.in_(prns))}
    for code in sorted(codes):
        subject=db.session.get(Subject,code)
        for row in report_rows(dict(c,subject_code=code,division='ALL',batch='ALL')):
            prn=row['student'].prn
            if prn not in prns:
                continue
            attendance.append(identity(len(attendance)+1,prn)+[subject.subject_name,code,row['present'],row['total']-row['present'],row['leave'],row['official'],row['total'],
                float(row['percentage'])/100 if row['percentage'] is not None else 'Pending','Review conflicts' if row['conflicts'] else 'Recorded' if row['total'] else 'Pending'])
    tables['Attendance']=(headers+['Subject','Code','Present','Absent','Leave','Official duty','Counted classes','Attendance percentage','Status'],attendance)
    curriculum=[]
    for record in StudentCurriculumTerm.query.filter_by(**c).filter(StudentCurriculumTerm.prn.in_(prns)).order_by(StudentCurriculumTerm.prn):
        curriculum.append(identity(len(curriculum)+1,record.prn)+[record.binding.version.title,record.binding.version.scheme,record.binding.regular_cohort,record.binding.group_name])
    tables['Curriculum']=(headers+['Approved version','Scheme','Regular cohort','Group'],curriculum)
    corrections=[]
    for r in MarkRevision.query.filter(MarkRevision.prn.in_(prns)).order_by(MarkRevision.created_at):
        if any(r.context.get(k)!=v for k,v in c.items()):
            continue
        corrections.append(identity(len(corrections)+1,r.prn)+[r.context['subject_code'],r.category,r.component,r.status,
            json.dumps(r.previous,default=str),json.dumps(r.proposed,default=str),r.reason,r.decision_note or '',str(r.created_at)])
    policies=AssessmentPolicy.query.filter_by(**c).order_by(AssessmentPolicy.subject_code).all()
    for change in AssessmentScoreChange.query.filter(AssessmentScoreChange.policy_token.in_([p.token for p in policies]),AssessmentScoreChange.prn.in_(prns)).order_by(AssessmentScoreChange.id):
        policy=next(p for p in policies if p.token==change.policy_token)
        corrections.append(identity(len(corrections)+1,change.prn)+[policy.subject_code,'CONFIGURED','All components','APPROVED',
            json.dumps(change.previous,ensure_ascii=False),json.dumps(change.approved,ensure_ascii=False),change.reason,'',str(change.created_at)])
    tables['Corrections']=(headers+['Code','Category','Component','Status','Previous evidence','Proposed / approved evidence','Reason','Decision note','Recorded at'],corrections)
    outcomes=[]
    for row in CourseOutcomeSelection.query.filter_by(semester_id=c['semester_id']).filter(CourseOutcomeSelection.prn.in_(prns)).order_by(CourseOutcomeSelection.prn,CourseOutcomeSelection.subject_code):
        evidence=row.evidence
        outcomes.append(identity(len(outcomes)+1,row.prn)+[evidence['subject_name'],row.subject_code,evidence['credits'],evidence['source'],evidence['teaching_year'],evidence['grade'],
            evidence['grade_point'],evidence['result'],row.source_key,row.revision,str(row.approved_at)])
    tables['Approved Outcomes']=(headers+['Subject','Code','Credits','Selected source','Selected teaching year','Grade','Grade point','Result','Evidence reference','Revision','Approved at'],outcomes)
    decisions=[]
    for row in ProgressionDecision.query.filter_by(**c).filter(ProgressionDecision.prn.in_(prns)).order_by(ProgressionDecision.prn,ProgressionDecision.revision):
        s=row.summary
        decisions.append(identity(len(decisions)+1,row.prn)+[row.decision,row.revision,row.target_year,s['registered_credits'],s['earned_credits'],s['backlog_subjects'],s['backlog_credits'],s['pending_courses'],row.reason,str(row.created_at)])
    tables['Progression']=(headers+['Decision','Revision','Next teaching year','Reviewed credits','Earned credits','Backlog subjects','Backlog credits','Pending courses','Reason','Approved at'],decisions)
    policy_rows=[]
    for policy in policies:
        components='\n'.join(f"{v['code']}: {v['label']} / {v['maximum']}; minimum {v.get('minimum',0)}; absence fails: {'yes' if v.get('absence_fails') else 'no'}" for v in policy.rules['components'])
        groups='\n'.join(f"{v['label']}: {v['method']} {v.get('best_n','')} of {', '.join(v['components'])}; output / {v['maximum']}; minimum {v.get('minimum',0)}; {v['category']}" for v in policy.rules['groups'])
        grades='\n'.join(f"{v['grade']}: from {v['minimum_percent']}%; {v['point']} points" for v in policy.rules['grades'])
        policy_rows.append([policy.subject_code,policy.title,policy.status,sum(float(g['maximum']) for g in policy.rules['groups']),
            policy.rules['total_minimum'],components,groups,grades,policy.source_reference])
    tables['Assessment Rules']=(['Code','Policy','Status','Total maximum','Total minimum','Component definitions','Groups','Grade bands','Source reference'],policy_rows)
    components=[]
    for score in AssessmentScore.query.filter(AssessmentScore.policy_token.in_([p.token for p in policies]),AssessmentScore.prn.in_(prns)).order_by(AssessmentScore.prn):
        policy=next(p for p in policies if p.token==score.policy_token)
        for component in policy.rules['components']:
            value=score.values.get(component['code'],{})
            components.append(identity(len(components)+1,score.prn)+[policy.subject_code,component['label'],component['code'],component['maximum'],
                float(value['value']) if value.get('value') is not None else 'Absent' if value.get('state')=='ABSENT' else 'Pending',value.get('state','PENDING'),score.revision])
    tables['Assessment Components']=(headers+['Subject code','Assessment','Component code','Maximum','Recorded mark','State','Revision'],components)
    portfolios=[]
    for row in StudentPortfolio.query.filter(StudentPortfolio.prn.in_(prns)).order_by(StudentPortfolio.prn):
        texts=[str(row.details.get(key) or '') for key in ('interests','skills','projects','certifications','achievements','goals')]
        for part in range(max(1,max((len(text)+599)//600 for text in texts))):
            portfolios.append(identity(len(portfolios)+1,row.prn)+[part+1]+[text[part*600:(part+1)*600] for text in texts])
    tables['Portfolios']=(headers+['Entry part','Interests','Skills','Projects','Certifications','Achievements','Goals'],portfolios)
    book=Workbook()
    summary=book.active
    summary.title='Summary'
    scope=f"Semester {c['semester_id']} / Teaching year {c['academic_year']}"
    summary.append(['ACADPULSE - Department master workbook'])
    summary.append(['Scope',scope])
    summary.append(['Actual admission cohort',raw_cohort or 'All enrolled cohorts'])
    summary.append(['Term division / batch',str(filters.get('division') or 'All')+' / '+str(filters.get('batch') or 'All')])
    summary.append(['Generated at',datetime.now(ZoneInfo('Asia/Kolkata')).strftime('%d %b %Y %H:%M IST')])
    summary.append(['Interpretation','Local calculations, printed university evidence and administrator-approved course selections are separate. Pending is not zero.'])
    summary.append(['Progression','Only approved institution policies determine progression. Next-term subject enrollment is a separate curriculum review.'])
    summary.append(['Portfolio','Recorded notes; certification entries are not credential verification.'])
    summary.append([])
    summary.append(['Measure','Value'])
    formulas={
        11:('Enrolled students',f'=COUNTA(Students!B2:B{len(profile_rows)+1})'),
        12:('Complete local passes',f'=COUNTIF(Students!L2:L{len(profile_rows)+1},"PASS")'),
        13:('Complete local failures',f'=COUNTIF(Students!L2:L{len(profile_rows)+1},"FAIL")'),
        14:('Pending local results',f'=COUNTIF(Students!L2:L{len(profile_rows)+1},"PENDING")'),
    }
    for _,(label,formula) in formulas.items():
        summary.append([label,formula])
    summary.append(['Linked university evidence rows',len(official)])
    summary.append(['Approved course selections (this semester)',len(outcomes)])
    summary.column_dimensions['A'].width=46
    summary.column_dimensions['B'].width=100
    summary.freeze_panes='B11'
    for title,(columns,rows) in tables.items():
        sheet=book.create_sheet(title)
        sheet.append(columns)
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes='D2' if columns[:3]==headers else 'B2'
        sheet.auto_filter.ref=sheet.dimensions
        for i,label in enumerate(columns,1):
            sheet.column_dimensions[get_column_letter(i)].width=(10 if label=='Sr. No.' else 24 if label=='PRN' else 34 if label=='Student Name' else
                45 if any(v in label.lower() for v in ('subject','reference','reason','groups','evidence','bands','source','projects','certifications','definitions')) else 24)
        sheet.sheet_view.showGridLines=False
        sheet.print_title_rows='1:1'
        sheet.sheet_properties.pageSetUpPr.fitToPage=True
        sheet.page_setup.orientation='landscape'
        sheet.page_setup.paperSize=sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth=1
        sheet.page_setup.fitToHeight=0
        for row in sheet.iter_rows():
            lines=max(sum(max(1,(len(line)+int(sheet.column_dimensions[cell.column_letter].width)-1)//max(1,int(sheet.column_dimensions[cell.column_letter].width)-3))
                for line in str(cell.value if cell.value is not None else '').split('\n')) for cell in row)
            sheet.row_dimensions[row[0].row].height=48 if row[0].row==1 else min(409,max(34,lines*15+8))
            for cell in row:
                if isinstance(cell.value,str):
                    cell.data_type='s'
                cell.font=Font(name='Calibri',size=11,bold=cell.row==1,color='FFFFFF' if cell.row==1 else '172B4D')
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                if cell.row==1:
                    cell.fill=PatternFill('solid',fgColor='14283F')
                elif cell.row%2==0:
                    cell.fill=PatternFill('solid',fgColor='F1F5F9')
        if title=='Attendance':
            for row in range(2,sheet.max_row+1):
                sheet.cell(row,11).number_format='0.0%'
    for row in summary:
        summary.row_dimensions[row[0].row].height=44 if row[0].row in {6,7,8} else 26
        for cell in row:
            cell.font=Font(name='Calibri',size=11,bold=cell.column==1 or cell.row==10)
            cell.alignment=Alignment(vertical='top',wrap_text=True)
            if isinstance(cell.value,str) and cell.data_type!='f':
                cell.data_type='s'
    summary['A1'].font=Font(name='Calibri',size=18,bold=True,color='14283F')
    summary.row_dimensions[1].height=40
    book.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
    stream=BytesIO()
    book.save(stream)
    stream.seek(0)
    return stream
