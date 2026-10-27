import io,re,unittest
from datetime import datetime,timedelta,timezone
from openpyxl import load_workbook
from werkzeug.datastructures import FileStorage
import test_workflows as base
from models import db,Student,Subject,Semester,Enrollment,TheoryMarks,InternalMarksSubmission,User
from teaching_models import TeachingAssignment
from workflow_models import WorkflowReview,AcceptedInternalMark
from workflow_services import roster_template


class MarksReviewWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.fixture=base.WorkflowTestCase();self.fixture.setUp();self.client=self.fixture.client
        self.context={'subject_code':'TH601','semester_id':6,'academic_year':'2025-26','division':'','batch':''}
        self.teacher=User.query.filter_by(username='teacher').one()
        db.session.add(TeachingAssignment(teacher_id=self.teacher.user_id,subject_code='TH601',semester_id=6,academic_year='2025-26'))
        db.session.commit()
    def tearDown(self):self.fixture.tearDown()
    def token(self):
        from flask import g
        g.pop('csrf_token',None)
        page=self.client.get('/profile/password')
        return re.search(rb'name="csrf_token" value="([^"]+)"',page.data).group(1).decode()
    def post(self,url,data):return self.client.post(url,data={**data,'csrf_token':self.token()})
    def marks_preview(self,marks=(0,'A',8),exam='CT1',**context):
        fields={**self.context,**context,'exam_type':exam}
        query='&'.join(f'{k}={v}' for k,v in fields.items())
        template=self.client.get('/download-ct1-template?'+query)
        self.assertEqual(template.status_code,200,template.data[:100])
        data=self.fixture._fill_ct_template(template,marks)
        return self.post('/workflow/marks/preview',{**fields,'ct1_file':(io.BytesIO(data),'marks.xlsx')})
    def confirm(self,response):
        self.assertEqual(response.status_code,302,response.data)
        self.assertIn('/workflow/reviews/',response.location)
        return self.post(response.location+'/confirm',{})

    def test_teacher_identity_is_automatic_and_url_cannot_impersonate(self):
        self.fixture.logout();self.fixture.login('teacher','teacher123')
        page=self.client.get('/ct1?subject_code=TH601&semester_id=6&academic_year=2025-26&exam_type=CT1')
        self.assertIn(b'Signed-in teacher',page.data);self.assertNotIn(b'Enter full name of subject teacher',page.data)
        response=self.client.get('/download-ct1-template?subject_code=TH601&semester_id=6&academic_year=2025-26&exam_type=CT1&subject_teacher=Impostor')
        book=load_workbook(io.BytesIO(response.data))
        values=[cell.value for row in book.active for cell in row]
        self.assertIn('teacher',values);self.assertNotIn('Impostor',values)

    def test_preview_saves_nothing_then_confirmation_records_zero_and_absence_and_locks(self):
        response=self.marks_preview()
        self.assertEqual(TheoryMarks.query.count(),0);self.assertEqual(InternalMarksSubmission.query.count(),0)
        page=self.client.get(response.location);self.assertIn(b'No marks have been saved yet',page.data)
        self.assertEqual(self.confirm(response).status_code,302)
        self.assertEqual(TheoryMarks.query.filter_by(prn='S001').one().ct1,0)
        self.assertIsNone(TheoryMarks.query.filter_by(prn='S002').one().ct1)
        self.assertEqual(AcceptedInternalMark.query.filter_by(prn='S002').one().state,'ABSENT')
        self.assertEqual(AcceptedInternalMark.query.count(),3)
        self.assertEqual(self.confirm(response).status_code,302);self.assertEqual(AcceptedInternalMark.query.count(),3)
        duplicate=self.marks_preview();self.assertNotIn('/workflow/reviews/',duplicate.location)
        self.assertEqual(TheoryMarks.query.filter_by(prn='S001').one().ct1,0)

    def test_blank_formula_out_of_range_and_missing_row_are_rejected_without_writes(self):
        for values in [(None,5,6),('=1+1',5,6),(11,5,6)]:
            response=self.marks_preview(values)
            self.assertNotIn('/workflow/reviews/',response.location)
        self.assertEqual(TheoryMarks.query.count(),0);self.assertEqual(WorkflowReview.query.count(),0)

    def test_roster_change_or_existing_mark_after_preview_blocks_commit_atomically(self):
        response=self.marks_preview((7,8,9))
        db.session.add(Student(prn='NEW',name='New student',year_of_admission=2023))
        db.session.add(Enrollment(prn='NEW',subject_code='TH601',semester_id=6,academic_year='2025-26'));db.session.commit()
        self.assertEqual(self.confirm(response).status_code,409)
        self.assertEqual(TheoryMarks.query.count(),0);self.assertEqual(AcceptedInternalMark.query.count(),0)

    def test_revoked_teacher_and_other_teacher_cannot_confirm(self):
        self.fixture.logout();self.fixture.login('teacher','teacher123')
        response=self.marks_preview((7,8,9))
        assignment=TeachingAssignment.query.one();assignment.is_active=False;db.session.commit()
        self.assertEqual(self.confirm(response).status_code,403);self.assertEqual(TheoryMarks.query.count(),0)
        assignment.is_active=True
        other=User(username='other_teacher',role='TEACHER',is_active=True);other.set_password('other-password');db.session.add(other);db.session.flush()
        db.session.add(TeachingAssignment(teacher_id=other.user_id,subject_code='TH601',semester_id=6,academic_year='2025-26'));db.session.commit()
        self.fixture.logout();self.fixture.login('other_teacher','other-password')
        self.assertEqual(self.client.get(response.location).status_code,403)
        self.assertEqual(self.confirm(response).status_code,403)

    def test_expired_preview_is_rejected(self):
        response=self.marks_preview((7,8,9));review=WorkflowReview.query.one()
        review.expires_at=datetime.now(timezone.utc)-timedelta(minutes=1);db.session.commit()
        self.assertEqual(self.confirm(response).status_code,409)
        self.assertEqual(TheoryMarks.query.count(),0)

    def test_subject_analysis_shows_partial_marks_without_external_results(self):
        self.confirm(self.marks_preview())
        page=self.client.get('/workflow/subject?subject_code=TH601&semester_id=6&academic_year=2025-26')
        self.assertEqual(page.status_code,200,page.data)
        self.assertIn(b'Asha Patil',page.data);self.assertIn(b'Absent',page.data);self.assertIn(b'Pending',page.data)
        self.assertIn(b'4.00',page.data) # numeric average includes zero and excludes absent.
        self.assertNotIn(b'Select semester...',page.data)

    def test_subject_analysis_zero_roster_and_internal_reports_show_setup_guidance(self):
        Enrollment.query.filter_by(subject_code='TH601').delete();db.session.commit()
        url='?subject_code=TH601&semester_id=6&academic_year=2025-26'
        page=self.client.get('/workflow/subject'+url)
        self.assertEqual(page.status_code,200);self.assertIn(b'No students enrolled',page.data)
        self.assertIn(b'Configure student roster',page.data)
        self.assertEqual(self.client.get('/download-internal-report/pdf'+url).status_code,302)

    def test_internal_pdf_and_excel_preserve_pending_instead_of_inventing_zero(self):
        self.confirm(self.marks_preview())
        query='?subject_code=TH601&semester_id=6&academic_year=2025-26'
        response=self.client.get('/download-internal-report/excel'+query)
        self.assertEqual(response.status_code,200)
        book=load_workbook(io.BytesIO(response.data));values=[c.value for row in book.active for c in row]
        self.assertIn('Pending',values);self.assertIn('Absent',values)
        response=self.client.get('/download-internal-report/pdf'+query)
        self.assertEqual(response.status_code,200)
        import fitz
        text=' '.join(p.get_text() for p in fitz.open(stream=response.data,filetype='pdf'))
        self.assertIn('Pending',text);self.assertIn('Absent',text)

    def test_roster_existing_selection_previews_then_enrolls_only_selected_students(self):
        Enrollment.query.filter_by(subject_code='TH601').delete();db.session.commit()
        response=self.post('/workflow/roster/preview',{**self.context,'student_prns':['S001','S003']})
        self.assertEqual(Enrollment.query.filter_by(subject_code='TH601').count(),0)
        self.assertEqual(self.confirm(response).status_code,302)
        self.assertEqual(Enrollment.query.filter_by(subject_code='TH601').count(),2)
        self.assertEqual(self.confirm(response).status_code,302)
        self.assertEqual(Enrollment.query.filter_by(subject_code='TH601').count(),2)
        self.assertEqual(Student.query.count(),3)

    def test_roster_new_student_import_preserves_text_prn_and_does_not_write_before_confirm(self):
        stream=roster_template(self.context);book=load_workbook(stream)
        sheet=book.active;sheet.delete_rows(2,1000);sheet.append(['0000000000000123','New Real Student',2023,'B','B2'])
        result=io.BytesIO();book.save(result);result.seek(0)
        response=self.post('/workflow/roster/preview',{**self.context,'roster_file':(result,'roster.xlsx')})
        self.assertIsNone(db.session.get(Student,'0000000000000123'))
        self.assertEqual(self.confirm(response).status_code,302)
        self.assertEqual(db.session.get(Student,'0000000000000123').name,'New Real Student')
        self.assertIsNotNone(Enrollment.query.filter_by(prn='0000000000000123',subject_code='TH601').first())

    def test_roster_conflicting_identity_and_missing_admission_year_are_blocked(self):
        from workflow_services import validate_roster_rows
        for row in [{'prn':'S001','name':'Different Person','admission_year':2023,'division':'','batch':''},
                    {'prn':'NEW','name':'New Real Student','admission_year':None,'division':'','batch':''}]:
            with self.assertRaises(ValueError):validate_roster_rows(self.context,[row])
        self.assertEqual(Student.query.count(),3)

    def test_teacher_cannot_manage_roster(self):
        self.fixture.logout();self.fixture.login('teacher','teacher123')
        self.assertEqual(self.client.get('/workflow/roster').status_code,403)
        self.assertEqual(self.post('/workflow/roster/preview',{**self.context,'student_prns':['S001']}).status_code,403)

    def test_original_workbook_is_retained_and_download_authorized(self):
        from hashlib import sha256
        self.fixture.logout();self.fixture.login('teacher','teacher123')
        response=self.marks_preview((7,8,9));review=WorkflowReview.query.one()
        downloaded=self.client.get(response.location+'/source')
        self.assertEqual(downloaded.status_code,200)
        self.assertEqual(sha256(downloaded.data).hexdigest(),review.source_hash)
        self.assertEqual(downloaded.data,review.source_workbook)
        self.fixture.logout()
        self.assertEqual(self.client.get(response.location+'/source').status_code,302)

    def test_overlapping_scope_cannot_overwrite_previously_confirmed_batch(self):
        response=self.marks_preview((7,),division='A',batch='A1')
        self.assertEqual(self.confirm(response).status_code,302)
        response=self.marks_preview((1,2,3))
        self.assertNotIn('/workflow/reviews/',response.location)
        self.assertEqual(TheoryMarks.query.filter_by(prn='S001').one().ct1,7)
        self.assertEqual(AcceptedInternalMark.query.count(),1)

    def test_roster_changed_profile_after_preview_blocks_confirmation(self):
        response=self.post('/workflow/roster/preview',{**self.context,'student_prns':['S001']})
        student=db.session.get(Student,'S001');student.name='Changed after preview';db.session.commit()
        self.assertEqual(self.confirm(response).status_code,409)
        self.assertEqual(db.session.get(Student,'S001').name,'Changed after preview')

    def test_review_and_roster_confirm_require_csrf(self):
        response=self.marks_preview((7,8,9))
        self.assertEqual(self.client.post(response.location+'/confirm').status_code,400)
        self.assertEqual(TheoryMarks.query.count(),0)

    def test_confirmed_university_pdf_prefills_only_identity_and_requires_actual_admission_year(self):
        job=self.fixture._create_result_pdf_dataset()
        job.status='READY';job.confirmed_at=datetime.now(timezone.utc);db.session.commit()
        stream=roster_template(self.context,job.id);book=load_workbook(stream)
        self.assertEqual(book.active['A2'].value,sorted(job.report.students,key=lambda s:s.seat_number)[0].seat_number)
        # The source cannot invent admission years or fill local coursework.
        self.assertEqual(TheoryMarks.query.count(),0)
        self.assertEqual(Enrollment.query.filter_by(subject_code='TH601').count(),3)
        job.confirmed_at=None;db.session.commit()
        with self.assertRaises(ValueError):roster_template(self.context,job.id)

    def test_missing_student_and_name_prn_mismatch_block_preview(self):
        query='?subject_code=TH601&semester_id=6&academic_year=2025-26&exam_type=CT1'
        for mismatch in (False,True):
            template=self.client.get('/download-ct1-template'+query)
            book=load_workbook(io.BytesIO(self.fixture._fill_ct_template(template,[7,8,9])))
            if mismatch:book.active.cell(11,4,'Wrong Student Name')
            else:book.active.cell(13,3,None);book.active['C13']=None
            stream=io.BytesIO();book.save(stream);stream.seek(0)
            response=self.post('/workflow/marks/preview',{**self.context,'exam_type':'CT1','ct1_file':(stream,'marks.xlsx')})
            self.assertNotIn('/workflow/reviews/',response.location)
        self.assertEqual(TheoryMarks.query.count(),0)


if __name__=='__main__':unittest.main()
