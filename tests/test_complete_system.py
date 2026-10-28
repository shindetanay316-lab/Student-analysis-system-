import copy
import csv
import io
import json
import re
import unittest
from decimal import Decimal
from datetime import timedelta
from openpyxl import load_workbook
import test_workflows as base
from models import db, Student, Subject, Semester, Enrollment, TheoryMarks, LabMarks
from completion_models import StudentTerm, PendingAdmission, AssessmentPolicy, AssessmentScore, AssessmentScoreChange, ProgressionPolicy, ProgressionDecision, CourseOutcomeSelection
from workflow_models import WorkflowReview, now
from teaching_models import TeachingAssignment
import assessment_services as assessments
import progression_services as progression
import term_services as terms


class CompleteSystemTests(unittest.TestCase):
    def setUp(self):
        self.fixture=base.WorkflowTestCase()
        self.fixture.setUp()
        self.client=self.fixture.client
        db.session.add(Semester(semester_id=7,academic_year='2026-27',is_active=False))
        db.session.commit()
        self.c=dict(semester_id=6,academic_year='2025-26')

    def tearDown(self):
        self.fixture.tearDown()

    def post(self,url,fields=None):
        from flask import g
        g.pop('csrf_token',None)
        token=re.search(rb'name="csrf_token" value="([^"]+)"',self.client.get('/profile/password').data).group(1).decode()
        return self.client.post(url,data={**(fields or {}),'csrf_token':token})

    def review(self,response):
        self.assertEqual(response.status_code,302,response.data)
        self.assertIn('/records/reviews/',response.location)
        return WorkflowReview.query.filter_by(token=response.location.rsplit('/',1)[-1]).one()

    def confirm(self,response):
        return self.post(response.location+'/confirm',dict(verified='on'))

    def policy(self,rules=None,approve=True):
        response=self.post('/records/assessments',dict(self.c,subject_code='TH601',title='Reviewed scheme',
            source_reference='Test syllabus, assessment scheme page 12.',rules=json.dumps(rules or assessments.DEFAULT_RULES)))
        self.assertEqual(response.status_code,302)
        policy=AssessmentPolicy.query.one()
        if approve:
            self.post(f'/records/assessments/{policy.token}/approve',dict(revision=policy.revision,verified='on'))
            self.assertEqual(policy.status,'APPROVED')
        return policy

    def score_csv(self,prn='S001',values=('8','9','7','15','42')):
        stream=io.StringIO()
        writer=csv.writer(stream)
        writer.writerow(['ACADPULSE configured assessment marks'])
        writer.writerows(assessments.template_metadata(AssessmentPolicy.query.one()).items())
        writer.writerow([])
        writer.writerows([['Sr. No.','PRN','Student Name','CT1','CT2','ASSIGNMENT','MSE','ESE'],[1,prn,db.session.get(Student,prn).name,*values]])
        return stream.getvalue().encode()

    def score_preview(self,policy,values=('8','9','7','15','42'),reason=''):
        return self.post(f'/records/assessments/{policy.token}/preview',dict(marks_file=(io.BytesIO(self.score_csv(values=values)),'marks.csv'),reason=reason))

    def progression_policy(self,**options):
        fields=dict(title='Verified progression scheme',from_semester=6,reviewed_semesters='6',maximum_backlog_subjects=0,
            maximum_backlog_credits=0,minimum_earned_percent=100,source_reference='Institution ordinance verified clause 7.',verified='on')
        fields.update(options)
        self.post('/records/progression/policies',fields)
        return ProgressionPolicy.query.order_by(ProgressionPolicy.approved_at.desc()).first()

    def only_theory(self):
        Enrollment.query.filter(Enrollment.prn=='S001',Enrollment.subject_code!='TH601').delete(synchronize_session=False)
        db.session.commit()

    def complete_local(self,external=42):
        self.only_theory()
        db.session.add(TheoryMarks(prn='S001',subject_code='TH601',**self.c,ct1=8,ct2=9,assignment=7,midsem=15,external=external))
        db.session.commit()

    def progression_preview(self,policy,**options):
        return self.post('/records/progression/preview',dict(prn='S001',policy=policy.token,academic_year='2025-26',target_year='2026-27',
            reason='Verified all available course attempts and institution rule.',**options))

    def test_all_new_screens_and_teacher_admin_boundaries(self):
        for url in ['/records/terms','/records/terms?semester_id=6&academic_year=2025-26','/records/admissions','/records/assessments','/records/progression','/records/master']:
            self.assertEqual(self.client.get(url).status_code,200,url)
        self.fixture.logout();self.fixture.login('teacher','teacher123')
        for url in ['/records/terms','/records/admissions','/records/progression','/records/master','/records/master/excel?semester_id=6&academic_year=2025-26']:
            self.assertEqual(self.client.get(url).status_code,403,url)
        self.assertEqual(self.client.get('/records/assessments').status_code,200)

    def test_term_placements_preview_and_confirm_preserve_other_year_and_identity(self):
        db.session.add(StudentTerm(prn='S001',semester_id=5,academic_year='2024-25',division='OLD',batch='OLD1',roll_number='17'))
        db.session.commit()
        response=self.post('/records/terms/preview',dict(self.c,prns=['S001'],roll_S001='01',division_S001='B',batch_S001='B2'))
        self.review(response)
        self.assertEqual(terms.placement('S001',6,'2025-26').division,'A')
        self.assertEqual(self.confirm(response).status_code,302)
        self.assertEqual(self.confirm(response).status_code,302)
        self.assertEqual(terms.placement('S001',6,'2025-26').division,'B')
        self.assertEqual(terms.placement('S001',5,'2024-25').division,'OLD')
        self.assertEqual(db.session.get(Student,'S001').division,'A')
        self.assertEqual(db.session.get(Student,'S001').year_of_admission,2023)
        self.assertEqual(self.client.get('/workflow/reviews/'+self.review(response).token).status_code,404)

    def test_term_filter_and_attendance_use_historical_placement(self):
        terms.placement('S001',6,'2025-26').division='C'
        terms.placement('S001',6,'2025-26').batch='C4'
        db.session.get(Student,'S001').division='Z'
        db.session.commit()
        import workflow_services as core
        c=dict(self.c,subject_code='TH601',division='C',batch='C4')
        self.assertEqual([s.prn for s in core.roster(c)],['S001'])
        from attendance_utils import students_for_attendance
        self.assertEqual([s.prn for s in students_for_attendance('TH601',6,'2025-26','C','C4')],['S001'])
        self.assertEqual(core.roster(dict(c,division='Z')),[])
        for url in ['/ct1','/lab','/external']:
            response=self.client.get(url,query_string=self.c)
            self.assertEqual(response.status_code,200)
            self.assertIn(b'value="C"',response.data)
            self.assertIn(b'value="C4"',response.data)

    def test_duplicate_roll_rejected_and_roll_swap_allowed(self):
        a,b=terms.placement('S001',6,'2025-26'),terms.placement('S002',6,'2025-26')
        a.roll_number,b.roll_number='01','02';db.session.commit()
        bad=self.post('/records/terms/preview',dict(self.c,prns=['S001'],roll_S001='02',division_S001='A',batch_S001='A1'))
        self.assertNotIn('/reviews/',bad.location)
        response=self.post('/records/terms/preview',dict(self.c,prns=['S001','S002'],roll_S001='02',division_S001='A',batch_S001='A1',
            roll_S002='01',division_S002='A',batch_S002='A2'))
        self.assertEqual(self.confirm(response).status_code,302)
        self.assertEqual((a.roll_number,b.roll_number),('02','01'))

    def test_stale_placement_preview_and_expiry_fail_atomically(self):
        response=self.post('/records/terms/preview',dict(self.c,prns=['S001'],division_S001='B',batch_S001='B1'))
        terms.placement('S001',6,'2025-26').revision+=1;db.session.commit()
        self.assertEqual(self.confirm(response).status_code,409)
        self.assertEqual(terms.placement('S001',6,'2025-26').division,'A')
        response=self.post('/records/terms/preview',dict(self.c,prns=['S001'],division_S001='B',batch_S001='B1'))
        self.review(response).expires_at=now()-timedelta(seconds=1);db.session.commit()
        self.assertEqual(self.confirm(response).status_code,409)

    def test_pending_admission_activation_preserves_real_admission_and_requires_real_unique_prn(self):
        self.post('/records/admissions',dict(self.c,admission_number='ADM-005',name='New Admission',admission_year=2024,division='D',batch='D1',roll_number='05'))
        item=PendingAdmission.query.one()
        self.assertIsNone(db.session.get(Student,'ADM-005'))
        bad=self.post(f'/records/admissions/{item.token}/preview',dict(prn='S001'))
        self.assertNotIn('/reviews/',bad.location)
        response=self.post(f'/records/admissions/{item.token}/preview',dict(prn='REAL005'))
        self.assertIsNone(db.session.get(Student,'REAL005'))
        self.assertEqual(self.confirm(response).status_code,302)
        self.assertEqual(self.confirm(response).status_code,302)
        student=db.session.get(Student,'REAL005')
        self.assertEqual(student.year_of_admission,2024)
        self.assertEqual(terms.placement('REAL005',6,'2025-26').roll_number,'05')
        self.assertEqual(Enrollment.query.filter_by(prn='REAL005').count(),0)

    def test_pending_admission_edit_invalidates_activation_preview(self):
        self.post('/records/admissions',dict(self.c,admission_number='ADM-1',name='New Admission',admission_year=2024))
        item=PendingAdmission.query.one()
        response=self.post(f'/records/admissions/{item.token}/preview',dict(prn='REAL1'))
        self.post(f'/records/admissions/{item.token}/edit',dict(self.c,name='Corrected Admission',admission_year=2024,revision=1))
        self.assertEqual(self.confirm(response).status_code,409)
        self.assertIsNone(db.session.get(Student,'REAL1'))

    def test_rule_validation_rejects_ambiguous_invalid_scales_and_nonfinite_marks(self):
        for modify in [lambda r:r['components'][0].update(maximum='NaN'),
            lambda r:r['components'][0].update(maximum=-1),lambda r:r['groups'][0].update(best_n=5),
            lambda r:r['groups'][1].update(components=['CT1']),lambda r:r['grades'][-1].update(minimum_percent=20),
            lambda r:r['grades'][0].update(point=11),lambda r:r['components'][0].update(absence_fails='yes')]:
            rules=copy.deepcopy(assessments.DEFAULT_RULES);modify(rules)
            with self.assertRaises(ValueError):assessments.validate_rules(rules)

    def test_assessment_draft_edit_and_immutable_approval(self):
        policy=self.policy(approve=False)
        rules=copy.deepcopy(policy.rules);rules['components'][-1]['minimum']=25
        self.post(f'/records/assessments/{policy.token}/edit',dict(revision=1,title='Corrected scheme',source_reference='Corrected source, page 13.',rules=json.dumps(rules)))
        self.assertEqual(policy.revision,2)
        self.post(f'/records/assessments/{policy.token}/approve',dict(revision=1,verified='on'))
        self.assertEqual(policy.status,'DRAFT')
        self.post(f'/records/assessments/{policy.token}/approve',dict(revision=2,verified='on'))
        self.assertEqual(policy.status,'APPROVED')
        self.post(f'/records/assessments/{policy.token}/edit',dict(revision=2,title='Should not change',source_reference='Another source reference.',rules=json.dumps(assessments.DEFAULT_RULES)))
        self.assertEqual(policy.title,'Corrected scheme')

    def test_existing_legacy_marks_block_new_policy_without_changing_values(self):
        self.complete_local()
        self.post('/records/assessments',dict(self.c,subject_code='TH601',title='Unsafe migration',source_reference='Test reference source.',rules=json.dumps(assessments.DEFAULT_RULES)))
        self.assertEqual(AssessmentPolicy.query.count(),0)
        self.assertEqual(TheoryMarks.query.one().external,42)

    def test_configured_score_review_calculates_best_two_and_integrates_history_and_dashboard(self):
        self.only_theory()
        policy=self.policy()
        response=self.score_preview(policy)
        self.review(response)
        self.assertEqual(AssessmentScore.query.count(),0)
        self.assertEqual(self.client.get(response.location).status_code,200)
        self.assertEqual(self.confirm(response).status_code,302)
        self.assertEqual(self.confirm(response).status_code,302)
        self.assertEqual(AssessmentScoreChange.query.count(),1)
        result=assessments.result_for(policy,'S001')
        self.assertEqual((result['internal'],result['total'],result['grade']),(32,74,'BC'))
        from final_result_utils import build_student_semester_result
        semester=build_student_semester_result('S001',6,'2025-26')
        self.assertEqual((semester['sgpa'],semester['result']),(7.5,'PASS'))
        self.assertIn(b'74',self.client.get('/academic/students/S001').data)
        self.assertEqual(self.client.get(f'/records/assessments/{policy.token}').status_code,200)
        from routes_analytics import _completed_semester_rows,_pending_count_for_semester
        self.assertEqual(len(_completed_semester_rows(6,'2025-26')),1)
        self.assertEqual(_pending_count_for_semester(6,'2025-26'),5)  # Other two students still have two courses plus S002's elective.

    def test_custom_scale_groups_and_separate_minimum(self):
        rules=dict(components=[dict(code='QUIZ',label='Quiz',maximum=20),dict(code='FINAL',label='Final',maximum=30,minimum=12,absence_fails=True)],
            groups=[dict(label='Internal',components=['QUIZ'],method='SUM',maximum=20,category='INTERNAL'),dict(label='External',components=['FINAL'],method='SUM',maximum=30,category='EXTERNAL')],
            total_minimum=20,grades=[dict(minimum_percent=80,grade='A',point=10),dict(minimum_percent=0,grade='P',point=5)])
        policy=self.policy(rules)
        result=assessments.evaluate(policy,dict(QUIZ=dict(value='20',state='PRESENT'),FINAL=dict(value='11',state='PRESENT')))
        self.assertEqual((result['total'],result['result']),(31,'FAIL'))
        result=assessments.evaluate(policy,dict(QUIZ=dict(value='20',state='PRESENT'),FINAL=dict(value='20',state='PRESENT')))
        self.assertEqual((result['total'],result['grade']),(40,'A'))

    def test_zero_pending_and_absence_are_distinct_and_dropped_absence_can_pass(self):
        policy=self.policy()
        for values,status in [(('0','9','7','15','42'),'PASS'),(('AB','9','7','15','42'),'PASS'),(('8','9','7','15','AB'),'FAIL'),(('8','9','7','','42'),'PENDING')]:
            rows=assessments.parse_scores(policy,self.score_csv(values=values))
            self.assertEqual(rows[0]['result']['result'],status,values)

    def test_legacy_entry_redirect_and_staged_legacy_write_rejected(self):
        policy=self.policy()
        query=dict(self.c,subject_code='TH601')
        response=self.client.get('/workflow/subject',query_string=query)
        self.assertEqual(response.location,'/records/assessments/'+policy.token)
        import workflow_services as core
        review=WorkflowReview(owner_id=1,
            kind='MARKS',context=dict(query,exam_type='CT1'),rows=[],snapshot='x'*64)
        with self.assertRaises(ValueError):core.confirm(review)
        self.assertEqual(TheoryMarks.query.count(),0)

    def test_score_corrections_require_admin_and_preserve_original_csv_and_history(self):
        policy=self.policy()
        self.confirm(self.score_preview(policy))
        db.session.add(TeachingAssignment(teacher_id=2,subject_code='TH601',**self.c));db.session.commit()
        self.fixture.logout();self.fixture.login('teacher','teacher123')
        response=self.score_preview(policy,values=('8','9','7','15','43'),reason='Corrected misread examination component.')
        review=self.review(response)
        self.assertTrue(review.context['correction'])
        self.assertEqual(self.confirm(response).status_code,403)
        self.assertEqual(AssessmentScore.query.one().values['ESE']['value'],'42')
        self.fixture.logout();self.fixture.login('admin','admin123')
        self.assertEqual(self.confirm(response).status_code,302)
        self.assertEqual(AssessmentScore.query.one().values['ESE']['value'],'43')
        self.assertEqual(AssessmentScoreChange.query.count(),2)
        self.assertEqual(self.client.get(response.location+'/source').data,self.score_csv(values=('8','9','7','15','43')))

    def test_stale_configured_preview_and_revoked_assignment_are_blocked(self):
        policy=self.policy()
        response=self.score_preview(policy)
        db.session.get(Student,'S001').name='Changed Name';db.session.commit()
        self.assertEqual(self.confirm(response).status_code,409)
        self.assertEqual(AssessmentScore.query.count(),0)
        db.session.get(Student,'S001').name='Asha Patil'
        assignment=TeachingAssignment(teacher_id=2,subject_code='TH601',**self.c);db.session.add(assignment);db.session.commit()
        self.fixture.logout();self.fixture.login('teacher','teacher123')
        response=self.score_preview(policy)
        assignment.is_active=False;db.session.commit()
        self.assertEqual(self.confirm(response).status_code,403)

    def test_pending_progression_records_hold_without_creating_next_term(self):
        policy=self.progression_policy()
        response=self.progression_preview(policy)
        review=self.review(response)
        self.assertEqual(review.context['decision'],'HOLD')
        self.assertEqual(self.client.get(response.location).status_code,200)
        self.assertEqual(self.confirm(response).status_code,302)
        self.assertIsNone(terms.placement('S001',7,'2026-27'))
        self.assertEqual(ProgressionDecision.query.one().summary['pending_courses'],3)

    def test_successful_progression_retains_attempt_snapshot_and_creates_blank_next_placement(self):
        self.complete_local()
        policy=self.progression_policy()
        response=self.progression_preview(policy)
        self.assertEqual(self.review(response).context['decision'],'PROMOTED')
        self.assertEqual(self.confirm(response).status_code,302)
        self.assertEqual(self.confirm(response).status_code,302)
        self.assertEqual(ProgressionDecision.query.count(),1)
        placement=terms.placement('S001',7,'2026-27')
        self.assertIsNone(placement.roll_number)
        self.assertIsNone(placement.division)
        self.assertEqual(CourseOutcomeSelection.query.one().evidence['grade'],'BC')
        self.assertEqual(db.session.get(Student,'S001').year_of_admission,2023)
        self.assertEqual(Enrollment.query.filter_by(prn='S001',semester_id=7).count(),0)
        second=self.progression_preview(policy)
        self.assertNotIn('/reviews/',second.location)

    def test_backlog_limits_and_source_minima_control_progression(self):
        self.complete_local(external=19)
        policy=self.progression_policy()
        response=self.progression_preview(policy)
        review=self.review(response)
        self.assertEqual(review.context['summary']['backlog_subjects'],1)
        self.assertEqual(review.context['decision'],'HOLD')
        self.confirm(response)
        self.assertEqual(CourseOutcomeSelection.query.one().evidence['result'],'FAIL')

    def test_progression_snapshot_rejects_changed_marks_without_partial_approval(self):
        self.complete_local()
        policy=self.progression_policy()
        response=self.progression_preview(policy)
        TheoryMarks.query.one().external=50;db.session.commit()
        self.assertEqual(self.confirm(response).status_code,409)
        self.assertEqual(ProgressionDecision.query.count(),0)
        self.assertEqual(CourseOutcomeSelection.query.count(),0)

    def test_progression_preview_does_not_persist_legacy_derived_mark_refresh(self):
        self.complete_local()
        mark=TheoryMarks.query.one()
        self.assertIsNone(mark.total_marks)
        policy=self.progression_policy()
        self.progression_preview(policy)
        db.session.expire_all()
        self.assertIsNone(TheoryMarks.query.one().total_marks)

    def test_master_workbook_identity_order_formulas_filters_and_real_grades(self):
        self.complete_local()
        policy=self.progression_policy()
        self.confirm(self.progression_preview(policy))
        response=self.client.get('/records/master/excel',query_string=dict(self.c,division='A'))
        self.assertEqual(response.status_code,200)
        book=load_workbook(io.BytesIO(response.data))
        self.assertEqual(len(book.sheetnames),12)
        self.assertEqual([c.value for c in book['Students'][1]][:3],['Sr. No.','PRN','Student Name'])
        self.assertEqual(book['Summary']['B11'].value,'=COUNTA(Students!B2:B3)')
        self.assertEqual(book['Summary']['B12'].data_type,'f')
        headers=[c.value for c in book['Local Results'][1]]
        rows=list(book['Local Results'].iter_rows(min_row=2,values_only=True))
        row=next(r for r in rows if r[1]=='S001')
        self.assertEqual(row[headers.index('Grade')],'BC')
        self.assertNotIn('S003',[r[1] for r in book['Students'].iter_rows(min_row=2,values_only=True)])
        self.assertEqual(self.client.get('/records/master/excel',query_string=dict(self.c,cohort='bad')).status_code,400)

    def test_xlsx_template_zero_context_formula_guard_and_original_retention(self):
        policy=self.policy()
        response=self.client.get(f'/records/assessments/{policy.token}/template')
        book=load_workbook(io.BytesIO(response.data))
        self.assertEqual(book['_meta'].sheet_state,'hidden')
        sheet=book['Configured Marks']
        for i,value in enumerate([0,9,7,15,42],4):
            sheet.cell(2,i).value=value
        stream=io.BytesIO();book.save(stream)
        content=stream.getvalue()
        response=self.post(f'/records/assessments/{policy.token}/preview',dict(marks_file=(io.BytesIO(content),'marks.xlsx')))
        review=self.review(response)
        self.assertEqual(review.rows[0]['values']['CT1']['value'],'0')
        self.assertEqual(self.confirm(response).status_code,302)
        self.assertEqual(self.client.get(response.location+'/source').data,content)
        book['_meta']['B3']='2026-27'
        stream=io.BytesIO();book.save(stream)
        with self.assertRaises(ValueError):assessments.parse_scores(policy,stream.getvalue(),'marks.xlsx')
        book['_meta']['B3']='2025-26';sheet['D2']='=1+1'
        stream=io.BytesIO();book.save(stream)
        with self.assertRaises(ValueError):assessments.parse_scores(policy,stream.getvalue(),'marks.xlsx')

    def test_configured_history_exports_include_components_without_teacher_scope_leak(self):
        policy=self.policy()
        self.confirm(self.score_preview(policy))
        response=self.client.get('/academic/students/S001/export/excel')
        book=load_workbook(io.BytesIO(response.data))
        self.assertIn('Configured Components',book.sheetnames)
        self.assertEqual(book['Configured Components']['K6'].value,42)
        self.assertEqual(self.client.get('/academic/students/S001/export/pdf').status_code,200)
        self.assertIn(b'Configured assessment components',self.client.get('/academic/students/S001').data)
        db.session.add(TeachingAssignment(teacher_id=2,subject_code='LB601',**self.c));db.session.commit()
        self.fixture.logout();self.fixture.login('teacher','teacher123')
        book=load_workbook(io.BytesIO(self.client.get('/academic/students/S001/export/excel').data))
        self.assertNotIn('Configured Components',book.sheetnames)
        self.assertEqual(self.client.get(f'/records/assessments/{policy.token}').status_code,403)
        payload=self.client.get('/api/analytics/student/S001').get_json()
        self.assertEqual({r['subject_code'] for r in payload['history']},{'LB601'})
        self.assertEqual(payload['cgpa'],'PENDING')

    def test_university_best_attempt_choice_uses_latest_source_revision_and_preserves_printed_history(self):
        self.only_theory()
        from test_academic_linking import AcademicLinkingTests
        job=AcademicLinkingTests.dataset(self,suffix='progression')
        response=self.post(f'/academic/linking/{job.id}/preview',dict(academic_year='2025-26'))
        self.post(response.location+'/confirm',dict(rows='0',verified='on'))
        policy=self.progression_policy()
        rows,summary=progression.analyze(db.session.get(Student,'S001'),policy,'2025-26')
        self.assertTrue(summary['eligible'])
        self.assertEqual(rows[0]['selected']['source'],'UNIVERSITY')
        self.confirm(self.progression_preview(policy))
        self.assertEqual(CourseOutcomeSelection.query.one().evidence['grade'],'AB')
        self.assertIsNone(TheoryMarks.query.first())

    def test_graduation_policy_cannot_allow_backlogs_and_requires_clearance_verification(self):
        self.post('/records/progression/policies',dict(title='Invalid graduation',from_semester=8,reviewed_semesters='8',maximum_backlog_subjects=1,
            maximum_backlog_credits=4,minimum_earned_percent=90,source_reference='Ordinance graduation clause.',verified='on'))
        self.assertEqual(ProgressionPolicy.query.count(),0)
        db.session.add(Semester(semester_id=8,academic_year='2026-27'))
        db.session.add(Subject(subject_code='TH801',subject_name='Final year theory',semester_id=8,subject_type='THEORY',credits=4))
        db.session.flush()
        db.session.add(Enrollment(prn='S001',subject_code='TH801',semester_id=8,academic_year='2026-27'));db.session.commit()
        self.post('/records/progression/policies',dict(title='Graduation clearance',from_semester=8,reviewed_semesters='8',maximum_backlog_subjects=0,
            maximum_backlog_credits=0,minimum_earned_percent=100,source_reference='Ordinance graduation clause.',verified='on'))
        policy=ProgressionPolicy.query.one()
        response=self.post('/records/progression/preview',dict(prn='S001',policy=policy.token,academic_year='2026-27',reason='Verified programme completion requirements.'))
        self.assertNotIn('/reviews/',response.location)

    def test_csv_template_is_context_bound_and_teacher_cannot_preview_other_course(self):
        policy=self.policy()
        content=self.client.get(f'/records/assessments/{policy.token}/template?format=csv').data
        self.assertEqual(len(assessments.parse_scores(policy,content)),3)
        with self.assertRaises(ValueError):assessments.parse_scores(policy,content.replace(b'2025-26',b'2026-27'))
        self.fixture.logout();self.fixture.login('teacher','teacher123')
        response=self.post(f'/records/assessments/{policy.token}/preview',dict(marks_file=(io.BytesIO(content),'marks.csv')))
        self.assertEqual(response.status_code,403)

    def test_progression_required_semester_missing_enrollment_blocks_approval(self):
        self.complete_local()
        policy=self.progression_policy(reviewed_semesters='5,6')
        response=self.progression_preview(policy)
        review=self.review(response)
        self.assertEqual(review.context['decision'],'HOLD')
        self.assertTrue(any('Missing required enrollment' in r for r in review.context['summary']['reasons']))

    def test_repeat_local_attempts_count_course_credit_once_and_suggest_complete_pass(self):
        self.complete_local()
        db.session.add(Semester(semester_id=6,academic_year='2024-25'))
        db.session.flush()
        db.session.add(Enrollment(prn='S001',subject_code='TH601',semester_id=6,academic_year='2024-25'))
        db.session.add(TheoryMarks(prn='S001',subject_code='TH601',semester_id=6,academic_year='2024-25',ct1=8,ct2=9,assignment=7,midsem=15,external=19))
        db.session.commit()
        policy=self.progression_policy()
        rows,summary=progression.analyze(db.session.get(Student,'S001'),policy,'2025-26')
        self.assertEqual(len(rows),1)
        self.assertEqual(len(rows[0]['options']),2)
        self.assertEqual(summary['registered_credits'],4)
        self.assertEqual(rows[0]['selected']['teaching_year'],'2025-26')
        self.assertTrue(summary['eligible'])


if __name__=='__main__':
    unittest.main()
