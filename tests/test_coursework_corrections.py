import io,re,unittest
from openpyxl import load_workbook
import test_workflows as base
from models import db,User,Enrollment,TheoryMarks,ExternalMarks,LabMarks,SgpaCgpa
from teaching_models import TeachingAssignment
from workflow_models import WorkflowReview,AcceptedInternalMark
from coursework_models import AcceptedCourseworkMark,MarkRevision
import coursework_services as service


class CourseworkCorrectionTests(unittest.TestCase):
    def setUp(self):
        self.fixture=base.WorkflowTestCase();self.fixture.setUp();self.client=self.fixture.client
        self.teacher=User.query.filter_by(username='teacher').one()
        for code in ['TH601','LB601']:db.session.add(TeachingAssignment(teacher_id=self.teacher.user_id,subject_code=code,semester_id=6,academic_year='2025-26'))
        db.session.commit()
        self.context={'subject_code':'TH601','semester_id':6,'academic_year':'2025-26','division':'','batch':''}
    def tearDown(self):self.fixture.tearDown()
    def token(self):
        from flask import g
        g.pop('csrf_token',None)
        return re.search(rb'name="csrf_token" value="([^"]+)"',self.client.get('/profile/password').data).group(1).decode()
    def post(self,url,data):return self.client.post(url,data={**data,'csrf_token':self.token()})
    def confirm(self,response):
        self.assertEqual(response.status_code,302,response.data)
        self.assertIn('/workflow/reviews/',response.location)
        return self.post(response.location+'/confirm',{})
    def preview(self,category,values=None,**filters):
        context={**self.context,**filters,'subject_code':'LB601' if category=='LAB' else 'TH601'}
        url='/download-lab-template' if category=='LAB' else '/download-external-template'
        response=self.client.get(url,query_string=context);self.assertEqual(response.status_code,200)
        book=load_workbook(io.BytesIO(response.data));sheet=book.active
        header=next(c.row for row in sheet for c in row if c.value=='PRN')
        if values is None:values=[(20,21,35)]*3 if category=='LAB' else [(45,)]*3
        for number,row in enumerate(values,header+1):
            for col,value in zip((5,6,8) if category=='LAB' else (5,),row):sheet.cell(number,col,value)
        stream=io.BytesIO();book.save(stream);stream.seek(0)
        return self.post('/coursework/'+('lab' if category=='LAB' else 'external')+'/preview',
            {**context,('lab_file' if category=='LAB' else 'file'):(stream,'marks.xlsx')})
    def request(self,category='EXTERNAL',component='EXTERNAL',value='50',prn='S001',**changes):
        data={**self.context,'subject_code':'LB601' if category=='LAB' else 'TH601','category':category,'component':component,
            'prn':prn,'new_value':value,'reason':'Corrected against the signed faculty marks sheet.',**changes}
        return self.post('/coursework/corrections/new',data)
    def decide(self,response,decision='APPROVE',note='Checked the original marks sheet.'):
        self.assertEqual(response.status_code,302)
        self.assertIn('/coursework/corrections/',response.location)
        return self.post(response.location+'/decide',{'decision':decision,'decision_note':note})

    def test_lab_preview_then_confirm_is_atomic_and_replay_safe(self):
        response=self.preview('LAB');self.assertEqual(LabMarks.query.count(),0)
        self.assertEqual(self.client.get(response.location).status_code,200)
        self.assertEqual(self.confirm(response).status_code,302)
        self.assertEqual(LabMarks.query.filter_by(prn='S001').one().total_marks,76)
        self.assertEqual(AcceptedCourseworkMark.query.count(),9)
        self.assertEqual(self.confirm(response).status_code,302);self.assertEqual(AcceptedCourseworkMark.query.count(),9)
        self.assertNotIn('/workflow/reviews/',self.preview('LAB').location)

    def test_external_absence_and_zero_are_preserved_without_empty_theory_rows(self):
        response=self.preview('EXTERNAL',[(0,),('A',),(45,)])
        self.assertEqual(ExternalMarks.query.count(),0)
        self.assertEqual(self.confirm(response).status_code,302)
        self.assertEqual(ExternalMarks.query.filter_by(prn='S001').one().external_marks,0)
        self.assertIsNone(ExternalMarks.query.filter_by(prn='S002').one().external_marks)
        self.assertTrue(all(r.locked for r in ExternalMarks.query))
        self.assertEqual(TheoryMarks.query.count(),0)
        page=self.client.get('/workflow/subject',query_string=self.context)
        self.assertIn(b'External ESE',page.data);self.assertIn(b'Absent',page.data)

    def test_lab_absence_does_not_become_zero_or_a_completed_result(self):
        response=self.preview('LAB',[('A',21,35),(0,0,0),(20,21,35)])
        self.assertEqual(self.confirm(response).status_code,302)
        row=LabMarks.query.filter_by(prn='S001').one()
        self.assertIsNone(row.ca1);self.assertIsNone(row.internal);self.assertIsNone(row.total_marks)
        page=self.client.get('/workflow/subject',query_string={**self.context,'subject_code':'LB601'})
        self.assertIn(b'Absent',page.data)

    def test_invalid_blank_formula_precision_and_ranges_leave_no_partial_marks(self):
        for category,values in [('LAB',[(31,21,35)]*3),('LAB',[(None,21,35)]*3),('EXTERNAL',[('=1+1',)]*3),('EXTERNAL',[(42.25,)]*3)]:
            response=self.preview(category,values)
            self.assertNotIn('/workflow/reviews/',response.location)
        self.assertEqual(LabMarks.query.count(),0);self.assertEqual(ExternalMarks.query.count(),0)

    def test_changed_roster_and_revoked_assignment_block_lab_confirmation(self):
        self.fixture.logout();self.fixture.login('teacher','teacher123')
        response=self.preview('LAB')
        assignment=TeachingAssignment.query.filter_by(subject_code='LB601').one();assignment.is_active=False;db.session.commit()
        self.assertEqual(self.confirm(response).status_code,403)
        assignment.is_active=True
        Enrollment.query.filter_by(prn='S003',subject_code='LB601').delete();db.session.commit()
        self.assertEqual(self.confirm(response).status_code,409);self.assertEqual(LabMarks.query.count(),0)

    def test_teacher_request_waits_for_admin_and_teacher_cannot_approve(self):
        self.confirm(self.preview('EXTERNAL'))
        original=AcceptedCourseworkMark.query.filter_by(prn='S001').one().value
        self.fixture.logout();self.fixture.login('teacher','teacher123')
        response=self.request();self.assertEqual(ExternalMarks.query.filter_by(prn='S001').one().external_marks,45)
        self.assertEqual(self.decide(response).status_code,403)
        self.fixture.logout();self.fixture.login('admin','admin123')
        self.assertEqual(self.decide(response).status_code,302)
        self.assertEqual(ExternalMarks.query.filter_by(prn='S001').one().external_marks,50)
        self.assertEqual(AcceptedCourseworkMark.query.filter_by(prn='S001').one().value,original)
        self.assertEqual(MarkRevision.query.one().status,'APPROVED')
        self.assertEqual(self.decide(response).status_code,302)
        self.assertEqual(MarkRevision.query.count(),1)

    def test_rejection_keeps_current_marks_and_requires_a_reason(self):
        self.confirm(self.preview('EXTERNAL'));response=self.request()
        self.assertEqual(self.decide(response,'REJECT','').status_code,409)
        self.assertEqual(self.decide(response,'REJECT','Source does not support the requested mark.').status_code,302)
        self.assertEqual(ExternalMarks.query.filter_by(prn='S001').one().external_marks,45)
        self.assertEqual(MarkRevision.query.one().status,'REJECTED')

    def test_stale_parallel_correction_is_blocked_after_first_approval(self):
        self.confirm(self.preview('EXTERNAL'))
        one=self.request(value='50');two=self.request(value='55')
        self.assertEqual(self.decide(one).status_code,302)
        self.assertEqual(self.decide(two).status_code,409)
        self.assertEqual(ExternalMarks.query.filter_by(prn='S001').one().external_marks,50)
        self.assertEqual(MarkRevision.query.filter_by(status='PENDING').count(),1)

    def test_approved_absence_and_later_numeric_correction_update_effective_state(self):
        self.confirm(self.preview('EXTERNAL'))
        self.assertEqual(self.decide(self.request(value='A')).status_code,302)
        self.assertIsNone(ExternalMarks.query.filter_by(prn='S001').one().external_marks)
        page=self.client.get('/workflow/subject',query_string=self.context);self.assertIn(b'Absent',page.data)
        self.assertEqual(self.decide(self.request(value='48')).status_code,302)
        self.assertEqual(ExternalMarks.query.filter_by(prn='S001').one().external_marks,48)
        self.assertEqual(service.effective_states(self.context)[('S001','EXTERNAL','EXTERNAL')],'PRESENT')

    def test_internal_correction_retains_original_acceptance_and_updates_analysis(self):
        context={**self.context,'exam_type':'CT1'}
        template=self.client.get('/download-ct1-template',query_string=context)
        data=self.fixture._fill_ct_template(template,[7,8,9])
        response=self.post('/workflow/marks/preview',{**context,'ct1_file':(io.BytesIO(data),'ct1.xlsx')});self.confirm(response)
        self.assertEqual(self.decide(self.request('INTERNAL','CT1','A')).status_code,302)
        self.assertIsNone(TheoryMarks.query.filter_by(prn='S001').one().ct1)
        self.assertEqual(AcceptedInternalMark.query.filter_by(prn='S001').one().value,7)
        self.assertEqual(service.internal_states(self.context)[('S001','ct1')],'ABSENT')

    def test_lab_correction_recalculates_total_without_changing_original_evidence(self):
        self.confirm(self.preview('LAB'))
        response=self.request('LAB','CA1','25')
        self.assertEqual(self.decide(response).status_code,302)
        row=LabMarks.query.filter_by(prn='S001').one();self.assertEqual(row.internal,46);self.assertEqual(row.total_marks,81)
        self.assertEqual(AcceptedCourseworkMark.query.filter_by(prn='S001',component='CA1').one().value,20)

    def test_correction_rejected_for_missing_component_noop_short_reason_and_unknown_student(self):
        self.confirm(self.preview('EXTERNAL'))
        for change in [{'value':'45'},{'reason':'short'},{'prn':'UNKNOWN'},{'category':'UNKNOWN'}]:
            response=self.request(**change);self.assertNotIn('/corrections/'+(MarkRevision.query.first().token if MarkRevision.query.first() else 'missing'),response.location)
        self.assertEqual(MarkRevision.query.count(),0)
        self.request('INTERNAL','CT1','5');self.assertEqual(MarkRevision.query.count(),0)

    def test_revoked_requester_blocks_approval_and_unrelated_teacher_cannot_read(self):
        self.confirm(self.preview('EXTERNAL'))
        self.fixture.logout();self.fixture.login('teacher','teacher123');response=self.request()
        assignment=TeachingAssignment.query.filter_by(subject_code='TH601').one();assignment.is_active=False;db.session.commit()
        self.fixture.logout();self.fixture.login('admin','admin123')
        self.assertEqual(self.decide(response).status_code,409)
        other=User(username='unrelated',role='TEACHER',is_active=True);other.set_password('unrelated-password');db.session.add(other);db.session.commit()
        self.fixture.logout();self.fixture.login('unrelated','unrelated-password')
        self.assertEqual(self.client.get(response.location).status_code,403)

    def test_api_save_stages_and_direct_lock_cannot_bypass_review(self):
        response=self.client.post('/external/save',json={**self.context,'marks':[{'prn':'S001','external_marks':42}]})
        self.assertEqual(response.status_code,200);self.assertEqual(response.json['saved'],0)
        self.assertEqual(ExternalMarks.query.count(),0)
        self.assertEqual(self.client.post('/external/lock',json=self.context).status_code,409)
        self.assertEqual(self.post(response.json['review_url']+'/confirm',{}).status_code,302)
        self.assertEqual(ExternalMarks.query.count(),1)

    def test_correction_queue_and_new_preview_require_login_and_csrf(self):
        self.assertEqual(self.client.post('/coursework/corrections/new',data=self.context).status_code,400)
        self.assertEqual(self.client.get('/coursework/corrections').status_code,200)
        self.assertEqual(self.client.get('/coursework/corrections/new',query_string=self.context).status_code,200)
        self.fixture.logout();self.assertEqual(self.client.get('/coursework/corrections').status_code,302)

    def test_admin_can_inspect_original_workbook_and_reject_inactive_subject_request(self):
        from models import Subject
        self.confirm(self.preview('EXTERNAL'));response=self.request()
        self.assertIn(b'Download original accepted workbook',self.client.get(response.location).data)
        subject=db.session.get(Subject,'TH601');subject.is_active=False;db.session.commit()
        self.assertEqual(self.decide(response).status_code,409)
        self.assertEqual(self.decide(response,'REJECT','The subject context is inactive.').status_code,302)
        self.assertEqual(ExternalMarks.query.filter_by(prn='S001').one().external_marks,45)

    def test_external_correction_refreshes_existing_final_grade_and_cgpa(self):
        self.fixture._complete_results()
        before=TheoryMarks.query.filter_by(prn='S001',subject_code='TH601').one()
        old_total=before.total_marks
        response=self.request(value='55')
        self.assertEqual(self.decide(response).status_code,302)
        row=TheoryMarks.query.filter_by(prn='S001',subject_code='TH601').one()
        self.assertEqual(row.external,55)
        self.assertGreater(row.total_marks,old_total)
        self.assertIsNotNone(SgpaCgpa.query.filter_by(prn='S001',semester_id=6,academic_year='2025-26').first())


if __name__=='__main__':unittest.main()
