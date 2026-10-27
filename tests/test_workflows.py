import io
import os
import re
import tempfile
import unittest
from datetime import date, time
from decimal import Decimal

import fitz
from openpyxl import load_workbook


TEST_DIR = tempfile.mkdtemp(prefix="student_analysis_tests_")
os.environ["DATABASE_URL"] = "sqlite:///" + os.path.join(TEST_DIR, "test.db")
os.environ["SECRET_KEY"] = "test-secret-key"
os.environ["CURRENT_ACADEMIC_YEAR"] = "2025-26"
os.environ["CURRENT_SEMESTER"] = "6"
os.environ["RESULT_DATA_UPLOAD_FOLDER"] = os.path.join(TEST_DIR, "result-uploads")

from app import app  # noqa: E402
from attendance_utils import generate_attendance_template  # noqa: E402
from calculations import update_internal_totals  # noqa: E402
from final_result_utils import recalculate_lab_result, recalculate_theory_result  # noqa: E402
from models import (  # noqa: E402
    Attendance,
    AttendanceEntry,
    Enrollment,
    ExternalMarks,
    LabMarks,
    Semester,
    Student,
    Subject,
    TheoryMarks,
    InternalMarksSubmission,
    TimetableSlot,
    User,
    db,
)
from result_data.models import (  # noqa: E402
    ExamReport,
    ImportJob as ResultImportJob,
    ReportStudent,
    ResultComponent,
    ResultSubject,
    StudentSummary,
    SubjectAssessment,
    SubjectResult,
)


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class WorkflowTestCase(unittest.TestCase):
    def setUp(self):
        app.config.update(TESTING=True)
        self.ctx = app.app_context()
        self.ctx.push()
        db.drop_all()
        db.create_all()
        self._seed()
        self.client = app.test_client()
        self.login("admin", "admin123")

    def tearDown(self):
        db.session.remove()
        self.ctx.pop()

    def _seed(self):
        db.session.add_all([
            Semester(semester_id=5, academic_year="2024-25", is_active=False),
            Semester(semester_id=6, academic_year="2025-26", is_active=True),
            Student(prn="S001", name="Asha Patil", year_of_admission=2023, division="A", batch="A1"),
            Student(prn="S002", name="Bharat More", year_of_admission=2023, division="A", batch="A2"),
            Student(prn="S003", name="Charu Shah", year_of_admission=2023, division="B", batch="B1"),
            Subject(
                subject_code="TH601", subject_name="Theory Subject", semester_id=6,
                subject_type="THEORY", credits=4, category="PCC", is_active=True,
            ),
            Subject(
                subject_code="LB601", subject_name="Lab Subject", semester_id=6,
                subject_type="LAB", credits=2, category="LC", is_active=True,
            ),
            Subject(
                subject_code="EL600", subject_name="Elective Parent", semester_id=6,
                subject_type="THEORY", credits=3, category="PEC", is_elective=True,
                elective_group="PEC-1", is_active=True,
            ),
            Subject(
                subject_code="EL601A", subject_name="Elective A", semester_id=6,
                subject_type="THEORY", credits=3, category="PEC", is_elective=True,
                elective_group="PEC-1", parent_subject_code="EL600", is_active=True,
            ),
            Subject(
                subject_code="EL601B", subject_name="Elective B", semester_id=6,
                subject_type="THEORY", credits=3, category="PEC", is_elective=True,
                elective_group="PEC-1", parent_subject_code="EL600", is_active=True,
            ),
            Subject(
                subject_code="ACT601", subject_name="Aptitude", semester_id=6,
                subject_type="AUDIT", credits=0, category="ATTEND", is_audit=True,
                is_attendance_only=True, is_active=True,
            ),
        ])

        admin = User(username="admin", role="ADMIN", is_active=True)
        admin.set_password("admin123")
        teacher = User(username="teacher", role="TEACHER", is_active=True)
        teacher.set_password("teacher123")
        db.session.add_all([admin, teacher])
        db.session.flush()

        for prn in ("S001", "S002", "S003"):
            for code in ("TH601", "LB601"):
                db.session.add(Enrollment(
                    prn=prn, subject_code=code, semester_id=6,
                    academic_year="2025-26",
                ))
        db.session.add_all([
            Enrollment(prn="S001", subject_code="EL601A", semester_id=6, academic_year="2025-26"),
            Enrollment(prn="S002", subject_code="EL601B", semester_id=6, academic_year="2025-26"),
        ])
        db.session.commit()

    def login(self, username, password):
        return self.client.post(
            "/login", data={"username": username, "password": password},
            follow_redirects=False,
        )

    def logout(self):
        return self.client.get("/logout", follow_redirects=False)

    def _xlsx_upload(self, url, fields, payload, filename="upload.xlsx"):
        data = dict(fields)
        data[next(iter(payload))] = (io.BytesIO(next(iter(payload.values()))), filename)
        response=self.client.post(url, data=data, content_type="multipart/form-data", follow_redirects=False)
        if response.status_code==302 and '/workflow/reviews/' in response.location:
            page=self.client.get(response.location)
            token=re.search(rb'name="csrf_token" value="([^"]+)"',page.data).group(1).decode()
            response=self.client.post(response.location+'/confirm',data={'csrf_token':token})
        return response

    def _fill_ct_template(self, response, marks):
        wb = load_workbook(io.BytesIO(response.data))
        ws = wb.active
        for offset, mark in enumerate(marks, start=11):
            ws.cell(offset, 5, mark)
        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    def _fill_lab_template(self, response, rows):
        wb = load_workbook(io.BytesIO(response.data))
        ws = wb.active
        for offset, (ca1, ca2, external) in enumerate(rows, start=11):
            ws.cell(offset, 5, ca1)
            ws.cell(offset, 6, ca2)
            ws.cell(offset, 8, external)
        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    def _complete_results(self):
        for prn, external in (("S001", 48), ("S002", 16), ("S003", 44)):
            row = TheoryMarks(
                prn=prn, subject_code="TH601", semester_id=6, academic_year="2025-26",
                ct1=8, ct2=9, assignment=7, midsem=16, external=external,
            )
            update_internal_totals(row)
            db.session.add(row)
            db.session.add(ExternalMarks(
                prn=prn, subject_code="TH601", semester_id=6,
                academic_year="2025-26", external_marks=external, locked=True,
            ))
            lab = LabMarks(
                prn=prn, subject_code="LB601", semester_id=6, academic_year="2025-26",
                ca1=25, ca2=24, external=35,
            )
            recalculate_lab_result(lab)
            db.session.add(lab)
        db.session.commit()
        for row in TheoryMarks.query.all():
            recalculate_theory_result(row)
        db.session.commit()

    def _create_result_pdf_dataset(self):
        job = ResultImportJob(
            original_filename="semester-result.pdf",
            stored_filename="integration-test-result.pdf",
            file_hash="a" * 64,
            file_size_bytes=2048,
            page_count=1,
            parser_version="1.0",
            status=ResultImportJob.STATUS_READY,
        )
        report = ExamReport(
            parser_name="DBATU Crystal Report",
            parser_version="1.0",
            university="Dr. Babasaheb Ambedkar Technological University",
            examination_title="B.Tech Semester Examination",
            programme="Bachelor of Technology",
            branch="Electronics and Computer Engineering",
            semester="6",
            examination_session="Summer 2026",
            examination_type="Regular",
            institute_code="2116",
            institute_name="CSMSS Chh. Shahu College of Engineering",
            maximum_marks=Decimal("100"),
            maximum_credits=Decimal("4"),
            maximum_grade_points=Decimal("40"),
            calculation_rule_version="DBATU-10-point-v1",
            validation_status="VALID",
            validation_error_count=0,
            validation_warning_count=0,
        )
        subject = ResultSubject(
            subject_code="BTECPE601",
            subject_name="Integration Subject",
            display_order=1,
            credits=Decimal("4"),
            is_audit=False,
            total_maximum=Decimal("100"),
            total_minimum=Decimal("40"),
        )
        assessment = SubjectAssessment(
            component_type="ESE",
            display_name="End Semester Exam",
            source_row="ESE",
            maximum_marks=Decimal("100"),
            minimum_marks=Decimal("40"),
            display_order=1,
        )
        student = ReportStudent(
            seat_number="ECE001",
            name="Integration Student",
            source_page=1,
            result_status="PASS",
            validation_status="VALID",
        )
        summary = StudentSummary(
            reported_sgpa=Decimal("8.50"),
            reported_total_grade_points=Decimal("34"),
            reported_earned_credits=Decimal("4"),
            reported_total_marks=Decimal("82"),
            reported_status="PASS",
            calculated_sgpa=Decimal("8.50"),
            calculated_total_grade_points=Decimal("34"),
            calculated_earned_credits=Decimal("4"),
            calculated_total_marks=Decimal("82"),
            calculated_status="PASS",
        )
        result = SubjectResult(
            source_page=1,
            total_raw="82",
            total_value=Decimal("82"),
            grade_raw="A",
            grade="A",
            grade_point=Decimal("8.5"),
            credit_points=Decimal("34"),
            component_sum=Decimal("82"),
            expected_grade="A",
            expected_grade_point=Decimal("8.5"),
            expected_credit_points=Decimal("34"),
            validation_status="VALID",
        )
        component = ResultComponent(
            component_type="ESE",
            raw_text="82",
            numeric_value=Decimal("82"),
        )

        job.report = report
        report.subjects.append(subject)
        subject.assessments.append(assessment)
        report.students.append(student)
        student.summary = summary
        student.subject_results.append(result)
        result.subject = subject
        result.components.append(component)
        component.assessment = assessment
        db.session.add(job)
        db.session.commit()
        return job

    def _upload_adaptive_pdf(self, document):
        pdf_bytes = document.tobytes()
        document.close()
        page = self.client.get("/result-data/imports/new")
        token = re.search(rb'name="csrf_token"[^>]*value="([^"]+)"', page.data).group(1).decode()
        response = self.client.post("/result-data/imports/new", data={
            "csrf_token": token,
            "pdf_file": (io.BytesIO(pdf_bytes), "adaptive-report.pdf"),
        }, content_type="multipart/form-data", follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        job = ResultImportJob.query.order_by(ResultImportJob.id.desc()).first()
        self.assertEqual(job.status, "REVIEW_REQUIRED", job.error_message)
        self.assertIsNotNone(job.report)
        return job

    def test_successful_pdf_upload_uses_result_upload_folder(self):
        from test_adaptive_result_parser import fixture_document
        job = self._upload_adaptive_pdf(fixture_document([5]))
        self.assertEqual(job.parser_version, "2.0.1")
        self.assertEqual(job.report.validation_error_count, 0)
        self.assertEqual(len(job.report.students), 3)
        self.assertTrue(os.path.isfile(os.path.join(app.config["RESULT_DATA_UPLOAD_FOLDER"], job.stored_filename)))
        self.assertFalse(os.path.isfile(os.path.join(app.config["UPLOAD_FOLDER"], job.stored_filename)))

    def test_pdf_electives_validation_coverage_and_reextraction(self):
        from test_adaptive_result_parser import fixture_document
        from result_data.services.acceptance_service import audit_import
        from result_data.services.analysis_service import build_report_analysis
        from result_data.services.extraction_service import extract_import_job
        from result_data.models import ResultSourceData
        job = self._upload_adaptive_pdf(fixture_document([3,4]))
        self.assertEqual(len(job.report.subjects), 10)
        self.assertEqual(job.report.validation_error_count, 0)
        self.assertFalse(any(i.code == "MISSING_SUBJECT_RESULT" for i in job.report.validation_issues))
        self.assertTrue(all(check.passed for check in audit_import(job)))
        self.assertEqual(build_report_analysis(job.report)["summary"]["completion_rate"], 100)
        self.assertEqual(job.report.expected_result_count, job.report.result_count)
        for student in job.report.students:
            response = self.client.get(f"/result-data/imports/{job.id}/students/{student.id}")
            self.assertEqual(response.status_code, 200)
        before = job.report.result_count
        extract_import_job(job, app.config["RESULT_DATA_UPLOAD_FOLDER"])
        self.assertEqual(job.report.result_count, before)
        self.assertEqual(ResultSourceData.query.count(), 1)
        for extension in ("xlsx", "csv", "json"):
            response=self.client.get(f"/result-data/exports/imports/{job.id}/results.{extension}")
            self.assertEqual(response.status_code, 200)

    def test_optional_pdf_summary_fields_survive_database_and_exports(self):
        from test_adaptive_result_parser import fixture_document
        from result_data.services.export_service import build_json_export, build_csv_export, build_detailed_excel_export
        import json
        job = self._upload_adaptive_pdf(fixture_document([1]))
        student = job.report.students[0]
        detail = self.client.get(f"/result-data/imports/{job.id}/students/{student.id}")
        self.assertIn(b"Ext Total(380)", detail.data)
        self.assertIn(b"Int Total(370)", detail.data)
        payload=json.loads(build_json_export(job).data)
        fields=payload["students"][0]["summary"]["additional_source_fields"]
        self.assertEqual(fields["external_total"]["raw_text"], "153")
        self.assertEqual(fields["internal_total"]["raw_text"], "192")
        self.assertIn(b"Source External Total", build_csv_export(job).data)
        wb=load_workbook(io.BytesIO(build_detailed_excel_export(job).data))
        self.assertTrue(any(cell.value == "Source Internal Total" for row in wb["Students"] for cell in row))
        db.session.expire_all()
        self.assertEqual(job.report.additional_summary_for(student.seat_number)["external_total"]["raw_text"], "153")

    def test_reextract_failure_retains_previous_report(self):
        from test_adaptive_result_parser import fixture_document
        from result_data.services.extraction_service import extract_import_job, ExtractionFailedError
        from result_data.parsers import PDFExtractionError
        job = self._upload_adaptive_pdf(fixture_document([5]))
        report_id, count = job.report.id, job.report.result_count
        class BadParser:
            def parse(self, path):
                raise PDFExtractionError("Test invalid source")
        with self.assertRaises(ExtractionFailedError):
            extract_import_job(job, app.config["RESULT_DATA_UPLOAD_FOLDER"], BadParser())
        self.assertEqual(job.report.id, report_id)
        self.assertEqual(job.report.result_count, count)

    def test_authentication_and_role_protection(self):
        self.logout()
        for path in (
            "/", "/ct1", "/external", "/lab", "/attendance", "/analytics",
            "/subjects", "/result-data/", "/result-data/imports/new",
        ):
            response = self.client.get(path)
            self.assertEqual(response.status_code, 302, path)

        self.login("teacher", "teacher123")
        for path in (
            "/", "/ct1", "/external", "/lab", "/attendance", "/analytics",
            "/result-data/", "/result-data/imports/new",
        ):
            self.assertEqual(self.client.get(path).status_code, 200, path)
        for path in ("/subjects", "/electives", "/batches"):
            response = self.client.get(path)
            self.assertEqual(response.status_code, 302, path)
            self.assertEqual(response.headers["Location"], "/")

    def test_all_main_pages_render_for_admin(self):
        for path in ("/", "/ct1", "/external", "/lab", "/attendance", "/analytics", "/subjects", "/electives", "/batches"):
            response = self.client.get(path)
            self.assertEqual(response.status_code, 200, path)

    def test_integrated_result_data_pages_database_and_exports(self):
        job = self._create_result_pdf_dataset()
        student_id = job.report.students[0].id

        self.assertIn("subjects", db.metadata.tables)
        self.assertIn("result_subjects", db.metadata.tables)
        self.assertNotEqual(Subject.__tablename__, ResultSubject.__tablename__)

        root = self.client.get("/")
        self.assertEqual(root.status_code, 200)
        self.assertIn(b"University Result PDF", root.data)

        for path in (
            "/result-data/",
            "/result-data/imports/",
            "/result-data/imports/new",
            f"/result-data/imports/{job.id}",
            f"/result-data/imports/{job.id}/students",
            f"/result-data/imports/{job.id}/students/{student_id}",
            f"/result-data/imports/{job.id}/students/{student_id}/edit",
            f"/result-data/analysis/imports/{job.id}",
            "/result-data/health",
        ):
            response = self.client.get(path)
            self.assertEqual(response.status_code, 200, path)

        for extension, mimetype in (
            ("xlsx", XLSX_MIME),
            ("csv", "text/csv"),
            ("json", "application/json"),
        ):
            response = self.client.get(
                f"/result-data/exports/imports/{job.id}/results.{extension}"
            )
            self.assertEqual(response.status_code, 200, extension)
            self.assertEqual(response.mimetype, mimetype)
            self.assertTrue(response.data)

        # Imported write routes require CSRF, while legacy corrected-project
        # forms keep their established behavior.
        response = self.client.post(f"/result-data/imports/{job.id}/validate")
        self.assertEqual(response.status_code, 400)

    def test_result_pdf_upload_is_stored_and_unsupported_layout_fails_safely(self):
        document = fitz.open()
        page = document.new_page()
        page.insert_text((72, 72), "Unsupported sample result layout")
        pdf_bytes = document.tobytes()
        document.close()

        upload_page = self.client.get("/result-data/imports/new")
        token_match = re.search(
            rb'name="csrf_token"[^>]*value="([^"]+)"', upload_page.data
        )
        self.assertIsNotNone(token_match)

        response = self.client.post(
            "/result-data/imports/new",
            data={
                "csrf_token": token_match.group(1).decode(),
                "pdf_file": (io.BytesIO(pdf_bytes), "unsupported.pdf"),
                "submit": "Upload PDF",
            },
            content_type="multipart/form-data",
            follow_redirects=False,
        )
        self.assertEqual(response.status_code, 302)

        job = ResultImportJob.query.one()
        self.assertEqual(job.status, ResultImportJob.STATUS_FAILED)
        self.assertTrue(job.error_message)
        source_path = os.path.join(
            app.config["RESULT_DATA_UPLOAD_FOLDER"], job.stored_filename
        )
        self.assertTrue(os.path.isfile(source_path))

    def test_active_semester_switches_an_exact_row(self):
        response = self.client.post(
            "/set-active-semester",
            data={"semester_id": "5", "academic_year": "2024-25"},
        )
        self.assertEqual(response.status_code, 302)
        active = Semester.query.filter_by(is_active=True).all()
        self.assertEqual([(s.semester_id, s.academic_year) for s in active], [(5, "2024-25")])

        response = self.client.post(
            "/set-active-semester",
            data={"semester_key": "6|2025-26"},
        )
        self.assertEqual(response.status_code, 302)
        active = Semester.query.filter_by(is_active=True).one()
        self.assertEqual((active.semester_id, active.academic_year), (6, "2025-26"))

        db.session.add(Semester(semester_id=6, academic_year="2024-25", is_active=False))
        db.session.commit()
        self.client.post("/set-active-semester", data={"semester_key": "6|2024-25"})
        active = Semester.query.filter_by(is_active=True).one()
        self.assertEqual((active.semester_id, active.academic_year), (6, "2024-25"))

    def test_batch_assignment_and_clear(self):
        response = self.client.post(
            "/batches/save",
            data={"division": "B", "batch": "B2", "student_prns": ["S001", "S002"]},
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(db.session.get(Student, "S001").batch, "B2")
        self.assertEqual(db.session.get(Student, "S002").division, "B")

        response = self.client.post("/batches/clear", data={"student_prns": ["S001"]})
        self.assertEqual(response.status_code, 302)
        self.assertIsNone(db.session.get(Student, "S001").division)
        self.assertIsNone(db.session.get(Student, "S001").batch)

    def test_subject_crud_and_safe_delete(self):
        response = self.client.post("/subjects/save", data={
            "subject_code": "NEW601", "subject_name": "New Subject", "semester_id": "6",
            "subject_type": "THEORY", "credits": "3", "is_active": "on",
        })
        self.assertEqual(response.status_code, 302)
        self.assertIsNotNone(db.session.get(Subject, "NEW601"))

        self.client.post("/subjects/NEW601/enroll-compulsory")
        self.assertEqual(Enrollment.query.filter_by(subject_code="NEW601").count(), 3)
        self.client.post("/subjects/NEW601/delete")
        self.assertFalse(db.session.get(Subject, "NEW601").is_active)

        self.client.post("/subjects/save", data={
            "subject_code": "EMPTY601", "subject_name": "Empty Subject", "semester_id": "6",
            "subject_type": "THEORY", "credits": "1", "is_active": "on",
        })
        self.client.post("/subjects/EMPTY601/delete")
        self.assertIsNone(db.session.get(Subject, "EMPTY601"))

        self.client.post("/subjects/TH601/disable")
        self.assertFalse(db.session.get(Subject, "TH601").is_active)
        self.client.post("/subjects/TH601/enable")
        self.assertTrue(db.session.get(Subject, "TH601").is_active)

        self.client.post("/subjects/save", data={
            "subject_code": "BAD601", "subject_name": "Invalid Subject", "semester_id": "6",
            "subject_type": "THEORY", "credits": "-1", "is_active": "on",
        })
        self.assertIsNone(db.session.get(Subject, "BAD601"))

    def test_elective_assignment_replaces_only_same_group(self):
        response = self.client.post("/electives/save", data={
            "semester_id": "6", "elective_group": "PEC-1",
            "subject_code": "EL601B", "student_prns": ["S001"],
        })
        self.assertEqual(response.status_code, 302)
        choices = Enrollment.query.filter(
            Enrollment.prn == "S001",
            Enrollment.subject_code.in_(["EL600", "EL601A", "EL601B"]),
        ).all()
        self.assertEqual([row.subject_code for row in choices], ["EL601B"])
        self.assertIsNotNone(Enrollment.query.filter_by(prn="S001", subject_code="TH601").first())

    def test_internal_template_upload_and_all_reports(self):
        query = "semester_id=6&exam_type=CT1&subject_code=TH601&division=A&batch=A1&subject_teacher=Teacher"
        template = self.client.get("/download-ct1-template?" + query)
        self.assertEqual(template.status_code, 200)
        self.assertEqual(template.mimetype, XLSX_MIME)
        filled = self._fill_ct_template(template, [8])
        upload = self._xlsx_upload(
            "/upload-ct1",
            {"subject_code": "TH601", "exam_type": "CT1", "semester_id": "6", "division": "A", "batch": "A1"},
            {"ct1_file": filled},
        )
        self.assertEqual(upload.status_code, 302)
        self.assertEqual(float(TheoryMarks.query.filter_by(prn="S001", subject_code="TH601").one().ct1), 8)
        self.assertEqual(InternalMarksSubmission.query.filter_by(batch="A1", exam_type="CT1").count(), 1)

        for fmt, mimetype in (("excel", XLSX_MIME), ("pdf", "application/pdf")):
            response = self.client.get(f"/download-ct1-report/{fmt}?{query}")
            self.assertEqual(response.status_code, 200, fmt)
            self.assertEqual(response.mimetype, mimetype)
            response = self.client.get(f"/download-internal-report/{fmt}?subject_code=TH601&semester_id=6&division=A&batch=A1")
            self.assertEqual(response.status_code, 200, fmt)
            self.assertEqual(response.mimetype, mimetype)

        for path, mimetype in (
            ("/download-ct1-summary-excel?semester_id=6&exam_type=CT1", XLSX_MIME),
            ("/download-ct1-summary-pdf?semester_id=6&exam_type=CT1", "application/pdf"),
        ):
            response = self.client.get(path)
            self.assertEqual(response.status_code, 200, path)
            self.assertEqual(response.mimetype, mimetype)

    def test_internal_submission_lock_is_scoped_to_selected_batch(self):
        a1_query = "semester_id=6&exam_type=CT1&subject_code=TH601&division=A&batch=A1&subject_teacher=Teacher"
        template = self.client.get("/download-ct1-template?" + a1_query)
        upload = self._xlsx_upload(
            "/upload-ct1",
            {"subject_code": "TH601", "exam_type": "CT1", "semester_id": "6", "division": "A", "batch": "A1"},
            {"ct1_file": self._fill_ct_template(template, [8])},
        )
        self.assertEqual(upload.status_code, 302)

        a1_page = self.client.get("/ct1?" + a1_query)
        self.assertIn(b"Marks have been successfully submitted and locked", a1_page.data)

        a2_query = "semester_id=6&exam_type=CT1&subject_code=TH601&division=A&batch=A2&subject_teacher=Teacher"
        a2_page = self.client.get("/ct1?" + a2_query)
        self.assertIn(b"Marks for this subject and examination have not been submitted", a2_page.data)
        template = self.client.get("/download-ct1-template?" + a2_query)
        upload = self._xlsx_upload(
            "/upload-ct1",
            {"subject_code": "TH601", "exam_type": "CT1", "semester_id": "6", "division": "A", "batch": "A2"},
            {"ct1_file": self._fill_ct_template(template, [7])},
        )
        self.assertEqual(upload.status_code, 302)
        self.assertEqual(float(TheoryMarks.query.filter_by(prn="S002", subject_code="TH601").one().ct1), 7)

    def test_internal_upload_rejects_unenrolled_student(self):
        Student.query.filter_by(prn="S003").one().division = "B"
        Enrollment.query.filter_by(prn="S003", subject_code="TH601").delete()
        db.session.commit()
        template = self.client.get(
            "/download-ct1-template?semester_id=6&exam_type=CT1&subject_code=TH601&division=A&subject_teacher=Teacher"
        )
        wb = load_workbook(io.BytesIO(template.data))
        ws = wb.active
        ws.cell(11, 3, "S003")
        ws.cell(11, 5, 8)
        output = io.BytesIO()
        wb.save(output)
        upload = self._xlsx_upload(
            "/upload-ct1",
            {"subject_code": "TH601", "exam_type": "CT1", "semester_id": "6"},
            {"ct1_file": output.getvalue()},
        )
        self.assertEqual(upload.status_code, 302)
        self.assertIsNone(TheoryMarks.query.filter_by(prn="S003", subject_code="TH601").first())

    def test_internal_upload_rejects_wrong_or_incomplete_template(self):
        query = "semester_id=6&exam_type=CT1&subject_code=TH601&division=A&subject_teacher=Teacher"
        template = self.client.get("/download-ct1-template?" + query)
        filled = self._fill_ct_template(template, [8, 7])
        upload = self._xlsx_upload(
            "/upload-ct1",
            {"subject_code": "TH601", "exam_type": "CT2", "semester_id": "6", "division": "A"},
            {"ct1_file": filled},
        )
        self.assertEqual(upload.status_code, 302)
        self.assertEqual(TheoryMarks.query.count(), 0)

        wb = load_workbook(io.BytesIO(filled))
        ws = wb.active
        ws.delete_rows(12, 1)
        output = io.BytesIO()
        wb.save(output)
        upload = self._xlsx_upload(
            "/upload-ct1",
            {"subject_code": "TH601", "exam_type": "CT1", "semester_id": "6", "division": "A"},
            {"ct1_file": output.getvalue()},
        )
        self.assertEqual(upload.status_code, 302)
        self.assertEqual(TheoryMarks.query.count(), 0)
        self.assertEqual(InternalMarksSubmission.query.count(), 0)

    def test_external_template_upload_lock_stats_and_reports(self):
        query = "semester_id=6&subject_code=TH601&academic_year=2025-26&division=A&batch=A1"
        template = self.client.get("/download-external-template?" + query)
        self.assertEqual(template.status_code, 200)
        wb = load_workbook(io.BytesIO(template.data))
        ws = wb.active
        header_row = next(cell.row for row in ws.iter_rows() for cell in row if cell.value and "PRN" == str(cell.value).strip())
        ws.cell(header_row + 1, 5, 48)
        output = io.BytesIO()
        wb.save(output)
        upload = self._xlsx_upload(
            "/upload-external",
            {"subject_code": "TH601", "semester_id": "6", "academic_year": "2025-26", "division": "A", "batch": "A1"},
            {"file": output.getvalue()},
        )
        self.assertEqual(upload.status_code, 302)
        row = ExternalMarks.query.filter_by(prn="S001", subject_code="TH601").one()
        self.assertEqual(float(row.external_marks), 48)
        self.assertTrue(row.locked)

        stats = self.client.get("/external/stats?semester_id=6&academic_year=2025-26")
        self.assertEqual(stats.status_code, 200)
        self.assertGreaterEqual(stats.get_json()["external_submitted"], 1)
        for path, mimetype in (
            ("/download-external-report/excel?" + query, XLSX_MIME),
            ("/download-external-report/pdf?" + query, "application/pdf"),
            ("/download-sgpa-cgpa/excel?" + query, XLSX_MIME),
        ):
            response = self.client.get(path)
            self.assertEqual(response.status_code, 200, path)
            self.assertEqual(response.mimetype, mimetype)

    def test_lab_template_upload_and_reports(self):
        query = "semester_id=6&subject_code=LB601&academic_year=2025-26&division=A&batch=A1"
        template = self.client.get("/download-lab-template?" + query)
        self.assertEqual(template.status_code, 200)
        filled = self._fill_lab_template(template, [(25, 24, 35)])
        upload = self._xlsx_upload(
            "/upload-lab",
            {"subject_code": "LB601", "semester_id": "6", "academic_year": "2025-26", "division": "A", "batch": "A1"},
            {"lab_file": filled},
        )
        self.assertEqual(upload.status_code, 302)
        row = LabMarks.query.filter_by(prn="S001", subject_code="LB601").one()
        self.assertEqual(float(row.total_marks), 84)
        for fmt, mimetype in (("excel", XLSX_MIME), ("pdf", "application/pdf")):
            response = self.client.get(f"/download-lab-report/{fmt}?{query}")
            self.assertEqual(response.status_code, 200, fmt)
            self.assertEqual(response.mimetype, mimetype)

    def test_attendance_timetable_template_upload_sheet_and_reports(self):
        template = self.client.get("/attendance/download-timetable-template")
        self.assertEqual(template.status_code, 200)
        wb = load_workbook(io.BytesIO(template.data))
        ws = wb["Class Timetable"]
        ws["B4"] = "2025-26"
        ws["D4"] = 6
        ws["H4"] = "A"
        for row in range(7, 13):
            for col in range(2, 10):
                ws.cell(row, col, "")
        ws["B7"] = "TH601 - Theory Subject [THEORY]"
        output = io.BytesIO()
        wb.save(output)
        upload = self._xlsx_upload(
            "/attendance/upload-timetable", {}, {"timetable_file": output.getvalue()}
        )
        self.assertEqual(upload.status_code, 302)
        self.assertEqual(TimetableSlot.query.filter_by(subject_code="TH601", is_active=True).count(), 1)

        attendance_template = self.client.get(
            "/attendance/download-template?semester_id=6&subject_code=TH601&division=A&batch=A1&start_date=2025-09-01&end_date=2025-09-01"
        )
        self.assertEqual(attendance_template.status_code, 200)
        wb = load_workbook(io.BytesIO(attendance_template.data))
        ws = wb["Attendance"]
        ws["D5"] = "P"
        output = io.BytesIO()
        wb.save(output)
        upload = self._xlsx_upload(
            "/attendance/upload", {}, {"attendance_file": output.getvalue()}
        )
        self.assertEqual(upload.status_code, 302)
        summary = Attendance.query.filter_by(prn="S001", subject_code="TH601").one()
        self.assertEqual(summary.total_lectures, 1)
        self.assertEqual(summary.attended_lectures, 1)

        for fmt, mimetype in (("excel", XLSX_MIME), ("pdf", "application/pdf")):
            response = self.client.get(
                f"/attendance/defaulters/{fmt}?semester_id=6&subject_code=TH601&division=A&batch=A1"
            )
            self.assertEqual(response.status_code, 200, fmt)
            self.assertEqual(response.mimetype, mimetype)

    def test_attendance_does_not_fallback_to_unenrolled_students_for_normal_subject(self):
        slot = TimetableSlot(
            semester_id=6, academic_year="2025-26", division="B", batch="B1",
            day_of_week="Monday", start_time=time(10, 0), end_time=time(11, 0),
            subject_code="TH601", session_type="THEORY", is_active=True,
        )
        db.session.add(slot)
        Enrollment.query.filter_by(prn="S003", subject_code="TH601").delete()
        db.session.commit()
        response = self.client.get(
            "/attendance/download-template?semester_id=6&subject_code=TH601&division=B&batch=B1&start_date=2025-09-01&end_date=2025-09-01"
        )
        self.assertEqual(response.status_code, 302)

    def test_timetable_reupload_replaces_old_slots_and_bad_files_do_not_crash(self):
        def timetable_bytes(subject_code):
            response = self.client.get("/attendance/download-timetable-template")
            wb = load_workbook(io.BytesIO(response.data))
            ws = wb["Class Timetable"]
            ws["B4"], ws["D4"], ws["H4"] = "2025-26", 6, "A"
            for row in range(7, 13):
                for col in range(2, 10):
                    ws.cell(row, col, "")
            ws["B7"] = f"{subject_code} - Subject"
            output = io.BytesIO()
            wb.save(output)
            return output.getvalue()

        for code in ("TH601", "LB601"):
            response = self._xlsx_upload(
                "/attendance/upload-timetable", {}, {"timetable_file": timetable_bytes(code)}
            )
            self.assertEqual(response.status_code, 302)

        self.assertEqual(TimetableSlot.query.filter_by(subject_code="TH601", is_active=True).count(), 0)
        self.assertEqual(TimetableSlot.query.filter_by(subject_code="LB601", is_active=True).count(), 1)

        for endpoint, field in (
            ("/attendance/upload-timetable", "timetable_file"),
            ("/attendance/upload", "attendance_file"),
        ):
            response = self._xlsx_upload(endpoint, {}, {field: b"not an xlsx workbook"})
            self.assertEqual(response.status_code, 302)

    def test_attendance_upload_rejects_student_outside_template_scope(self):
        db.session.add(TimetableSlot(
            semester_id=6, academic_year="2025-26", division="A", batch="A1",
            day_of_week="Monday", start_time=time(10, 0), end_time=time(11, 0),
            subject_code="TH601", session_type="THEORY", is_active=True,
        ))
        db.session.commit()
        response = self.client.get(
            "/attendance/download-template?semester_id=6&subject_code=TH601&division=A&batch=A1&start_date=2025-09-01&end_date=2025-09-01"
        )
        wb = load_workbook(io.BytesIO(response.data))
        ws = wb["Attendance"]
        ws["B5"] = "S003"
        ws["D5"] = "P"
        output = io.BytesIO()
        wb.save(output)
        upload = self._xlsx_upload("/attendance/upload", {}, {"attendance_file": output.getvalue()})
        self.assertEqual(upload.status_code, 302)
        self.assertEqual(AttendanceEntry.query.count(), 0)

    def test_json_and_auxiliary_workflows(self):
        for path in (
            "/api/theory-subjects?semester_id=6",
            "/api/lab-subjects?semester_id=6",
            "/attendance/api/subjects?semester_id=6",
            "/external/stats?semester_id=6&academic_year=2025-26",
        ):
            response = self.client.get(path)
            self.assertEqual(response.status_code, 200, path)
            self.assertTrue(response.is_json, path)

        response = self.client.post("/external/save", json={
            "subject_code": "TH601", "semester_id": 6, "academic_year": "2025-26",
            "marks": [{"student_id": "S001", "external_marks": 42}],
        })
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["saved"], 0)
        self.assertEqual(ExternalMarks.query.count(),0)
        review_url=response.get_json()['review_url']
        page=self.client.get(review_url)
        token=re.search(rb'name="csrf_token" value="([^"]+)"',page.data).group(1).decode()
        self.assertEqual(self.client.post(review_url+'/confirm',data={'csrf_token':token}).status_code,302)
        self.assertTrue(ExternalMarks.query.filter_by(prn="S001").one().locked)
        response=self.client.post('/external/lock',json={'subject_code':'TH601','semester_id':6,'academic_year':'2025-26'})
        self.assertEqual(response.status_code,409)

        from coursework_models import AcceptedCourseworkMark
        AcceptedCourseworkMark.query.delete()
        ExternalMarks.query.delete()

        Enrollment.query.filter_by(prn="S003", subject_code="TH601").delete()
        db.session.commit()
        response = self.client.post("/external/save", json={
            "subject_code": "TH601", "semester_id": 6, "academic_year": "2025-26",
            "marks": [{"student_id": "S003", "external_marks": 40}],
        })
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.get_json()["saved"], 0)
        self.assertTrue(response.get_json()["errors"])
        self.assertEqual(ExternalMarks.query.count(), 0)

        response = self.client.post("/external/lock", json={
            "subject_code": "TH601", "semester_id": 6, "academic_year": "2025-26",
        })
        self.assertEqual(response.status_code, 409)

    def test_analytics_endpoints(self):
        self._complete_results()
        paths = (
            "/api/analytics/subjects/6?academic_year=2025-26",
            "/api/analytics/visual?semester=6&academic_year=2025-26",
            "/api/analytics/semester/6?academic_year=2025-26",
            "/api/analytics/subject/TH601",
            "/api/analytics/student/S001",
        )
        for path in paths:
            response = self.client.get(path)
            self.assertEqual(response.status_code, 200, path)
            self.assertTrue(response.is_json, path)


if __name__ == "__main__":
    unittest.main()
