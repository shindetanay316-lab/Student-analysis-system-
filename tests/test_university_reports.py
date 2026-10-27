import io
import json
import re
import unittest
from datetime import datetime, timezone
from decimal import Decimal

import fitz
from openpyxl import load_workbook
import test_workflows as base
from result_data.models import ImportJob, ExamReport, ResultSubject, SubjectAssessment, ReportStudent, StudentSummary, SubjectResult, ResultComponent
from result_data.services.analysis_service import build_report_analysis, SGPA_BANDS
from result_data.services.analysis_filter_service import filter_records
from result_data.services.export_service import build_excel_export, build_pdf_export
from result_data.services.faculty_report_service import build_faculty_report


class UniversityReportTests(unittest.TestCase):
    def setUp(self):
        self.fixture=base.WorkflowTestCase();self.fixture.setUp()
        self.app=base.app;self.db=base.db;self.client=self.fixture.client
        self.job=self.dataset()

    def tearDown(self):self.fixture.tearDown()

    def dataset(self):
        job=ImportJob(original_filename='Anonymized university report.pdf',stored_filename='report-fixture.pdf',file_hash='9'*64,
                      file_size_bytes=100,status='READY',confirmed_at=datetime.now(timezone.utc),parser_version='2.0.0')
        report=ExamReport(parser_name='DBATU Crystal Reports',parser_version='2.0.0',university='DBATU',
             examination_title='Example Engineering Semester 6 Summer 2026',programme='Bachelor of Technology',branch='Example Engineering',
             semester='6',examination_session='Summer 2026',examination_type='REGULAR',maximum_credits=Decimal('8'),
             maximum_marks=Decimal('200'),maximum_grade_points=Decimal('80'),validation_status='VALID',validation_error_count=0,validation_warning_count=0)
        job.report=report
        subjects=[]
        for i in (1,2):
            subject=ResultSubject(subject_code='SUB'+str(i),subject_name='Example subject '+str(i),display_order=i,credits=Decimal('4'),
                 is_audit=False,total_maximum=Decimal('100'),total_minimum=Decimal('40'))
            subject.assessments.append(SubjectAssessment(component_type='ESE',display_name='End semester',source_row='ESE_0',maximum_marks=100,minimum_marks=40,display_order=1))
            report.subjects.append(subject);subjects.append(subject)
        for i in range(1,16):
            status='FAIL' if i==13 else 'CANCEL' if i==14 else 'PASS'
            sgpa=Decimal('9.50') if i in (1,2) else Decimal('9.99') if i==13 else Decimal('10') if i==14 else None if i==15 else Decimal('9.50')-Decimal(i)/10
            student=ReportStudent(seat_number=str(10000000000000+i),name=f'Student {i:02}',source_page=1,result_status=status,validation_status='WARNING' if i==5 else 'VALID')
            student.summary=StudentSummary(reported_sgpa=sgpa,reported_total_marks=Decimal('164'),reported_status=status,calculated_sgpa=Decimal('1.0'))
            report.students.append(student)
            for j,subject in enumerate(subjects):
                if i==15 and j==1:continue
                result=SubjectResult(subject=subject,source_page=1,total_value=Decimal('82'),total_raw='082',grade='FF' if i==13 and j==0 else 'BB',
                    grade_raw='8/BB/32',grade_point=8,credit_points=32,expected_grade='FF',validation_status='VALID')
                component=ResultComponent(assessment=subject.assessments[0],component_type='ESE',raw_text='082',numeric_value=82)
                result.components.append(component);student.subject_results.append(result)
        self.db.session.add(job);self.db.session.commit()
        return job

    def test_top_ten_uses_reported_sgpa_pass_only_and_ties(self):
        a=build_report_analysis(self.job.report)
        self.assertEqual(len(a['top_performers']),10)
        self.assertEqual(a['summary']['highest_sgpa'],9.5)
        self.assertEqual(len(a['summary']['highest_students']),2)
        self.assertIn('Student 01',a['summary']['highest_student_names'])
        self.assertEqual([r['rank'] for r in a['top_performers'][:3]],[1,1,3])
        self.assertNotIn('attention_students',a)
        self.assertTrue(all(r['result_status']=='PASS' for r in a['top_performers']))

    def test_each_distribution_matches_click_filter(self):
        report=self.job.report;a=build_report_analysis(report)
        for key,kind in [('result_status','result'),('validation_status','validation')]:
            for group in a[key]:
                self.assertEqual(filter_records(report,{'kind':kind,'value':group['key']})['count'],group['count'])
        for row in a['sgpa_distribution']:
            self.assertEqual(filter_records(report,{'kind':'sgpa','value':row['key']})['count'],row['count'])
        for row in a['grade_distribution']:
            self.assertEqual(filter_records(report,{'kind':'grade','value':row['grade']})['count'],row['count'])
        for row in a['component_performance']:
            self.assertEqual(filter_records(report,{'kind':'component','value':row['key']})['count'],row['records'])
        for row in a['subject_performance']:
            self.assertEqual(filter_records(report,{'kind':'subject','value':row['code'],'outcome':'PASS'})['count'],row['passed'])
            self.assertEqual(filter_records(report,{'kind':'subject','value':row['code'],'outcome':'FAIL'})['count'],row['failed'])
            self.assertEqual(filter_records(report,{'kind':'subject','value':row['code'],'measure':'average'})['count'],row['students'])

    def test_invalid_filter_rejected_and_search_preserves_scope(self):
        url=f'/result-data/analysis/imports/{self.job.id}/records'
        for query in ('kind=bogus','kind=subject&value=OTHER','kind=sgpa&value=99','kind=grade&value=UNKNOWN'):
            self.assertEqual(self.client.get(url+'?'+query).status_code,400)
        filtered=filter_records(self.job.report,{'kind':'result','value':'PASS','q':'Student 13'})
        self.assertEqual(filtered['count'],0)
        self.assertEqual(self.client.get(url+'?kind=result&value=FAIL').status_code,200)
        self.assertEqual(self.client.get(url+'?kind=grade&value=BB&page=2').status_code,200)

    def test_sgpa_band_boundaries_and_cancelled_exclusion(self):
        student=self.job.report.students[0]
        for sgpa,index in [('0',0),('4.99',0),('5',1),('8.99',4),('9',5),('10',5)]:
            student.summary.reported_sgpa=Decimal(sgpa)
            matches=filter_records(self.job.report,{'kind':'sgpa','value':str(index)})['records']
            self.assertIn(student.id,[r['id'] for r in matches])
        student.result_status='CANCEL'
        self.assertNotIn(student.id,[r['id'] for r in filter_records(self.job.report,{'kind':'sgpa','value':'5'})['records']])

    def test_faculty_workbook_order_grades_only_and_no_internal_fields(self):
        workbook=load_workbook(io.BytesIO(build_excel_export(self.job).data))
        self.assertEqual(workbook.sheetnames,['Result Summary','Student Results','Subject Analysis','Top 10','Failed Students','Charts'])
        sheet=workbook['Student Results']
        self.assertEqual([sheet.cell(10,i).value for i in (1,2,3)],['Sr. No.','PRN','Student Name'])
        self.assertEqual(sheet['D11'].value,'BB')
        self.assertEqual(sheet['B11'].data_type,'s')
        self.assertEqual(sheet.freeze_panes,'D11')
        self.assertIn('SUB1',sheet['D10'].value)
        self.assertTrue(sheet.auto_filter.ref)
        self.assertTrue(sheet.conditional_formatting)
        self.assertTrue(any('"FAIL"' in rule.formula[0] for rules in sheet.conditional_formatting._cf_rules.values() for rule in rules))
        self.assertEqual(sheet.cell(25,5).value,'MISSING')  # Student 15 has no SUB2 elective record.
        self.assertEqual(len(workbook['Charts']._charts),4)
        self.assertEqual(workbook['Result Summary']['H11'].value,'=IF(C11+D11=0,"-",C11/(C11+D11))')
        self.assertNotIn('Subject ID',[c.value for c in sheet[10]])

    def test_settings_persist_and_are_shared_by_pdf_and_excel(self):
        url=f'/result-data/exports/imports/{self.job.id}/report-settings'
        page=self.client.get(url);self.assertEqual(page.status_code,200)
        token=re.search(rb'name="csrf_token"[^>]*value="([^"]+)"',page.data).group(1).decode()
        subject=self.job.report.subjects[0]
        data={'csrf_token':token,'report_date':'2026-10-27','signature_coordinator':'Prof. Sample Coordinator',
              'signature_hod':'Dr. Sample HOD',f'faculty_{subject.id}':'Prof. Sample Faculty'}
        response=self.client.post(url,data=data,follow_redirects=True);self.assertEqual(response.status_code,200)
        model=build_faculty_report(self.job)
        self.assertEqual(model['subject_rows'][0]['faculty'],'Prof. Sample Faculty')
        workbook=load_workbook(io.BytesIO(build_excel_export(self.job).data))
        self.assertEqual(workbook['Subject Analysis']['D11'].value,'Prof. Sample Faculty')
        pdf=fitz.open(stream=build_pdf_export(self.job).data,filetype='pdf')
        text='\n'.join(p.get_text() for p in pdf)
        self.assertIn('Prof. Sample Faculty',' '.join(text.split()))
        self.assertIn('Prof. Sample Coordinator',' '.join(text.split()))
        self.assertIn('ALL STUDENTS RESULTS',text)
        self.assertIn('TOPPER LIST',text)
        self.assertNotIn('Students requiring attention',text)

    def test_pdf_and_all_report_pages_download(self):
        base_url=f'/result-data/exports/imports/{self.job.id}'
        for path,kind in [('results.pdf','application/pdf'),('results.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),('details.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')]:
            result=self.client.get(base_url+'/'+path)
            self.assertEqual(result.status_code,200,path)
            self.assertEqual(result.mimetype,kind)
        dashboard=self.client.get(f'/result-data/analysis/imports/{self.job.id}')
        self.assertEqual(dashboard.status_code,200)
        self.assertIn(b'Top 10 performers',dashboard.data)
        self.assertNotIn(b'Students requiring attention',dashboard.data)
        self.assertIn(b'Student 01, Student 02',dashboard.data)

    def test_source_text_is_not_an_excel_formula(self):
        self.job.report.students[0].name='=1+1'
        workbook=load_workbook(io.BytesIO(build_excel_export(self.job).data))
        self.assertEqual(workbook['Student Results']['C11'].value,'=1+1')
        self.assertEqual(workbook['Student Results']['C11'].data_type,'s')

    def test_reported_grades_are_used_instead_of_calculation_warnings(self):
        analysis=build_report_analysis(self.job.report)
        counts={r['grade']:r['count'] for r in analysis['grade_distribution']}
        self.assertEqual(counts['FF'],1)
        self.assertEqual(counts['BB'],26)
        self.job.report.students[0].subject_results[0].grade=None
        updated=build_report_analysis(self.job.report)
        self.assertEqual(next(r for r in updated['grade_distribution'] if r['grade']=='BB')['count'],25)
        self.assertEqual(build_faculty_report(self.job)['students'][0]['grades'][0],'BB')

    def test_no_eligible_sgpa_renders_empty_merit_without_invented_values(self):
        for student in self.job.report.students:student.summary.reported_sgpa=None
        analysis=build_report_analysis(self.job.report)
        self.assertEqual(analysis['top_performers'],[])
        self.assertIsNone(analysis['summary']['highest_sgpa'])
        self.assertEqual(filter_records(self.job.report,{'kind':'highest'})['count'],0)
        self.assertTrue(build_pdf_export(self.job).data.startswith(b'%PDF'))

    def test_wide_curriculum_pdf_retains_all_subjects_and_identifiers(self):
        for i in range(3,19):
            self.job.report.subjects.append(ResultSubject(subject_code='SUB'+str(i),
                subject_name='Long Example Engineering Subject '+str(i),display_order=i,credits=4,
                is_audit=False,total_maximum=100,total_minimum=40))
        self.db.session.flush()
        document=fitz.open(stream=build_pdf_export(self.job).data,filetype='pdf')
        text='\n'.join(page.get_text() for page in document)
        self.assertTrue(all(page.rect.width<page.rect.height for page in document))
        self.assertIn('SUB18',text)
        self.assertIn(self.job.report.students[0].seat_number,text)

    def test_settings_post_requires_csrf(self):
        self.assertEqual(self.client.post(f'/result-data/exports/imports/{self.job.id}/report-settings',data={'report_date':'2026-10-27'}).status_code,400)


if __name__=='__main__':unittest.main()
