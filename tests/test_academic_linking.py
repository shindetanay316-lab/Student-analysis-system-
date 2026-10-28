import io
import re
import unittest
from datetime import timedelta
from decimal import Decimal
from hashlib import sha256
from pathlib import Path
from openpyxl import load_workbook
import test_workflows as base
from models import db, Student, Subject, Enrollment, TheoryMarks, ExternalMarks, User, Semester
from result_data.models import ImportJob, ExamReport, ResultSubject, SubjectAssessment, ReportStudent, SubjectResult, ResultComponent, StudentSummary
from workflow_models import WorkflowReview, now
from coursework_models import AcceptedCourseworkMark
from teaching_models import TeachingAssignment, StaffAccessAudit
from academic_link_models import UniversityLinkedResult, UniversityPublication, StudentPortfolio


class AcademicLinkingTests(unittest.TestCase):
    def setUp(self):
        self.fixture = base.WorkflowTestCase()
        self.fixture.setUp()
        self.client = self.fixture.client
        self.job = self.dataset()

    def tearDown(self):
        self.fixture.tearDown()

    def token(self):
        from flask import g
        g.pop('csrf_token', None)
        return re.search(rb'name="csrf_token" value="([^"]+)"', self.client.get('/profile/password').data).group(1).decode()

    def post(self, path, fields=None):
        return self.client.post(path, data={**(fields or {}), 'csrf_token': self.token()})

    def dataset(self, suffix='', exam='REGULAR', semester='6', code='TH601', name='Theory Subject', ese=42, maximum=60):
        content = ('%PDF-1.4\nsynthetic source for linking tests ' + suffix).encode()
        job = ImportJob(original_filename='source' + suffix + '.pdf', stored_filename='link-source' + suffix + '.pdf', file_hash=sha256(content).hexdigest(),
            file_size_bytes=len(content), status='READY', confirmed_at=now())
        folder = Path(base.app.config['RESULT_DATA_UPLOAD_FOLDER'])
        folder.mkdir(exist_ok=True, parents=True)
        (folder / job.stored_filename).write_bytes(content)
        report = ExamReport(parser_name='DBATU', parser_version='2.0.1', university='DBATU', programme='BTech',
            examination_title='Semester ' + semester, semester=semester, examination_session='Summer 2026', examination_type=exam,
            validation_status='VALID', validation_error_count=0, validation_warning_count=0)
        job.report = report
        subject = ResultSubject(subject_code=code, subject_name=name, display_order=1, credits=4,
            is_audit=False, total_maximum=100, total_minimum=40)
        assessment = SubjectAssessment(component_type='ESE', display_name='ESE', source_row='ESE_0', maximum_marks=maximum, minimum_marks=20, display_order=1)
        subject.assessments.append(assessment)
        report.subjects.append(subject)
        for prn in ('S001', 'S002', 'S003'):
            local = db.session.get(Student, prn)
            student = ReportStudent(seat_number=prn, name=local.name, source_page=1, result_status='PASS', registration_scope='Whole', validation_status='VALID')
            student.summary = StudentSummary(reported_sgpa=8, reported_status='PASS')
            result = SubjectResult(subject=subject, source_page=1, total_raw='082', total_value=82, grade='AB', grade_raw='9/AB/36',
                grade_point=9, credit_points=36, validation_status='VALID', total_is_carried_forward=False)
            result.components.append(ResultComponent(assessment=assessment, component_type='ESE', raw_text=str(ese) if ese is not None else '--', numeric_value=ese, is_carried_forward=False))
            student.subject_results.append(result)
            report.students.append(student)
        db.session.add(job)
        db.session.commit()
        return job

    def preview(self, job=None, **fields):
        job = job or self.job
        return self.post(f'/academic/linking/{job.id}/preview', {'academic_year':'2025-26', **fields})

    def review(self, response):
        self.assertEqual(response.status_code, 302, response.data)
        self.assertIn('/academic/linking/reviews/', response.location)
        return WorkflowReview.query.filter_by(token=response.location.rsplit('/', 1)[-1]).one()

    def confirm(self, response, rows=('0', '1', '2'), copy=()):
        return self.post(response.location + '/confirm', dict(rows=list(rows), copy_ese=list(copy), verified='on'))

    def test_pages_render_and_unconfirmed_source_is_not_linkable(self):
        for path in ['/academic/linking', f'/academic/linking/{self.job.id}', '/academic/students', '/academic/students/S001']:
            self.assertEqual(self.client.get(path).status_code, 200, path)
        self.job.confirmed_at = None
        db.session.commit()
        response = self.preview()
        self.assertNotIn('/reviews/', response.location)
        self.assertEqual(WorkflowReview.query.count(), 0)

    def test_preview_no_academic_writes_then_history_only_confirm_and_idempotency(self):
        response = self.preview()
        review = self.review(response)
        self.assertTrue(all(r['ready'] for r in review.rows))
        self.assertEqual(UniversityLinkedResult.query.count(), 0)
        self.assertEqual(ExternalMarks.query.count(), 0)
        self.assertEqual(self.client.get(response.location).status_code, 200)
        self.assertEqual(self.confirm(response).status_code, 302)
        self.assertEqual(UniversityPublication.query.count(), 1)
        self.assertEqual(UniversityLinkedResult.query.count(), 3)
        self.assertEqual(ExternalMarks.query.count(), 0)
        self.assertEqual(self.client.get(response.location + '/source').data, review.source_workbook)
        self.assertEqual(self.confirm(response).status_code, 302)
        self.assertEqual(UniversityLinkedResult.query.count(), 3)
        self.assertIn(b'Approved university evidence', self.client.get('/academic/students/S001').data)
        self.assertIn(b'Linked', self.client.get(response.location).data)

    def test_explicit_ese_copy_preserves_internal_marks_and_uses_existing_calculation(self):
        mark = TheoryMarks(prn='S001', subject_code='TH601', semester_id=6, academic_year='2025-26', ct1=8, ct2=9, assignment=7, midsem=15)
        db.session.add(mark)
        db.session.commit()
        response = self.preview()
        self.assertEqual(self.confirm(response, rows=('0',), copy=('0',)).status_code, 302)
        external = ExternalMarks.query.one()
        self.assertEqual(external.external_marks, 42)
        self.assertTrue(external.locked)
        self.assertEqual((mark.ct1, mark.ct2, mark.assignment, mark.midsem), (8, 9, 7, 15))
        self.assertEqual(mark.total_marks, 74)  # Best two internal scores 17 + MSE 15 + ESE 42.
        self.assertEqual(AcceptedCourseworkMark.query.one().value, 42)
        self.assertTrue(UniversityLinkedResult.query.one().ese_copied)
        self.assertEqual(self.client.get('/workflow/subject?subject_code=TH601&semester_id=6&academic_year=2025-26').status_code, 200)

    def test_existing_mark_absence_and_legacy_theory_external_cannot_be_overwritten(self):
        db.session.add(ExternalMarks(prn='S001', subject_code='TH601', semester_id=6, academic_year='2025-26', external_marks=35, locked=True))
        db.session.add(ExternalMarks(prn='S002', subject_code='TH601', semester_id=6, academic_year='2025-26', external_marks=None, locked=True))
        db.session.add(TheoryMarks(prn='S003', subject_code='TH601', semester_id=6, academic_year='2025-26', external=31))
        db.session.commit()
        response = self.preview()
        self.assertTrue(all(r['ready'] and r['ese'] is None for r in self.review(response).rows))
        self.assertEqual(self.confirm(response, copy=('0',)).status_code, 409)
        self.assertEqual(UniversityLinkedResult.query.count(), 0)
        self.assertEqual(self.confirm(response).status_code, 302)
        self.assertEqual(ExternalMarks.query.filter_by(prn='S001').one().external_marks, 35)
        self.assertIsNone(ExternalMarks.query.filter_by(prn='S002').one().external_marks)

    def test_repeat_grade_only_carried_cancelled_and_nonstandard_scale_are_history_only(self):
        for suffix, modify in [
            ('repeat', lambda j: setattr(j.report, 'examination_type', 'REPEATER')),
            ('grade', lambda j: setattr(j.report.students[0].subject_results[0].components[0], 'numeric_value', None)),
            ('carried', lambda j: setattr(j.report.students[0].subject_results[0].components[0], 'is_carried_forward', True)),
            ('cancel', lambda j: setattr(j.report.students[0], 'result_status', 'CANCEL')),
            ('scale', lambda j: setattr(j.report.subjects[0].assessments[0], 'maximum_marks', 100)),
            ('precision', lambda j: setattr(j.report.students[0].subject_results[0].components[0], 'numeric_value', Decimal('42.25'))),
        ]:
            job = self.dataset(suffix)
            modify(job)
            db.session.commit()
            response = self.preview(job)
            row = self.review(response).rows[0]
            self.assertTrue(row['ready'], suffix)
            self.assertIsNone(row['ese'], suffix)
            self.assertEqual(self.confirm(response, rows=('0',)).status_code, 302)
        self.assertEqual(ExternalMarks.query.count(), 0)

    def test_zero_ese_is_numeric_and_does_not_become_pending(self):
        self.job.report.students[0].subject_results[0].components[0].numeric_value = 0
        db.session.commit()
        response = self.preview()
        self.assertEqual(Decimal(self.review(response).rows[0]['ese']), 0)
        self.confirm(response, rows=('0',), copy=('0',))
        self.assertEqual(ExternalMarks.query.one().external_marks, 0)

    def test_printed_grade_fallback_respects_intentionally_cleared_grades(self):
        from result_data.models import DataCorrection
        from academic_link_services import source_rows
        result = self.job.report.students[0].subject_results[0]
        result.grade = None
        db.session.commit()
        self.assertEqual(source_rows(self.job)[0]['grade_display'], 'AB')
        result.grade_raw = '|'
        db.session.commit()
        self.assertEqual(source_rows(self.job)[0]['grade_display'], 'NOT PRINTED')
        result.grade_raw = '9/AB/36'
        self.job.report.corrections.append(DataCorrection(
            student_id=result.student.id, subject_result_id=result.id,
            entity_type='subject_result', entity_label='S001 / TH601', field_name='grade',
            previous_value='AB', new_value='', reason='Reviewer cleared an ambiguous printed grade.'))
        db.session.commit()
        self.assertEqual(source_rows(self.job)[0]['grade_display'], 'REVIEW')
        response = self.preview()
        self.assertEqual(self.confirm(response, rows=('0',)).status_code, 302)
        self.assertEqual(UniversityLinkedResult.query.one().source['grade_display'], 'REVIEW')
        workbook = load_workbook(io.BytesIO(self.client.get('/academic/students/S001/export/excel').data))
        self.assertIn('REVIEW', [cell.value for row in workbook['University Evidence'] for cell in row])

    def test_name_prn_enrollment_and_credit_mismatches_remain_unmatched(self):
        self.job.report.students[0].name = 'Another person'
        self.job.report.students[1].seat_number = 'UNKNOWN'
        Enrollment.query.filter_by(prn='S003', subject_code='TH601').delete()
        db.session.commit()
        response = self.preview()
        self.assertTrue(all(not r['ready'] for r in self.review(response).rows))
        self.assertEqual(self.confirm(response).status_code, 409)
        self.assertEqual(UniversityLinkedResult.query.count(), 0)

    def test_alias_mapping_requires_reason_and_does_not_modify_source_or_local_subject(self):
        source = self.job.report.subjects[0]
        source.subject_code = 'PRINT601'
        source.subject_name = 'Printed subject alias'
        db.session.commit()
        response = self.preview(**{f'map_{source.id}': 'TH601'})
        self.assertFalse(self.review(response).rows[0]['ready'])
        response = self.preview(**{f'map_{source.id}': 'TH601', f'note_{source.id}': 'Verified university-to-college subject code alias.'})
        self.assertTrue(self.review(response).rows[0]['ready'])
        self.confirm(response, rows=('0',))
        self.assertEqual(UniversityLinkedResult.query.one().source['subject_code'], 'PRINT601')
        self.assertEqual(source.subject_code, 'PRINT601')
        self.assertEqual(db.session.get(Subject, 'TH601').subject_name, 'Theory Subject')

    def test_stale_source_and_target_marks_previews_fail_without_partial_changes(self):
        response = self.preview()
        self.job.report.students[0].subject_results[0].grade = 'BB'
        db.session.commit()
        self.assertEqual(self.confirm(response).status_code, 409)
        response = self.preview()
        db.session.add(ExternalMarks(prn='S001', subject_code='TH601', semester_id=6, academic_year='2025-26', external_marks=30, locked=True))
        db.session.commit()
        self.assertEqual(self.confirm(response, copy=('0',)).status_code, 409)
        self.assertEqual(UniversityLinkedResult.query.count(), 0)
        self.assertEqual(UniversityPublication.query.count(), 0)

    def test_tampered_missing_source_and_expired_preview_fail(self):
        response = self.preview()
        review = self.review(response)
        review.expires_at = now() - timedelta(seconds=1)
        db.session.commit()
        self.assertEqual(self.confirm(response).status_code, 409)
        path = Path(base.app.config['RESULT_DATA_UPLOAD_FOLDER']) / self.job.stored_filename
        path.write_bytes(b'changed')
        self.assertNotIn('/reviews/', self.preview().location)
        path.unlink()
        self.assertNotIn('/reviews/', self.preview().location)
        self.assertEqual(UniversityLinkedResult.query.count(), 0)

    def test_evidence_survives_source_edit_and_deletion_with_earlier_versions_retained(self):
        first = self.preview()
        self.confirm(first, rows=('0',))
        old = UniversityLinkedResult.query.one()
        self.job.report.students[0].subject_results[0].grade = 'BB'
        db.session.commit()
        second = self.preview()
        self.assertTrue(self.review(second).rows[0]['ready'])
        self.confirm(second, rows=('0',))
        self.assertEqual(UniversityLinkedResult.query.count(), 2)
        self.assertEqual(old.source['grade'], 'AB')
        from result_data.services.import_service import delete_import_job
        delete_import_job(self.job, base.app.config['RESULT_DATA_UPLOAD_FOLDER'])
        self.assertEqual(self.client.get('/academic/students/S001').status_code, 200)
        self.assertEqual(self.client.get(first.location + '/source').status_code, 200)

    def test_duplicate_import_preview_and_changed_context_cannot_duplicate_evidence(self):
        first = self.preview()
        self.confirm(first, rows=('0',))
        second = self.preview()
        self.assertTrue(self.review(second).rows[0]['duplicate'])
        self.assertEqual(self.confirm(second, rows=('0',)).status_code, 409)
        self.assertEqual(self.confirm(second, rows=('1', '2')).status_code, 302)
        self.assertEqual(UniversityLinkedResult.query.count(), 3)

    def test_selection_and_ese_tampering_require_explicit_verification(self):
        response = self.preview()
        for fields in [dict(rows=['0'], copy_ese=['1'], verified='on'), dict(rows=['999'], verified='on'), dict(rows=['0'])]:
            self.assertEqual(self.post(response.location + '/confirm', fields).status_code, 409)
        self.assertEqual(UniversityLinkedResult.query.count(), 0)
        self.assertEqual(self.client.post(response.location + '/confirm', data={}).status_code, 400)
        token = response.location.rsplit('/', 1)[-1]
        self.assertEqual(self.client.get('/workflow/reviews/' + token).status_code, 404)

    def test_teacher_scope_applies_to_history_exports_sources_and_portfolio_writes(self):
        response = self.preview()
        self.confirm(response)
        teacher = User.query.filter_by(username='teacher').one()
        db.session.add(TeachingAssignment(teacher_id=teacher.user_id, subject_code='LB601', semester_id=6, academic_year='2025-26'))
        Enrollment.query.filter_by(prn='S003', subject_code='LB601').delete()
        db.session.commit()
        self.fixture.logout()
        self.fixture.login('teacher', 'teacher123')
        self.assertEqual(self.client.get('/academic/linking').status_code, 403)
        self.assertEqual(self.client.get(response.location + '/source').status_code, 403)
        self.assertEqual(self.post(response.location + '/confirm', dict(verified='on')).status_code, 403)
        page = self.client.get('/academic/students/S001')
        self.assertEqual(page.status_code, 200)
        self.assertIn(b'Lab Subject', page.data)
        self.assertNotIn(b'Theory Subject', page.data)
        self.assertEqual(self.client.get('/academic/students/S003').status_code, 404)
        self.assertEqual(self.client.get('/academic/students/S003/export/excel').status_code, 404)
        self.assertEqual(self.post('/academic/students/S001/portfolio', dict(skills='Forged')).status_code, 403)
        export = self.client.get('/academic/students/S001/export/excel')
        book = load_workbook(io.BytesIO(export.data))
        text = str([[c.value for row in s for c in row] for s in book])
        self.assertNotIn('Theory Subject', text)
        self.assertNotIn('source.pdf', text)

    def test_portfolio_is_editable_audited_and_stale_edits_do_not_overwrite(self):
        self.assertEqual(self.post('/academic/students/S001/portfolio', dict(revision=0, interests='Robotics', skills='Java', projects='Attendance app')).status_code, 302)
        profile = StudentPortfolio.query.one()
        self.assertEqual(profile.details['skills'], 'Java')
        self.assertEqual(self.post('/academic/students/S001/portfolio', dict(revision=0, skills='Stale')).status_code, 303)
        self.assertEqual(profile.details['skills'], 'Java')
        self.assertEqual(self.post('/academic/students/S001/portfolio', dict(revision=1, skills='Python')).status_code, 302)
        self.assertEqual(profile.details['skills'], 'Python')
        self.assertEqual(StaffAccessAudit.query.filter_by(action='student_portfolio_updated').count(), 2)

    def test_history_exports_preserve_identity_order_source_grades_pending_and_text_safety(self):
        response = self.preview()
        self.confirm(response, rows=('0',))
        self.post('/academic/students/S001/portfolio', dict(revision=0, interests='=1+1', projects='Student project'))
        export = self.client.get('/academic/students/S001/export/excel')
        self.assertEqual(export.status_code, 200)
        book = load_workbook(io.BytesIO(export.data))
        self.assertEqual(book.sheetnames, ['Student Profile', 'Local Results', 'University Evidence', 'Attendance', 'Curriculum', 'Corrections'])
        self.assertEqual([book['University Evidence'].cell(1, i).value for i in range(1, 4)], ['Sr. No.', 'PRN', 'Student Name'])
        self.assertEqual(book['Student Profile']['B8'].data_type, 's')
        self.assertEqual(book['Student Profile']['B8'].value, '=1+1')
        self.assertIn('AB', [c.value for row in book['University Evidence'] for c in row])
        self.assertIn('Pending', [c.value for row in book['Local Results'] for c in row])
        pdf = self.client.get('/academic/students/S001/export/pdf')
        self.assertEqual(pdf.status_code, 200)
        import fitz
        doc = fitz.open(stream=pdf.data, filetype='pdf')
        text = ''.join(page.get_text() for page in doc)
        self.assertIn('Asha Patil', text)
        self.assertIn('University Evidence', text)
        self.assertIn('Student project', text)

    def test_actual_pdf_parser_confirmation_linking_and_history_end_to_end(self):
        from test_adaptive_result_parser import custom_document
        from result_data.services.import_service import create_import_job
        from result_data.services.extraction_service import extract_import_job
        from result_data.services.review_service import confirm_import
        from werkzeug.datastructures import FileStorage
        doc = custom_document()
        page = doc[0]
        boxes = page.search_for('JUNE - 2026')
        for box in boxes:
            page.add_redact_annot(box, fill=(1, 1, 1))
        page.apply_redactions()
        for box in boxes:
            page.insert_text((box.x0, box.y1 - 1.5), 'Summer - 2026', fontsize=7)
        content = doc.tobytes()
        doc.close()
        job = create_import_job(FileStorage(stream=io.BytesIO(content), filename='anonymized-integration.pdf'), base.app.config['RESULT_DATA_UPLOAD_FOLDER'])
        report = extract_import_job(job, base.app.config['RESULT_DATA_UPLOAD_FOLDER'])
        self.assertEqual(report.validation_error_count, 0)
        confirm_import(job, User.query.filter_by(username='admin').one(), acknowledge_warnings=True)
        db.session.add(Semester(semester_id=1, academic_year='2025-26', is_active=False))
        for source in report.subjects:
            db.session.add(Subject(subject_code=source.subject_code, subject_name=source.subject_name, semester_id=1,
                subject_type='THEORY', credits=int(source.credits), is_active=True))
        for source in report.students:
            db.session.add(Student(prn=source.seat_number, name=source.name, year_of_admission=2025))
        db.session.flush()
        for source in report.students:
            for subject in report.subjects:
                db.session.add(Enrollment(prn=source.seat_number, subject_code=subject.subject_code, semester_id=1, academic_year='2025-26'))
        db.session.commit()
        response = self.preview(job)
        review = self.review(response)
        self.assertTrue(all(r['ready'] and r['ese'] is not None for r in review.rows), review.rows)
        selected = tuple(str(r['index']) for r in review.rows)
        self.assertEqual(self.confirm(response, rows=selected, copy=selected).status_code, 302)
        self.assertEqual(ExternalMarks.query.count(), 2)
        self.assertEqual(UniversityLinkedResult.query.count(), 2)
        self.assertTrue(all(r.source['grade'] == 'CC' for r in UniversityLinkedResult.query))
        self.assertEqual(self.client.get('/academic/students/' + report.students[0].seat_number).status_code, 200)

    def test_credit_year_and_subject_change_protections_preserve_history(self):
        self.job.report.subjects[0].credits = 3
        db.session.commit()
        response = self.preview()
        self.assertTrue(all(not r['ready'] for r in self.review(response).rows))
        self.job.report.subjects[0].credits = 4
        db.session.add(Semester(semester_id=6, academic_year='2024-25', is_active=False))
        db.session.commit()
        wrong_year = self.preview(academic_year='2024-25')
        self.assertTrue(all(not r['ready'] for r in self.review(wrong_year).rows))
        response = self.preview()
        self.confirm(response, rows=('0',))
        Enrollment.query.filter_by(subject_code='TH601').delete()
        db.session.commit()
        self.client.post('/subjects/TH601/delete')
        self.assertIsNotNone(db.session.get(Subject, 'TH601'))
        self.assertFalse(db.session.get(Subject, 'TH601').is_active)

    def test_copied_ese_uses_existing_correction_workflow_and_retains_pdf_evidence(self):
        response = self.preview()
        self.confirm(response, rows=('0',), copy=('0',))
        request = self.post('/coursework/corrections/new', dict(subject_code='TH601', semester_id=6, academic_year='2025-26',
            prn='S001', category='EXTERNAL', component='EXTERNAL', new_value='41', reason='Verified against signed corrected university marks.'))
        self.assertIn('/coursework/corrections/', request.location)
        page = self.client.get(request.location)
        self.assertIn((response.location + '/source').encode(), page.data)
        self.assertEqual(self.post(request.location + '/decide', dict(decision='APPROVE', decision_note='Verified correction.')).status_code, 302)
        self.assertEqual(ExternalMarks.query.one().external_marks, 41)
        self.assertEqual(AcceptedCourseworkMark.query.one().value, 42)
        self.assertEqual(UniversityLinkedResult.query.one().source['components'][0]['value'], '42.00')


if __name__ == '__main__':
    unittest.main()
