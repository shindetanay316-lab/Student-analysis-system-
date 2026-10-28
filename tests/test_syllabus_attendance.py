import io
import re
import unittest
from copy import deepcopy
from datetime import date, time, timedelta
from pathlib import Path
from unittest.mock import patch
from urllib.parse import urlencode
from werkzeug.datastructures import FileStorage
from openpyxl import load_workbook
import test_workflows as base
from models import db, User, Subject, Student, Semester, Enrollment, TimetableSlot, AttendanceSession, AttendanceEntry, Attendance
from curriculum_models import CurriculumVersion, CurriculumBinding, StudentCurriculumTerm
from attendance_models import AttendanceCalendar, AttendanceHoliday, AttendanceChange
from workflow_models import WorkflowReview, now
from teaching_models import TeachingAssignment
from syllabus_parser import extract_syllabus
import curriculum_services as curricula
import curriculum_enrollment as curriculum_roster
import attendance_services as attendance


def course(code='NEW301', **changes):
    row = dict(row_id=1, semester_id=3, group='COMMON', code=code, name='New course', category='PCC', subject_type='THEORY',
        credits='3', l='3', t='0', p='0', ca='20', mse='20', ese='60', total='100', is_audit=False,
        choice_group='', is_bucket=False, include=True, evidence='', page=1, source_text='Printed evidence', issues=[])
    row.update(changes)
    return row


class SyllabusParserTests(unittest.TestCase):
    def test_real_first_year_schemes_keep_missing_group_b_codes_and_options(self):
        rows, warnings = extract_syllabus((Path(__file__).parent / 'fixtures/first_year_syllabus_schemes.pdf').read_bytes())
        self.assertEqual(len(rows), 54)
        self.assertEqual(sum(r['group'] == 'B' for r in rows), 27)
        self.assertTrue(all(not r['code'] and not r['include'] for r in rows if r['group'] == 'B'))
        coded = {r['code']: r for r in rows if r['code']}
        self.assertEqual(coded['24AF1000CC111B']['name'], 'NSS-I')
        self.assertEqual(coded['24AF1CHEBS102']['name'], 'Engineering Chemistry')
        self.assertEqual(coded['24AF1EMES104']['l'], '3')
        self.assertEqual(coded['24AF1EEEES209']['credits'], '0')
        self.assertTrue(coded['24AF1EEEES209']['is_audit'])
        self.assertEqual(warnings, [])

    def test_real_second_year_schemes_join_multiline_names_and_choice_rows(self):
        rows, warnings = extract_syllabus((Path(__file__).parent / 'fixtures/second_year_syllabus_schemes.pdf').read_bytes())
        self.assertEqual(len(rows), 27)
        coded = {r['code']: r for r in rows}
        self.assertEqual(coded['25AF1844PC302']['name'], 'Analog and Digital Electronics')
        self.assertEqual(coded['25AF1000AE410A']['name'], 'Marathi')
        self.assertEqual(coded['25AF1000AE410C']['name'], 'Sanskrit')
        self.assertEqual(coded['25AF1844VS411A']['p'], '4')
        self.assertEqual(coded['25AF1844VS411B']['name'], 'PCB Designing')
        self.assertTrue(coded['25AF1XXXOE305']['is_bucket'])
        self.assertEqual(coded['25AF1COIVE407']['credits'], '0')
        self.assertEqual(warnings, [])

    def test_bad_pdf_has_actionable_error(self):
        with self.assertRaisesRegex(ValueError, 'PDF'):
            extract_syllabus(b'not a pdf')


class CombinedWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.fixture = base.WorkflowTestCase()
        self.fixture.setUp()
        self.client = self.fixture.client
        self.c = dict(subject_code='TH601', semester_id=6, academic_year='2025-26', division='A', batch='ALL')
        self.teacher = User.query.filter_by(username='teacher').one()
        db.session.add(TeachingAssignment(teacher_id=self.teacher.user_id, subject_code='TH601', semester_id=6, academic_year='2025-26'))
        self.slot = TimetableSlot(**self.c, day_of_week='Monday', start_time=time(10), end_time=time(11), session_type='THEORY', is_active=True)
        db.session.add(self.slot)
        db.session.commit()

    def tearDown(self):
        self.fixture.tearDown()

    def token(self):
        from flask import g
        g.pop('csrf_token', None)
        page = self.client.get('/profile/password')
        return re.search(rb'name="csrf_token" value="([^"]+)"', page.data).group(1).decode()

    def post(self, url, data=None):
        return self.client.post(url, data={**(data or {}), 'csrf_token': self.token()})

    def version(self, rows=None):
        rows = rows or [course()]
        v = CurriculumVersion(title='Test curriculum', branch='ECE', scheme='NEP/MIS', effective_year='2024-25',
            owner_id=User.query.filter_by(username='admin').one().user_id, source_name='source.pdf', source_hash='a' * 64,
            source_pdf=b'source pdf bytes', courses=deepcopy(rows), extracted=deepcopy(rows), warnings=[])
        db.session.add(v)
        db.session.commit()
        return v

    def approve(self, v):
        return self.post(f'/curriculum/{v.token}/approve', dict(revision=v.revision, verified='on'))

    def workbook(self, c=None, start='2025-09-01', end='2025-09-01', values=None):
        c = c or self.c
        response = self.client.get('/attendance/download-template?' + urlencode(dict(c, start_date=start, end_date=end)))
        self.assertEqual(response.status_code, 200, response.data[:200])
        wb = load_workbook(io.BytesIO(response.data))
        if values is not None:
            for i, row in enumerate(values, 5):
                for j, value in enumerate(row, 4):
                    wb['Attendance'].cell(i, j, value)
        stream = io.BytesIO()
        wb.save(stream)
        return stream.getvalue()

    def preview(self, content):
        return self.post('/attendance/upload', {'attendance_file': (io.BytesIO(content), 'attendance.xlsx')})

    def confirm(self, response):
        self.assertEqual(response.status_code, 302)
        self.assertIn('/attendance/reviews/', response.location)
        return self.post(response.location + '/confirm')

    def test_new_pages_and_forms_render_with_csrf(self):
        for path in ['/curriculum', '/attendance', '/attendance?' + urlencode(self.c)]:
            response = self.client.get(path)
            self.assertEqual(response.status_code, 200, path)
            self.assertIn(b'csrf_token', response.data)
        self.assertEqual(self.client.post('/attendance/calendar', data={}).status_code, 400)
        self.assertEqual(self.client.post('/curriculum/upload', data={}).status_code, 400)

    def test_syllabus_upload_retains_pdf_but_does_not_create_subjects_until_approval(self):
        content = (Path(__file__).parent / 'fixtures/second_year_syllabus_schemes.pdf').read_bytes()
        before = Subject.query.count()
        response = self.post('/curriculum/upload', dict(title='NEP second year', branch='ECE', scheme='MIS', effective_year='2025-2026',
            syllabus=(io.BytesIO(content), 'syllabus.pdf')))
        self.assertEqual(response.status_code, 302)
        version = CurriculumVersion.query.one()
        self.assertEqual(version.source_pdf, content)
        self.assertEqual(version.scheme, 'NEP/MIS')
        self.assertEqual(Subject.query.count(), before)
        self.assertEqual(self.client.get(response.location).status_code, 200)
        self.approve(version)
        self.assertEqual(version.status, 'APPROVED')
        self.assertEqual(Subject.query.count(), before + 27)
        self.assertEqual(self.client.get(f'/curriculum/{version.token}/source').data, content)

    def test_approval_is_immutable_idempotent_and_revision_preserves_original(self):
        v = self.version()
        self.approve(v)
        subject = db.session.get(Subject, 'NEW301')
        self.assertIsNotNone(subject)
        self.approve(v)
        self.assertEqual(Subject.query.filter_by(subject_code='NEW301').count(), 1)
        response = self.post(f'/curriculum/{v.token}/save', {'revision': v.revision, 'r1_name': 'Forged'})
        self.assertEqual(response.status_code, 303)
        self.assertEqual(v.courses[0]['name'], 'New course')
        revised = self.post(f'/curriculum/{v.token}/revise')
        self.assertEqual(revised.status_code, 302)
        self.assertEqual(CurriculumVersion.query.count(), 2)
        draft = CurriculumVersion.query.filter_by(status='DRAFT').one()
        self.assertEqual(draft.predecessor, v.token)
        self.assertEqual(draft.source_pdf, v.source_pdf)

    def test_missing_codes_credit_errors_and_subject_conflicts_cannot_overwrite(self):
        for row in [course(code=''), course(credits='nan'), course(credits='2.5'), course(code='TH601')]:
            v = self.version([row])
            response = self.approve(v)
            self.assertEqual(response.status_code, 303)
            self.assertEqual(v.status, 'DRAFT')
        self.assertEqual(db.session.get(Subject, 'TH601').subject_name, 'Theory Subject')
        self.assertEqual(db.session.get(Subject, 'TH601').credits, 4)

    def test_missing_code_requires_evidence_and_stale_draft_cannot_replace_newer(self):
        v = self.version([course(code='', include=False)])
        saved = self.post(f'/curriculum/{v.token}/save', {'revision': 1, 'r1_code': 'CHECK301', 'r1_include': 'on'})
        self.assertEqual(saved.status_code, 302)
        self.assertEqual(v.revision, 2)
        self.assertEqual(self.approve(v).status_code, 303)
        stale = self.post(f'/curriculum/{v.token}/save', {'revision': 1, 'r1_name': 'Stale edit'})
        self.assertEqual(stale.status_code, 303)
        self.assertEqual(v.courses[0]['code'], 'CHECK301')
        rows = deepcopy(v.courses)
        rows[0]['evidence'] = 'Official result sheet, page 1'
        v.courses = rows
        db.session.commit()
        self.assertEqual(self.approve(v).status_code, 302)
        self.assertEqual(v.status, 'APPROVED')

    def test_binding_and_reviewed_dse_enrollment_preserve_actual_admission_and_existing_marks(self):
        v = self.version()
        self.approve(v)
        self.post(f'/curriculum/{v.token}/bind', dict(regular_cohort=2023, semester_id=3, academic_year='2024-25', group='COMMON'))
        binding = CurriculumBinding.query.one()
        Student.query.filter_by(prn='S002').one().year_of_admission = 2024
        db.session.commit()
        response = self.post(f'/curriculum/bindings/{binding.id}/preview', dict(prns=['S001', 'S002']))
        self.assertEqual(StudentCurriculumTerm.query.count(), 0)
        self.assertEqual(Enrollment.query.filter_by(subject_code='NEW301').count(), 0)
        self.assertEqual(self.client.get(response.location).status_code, 200)
        confirmed = self.post(response.location + '/confirm')
        self.assertEqual(confirmed.status_code, 302)
        self.assertEqual(StudentCurriculumTerm.query.count(), 2)
        self.assertEqual(Enrollment.query.filter_by(subject_code='NEW301').count(), 2)
        self.assertEqual(db.session.get(Student, 'S002').year_of_admission, 2024)
        self.post(response.location + '/confirm')
        self.assertEqual(Enrollment.query.filter_by(subject_code='NEW301').count(), 2)
        self.assertTrue(db.session.get(Semester, (6, '2025-26')).is_active)

    def test_curriculum_choices_require_one_option_and_reject_conflicting_existing_enrollment(self):
        v = self.version([course(), course(code='OP301A', row_id=2, name='Option A', choice_group='S3-COMMON-7'),
            course(code='OP301B', row_id=3, name='Option B', choice_group='S3-COMMON-7')])
        self.approve(v)
        self.post(f'/curriculum/{v.token}/bind', dict(regular_cohort=2023, semester_id=3, academic_year='2024-25', group='COMMON'))
        binding = CurriculumBinding.query.one()
        bad = self.post(f'/curriculum/bindings/{binding.id}/preview', dict(prns=['S001']))
        self.assertNotIn('roster-reviews', bad.location)
        good = self.post(f'/curriculum/bindings/{binding.id}/preview', dict(prns=['S001'], options=['OP301A']))
        self.post(good.location + '/confirm')
        bad = self.post(f'/curriculum/bindings/{binding.id}/preview', dict(prns=['S001'], options=['OP301B']))
        self.assertNotIn('roster-reviews', bad.location)
        self.assertEqual(Enrollment.query.filter_by(subject_code='OP301B').count(), 0)

    def test_teacher_cannot_manage_syllabus_or_calendar_and_assignment_limits_all_attendance(self):
        v = self.version()
        self.fixture.logout()
        self.fixture.login('teacher', 'teacher123')
        self.assertEqual(self.client.get('/curriculum').status_code, 403)
        self.assertEqual(self.client.get('/curriculum/' + v.token).status_code, 403)
        for path in ['/attendance/calendar', '/attendance/holidays', '/curriculum/' + v.token + '/approve']:
            self.assertEqual(self.post(path, {}).status_code, 403)
        denied = dict(self.c, subject_code='LB601')
        for path in ['/attendance/report/excel', '/attendance/report/pdf']:
            self.assertEqual(self.client.get(path + '?' + urlencode(denied)).status_code, 403)
        self.assertEqual(self.post('/attendance/sessions/makeup', dict(denied, date='2025-09-02', start_time='12:00', end_time='13:00', reason='Extra class')).status_code, 403)

    def test_attendance_preview_no_writes_confirm_idempotent_and_pending_columns_ignored(self):
        content = self.workbook(end='2025-09-08', values=[['P', None], ['A', None]])
        response = self.preview(content)
        self.assertEqual(AttendanceEntry.query.count(), 0)
        self.assertEqual(AttendanceSession.query.count(), 0)
        self.assertEqual(self.client.get(response.location).status_code, 200)
        self.confirm(response)
        self.assertEqual(AttendanceEntry.query.count(), 2)
        self.assertEqual(AttendanceSession.query.count(), 1)
        self.assertEqual(Attendance.query.filter_by(prn='S001').one().percentage, 100)
        self.assertEqual(Attendance.query.filter_by(prn='S002').one().percentage, 0)
        self.assertEqual(self.client.get(response.location + '/source').data, content)
        self.confirm(response)
        self.assertEqual(AttendanceEntry.query.count(), 2)
        self.assertEqual(AttendanceChange.query.count(), 1)

    def test_blank_partial_formula_bad_names_and_future_attendance_rejected_atomically(self):
        for values in [[[None], [None]], [['P'], [None]], [['=1+1'], ['A']], [['BAD'], ['P']]]:
            response = self.preview(self.workbook(values=values))
            self.assertNotIn('/attendance/reviews/', response.location)
            self.assertEqual(AttendanceEntry.query.count(), 0)
            self.assertEqual(AttendanceSession.query.count(), 0)
        with patch('attendance_services.local_now') as clock:
            from datetime import datetime
            from zoneinfo import ZoneInfo
            clock.return_value = datetime(2025, 9, 1, 10, 30, tzinfo=ZoneInfo('Asia/Kolkata'))
            response = self.preview(self.workbook(values=[['P'], ['A']]))
            self.assertNotIn('/attendance/reviews/', response.location)
            legacy = AttendanceSession(**self.c, lecture_date=date(2025, 9, 1), start_time=time(10), end_time=time(11), session_type='THEORY', status='COMPLETED')
            db.session.add(legacy)
            db.session.flush()
            db.session.add(AttendanceEntry(session_id=legacy.id, prn='S001', status='P'))
            db.session.flush()
            self.assertEqual(attendance.report_rows(self.c)[0]['total'], 0)
            db.session.rollback()
        content = self.workbook(values=[['P'], ['A']])
        wb = load_workbook(io.BytesIO(content))
        wb['Attendance']['C5'] = 'Different student'
        stream = io.BytesIO()
        wb.save(stream)
        self.assertNotIn('/attendance/reviews/', self.preview(stream.getvalue()).location)

    def test_same_day_distinct_classes_count_twice_and_batch_reuploads_do_not_duplicate(self):
        db.session.add(TimetableSlot(**self.c, day_of_week='Monday', start_time=time(12), end_time=time(13), session_type='THEORY', is_active=True))
        db.session.commit()
        self.confirm(self.preview(self.workbook(values=[['P', 'A'], ['A', 'P']])))
        self.assertEqual(AttendanceSession.query.count(), 2)
        self.assertEqual(Attendance.query.filter_by(prn='S001').one().percentage, 50)
        batch = dict(self.c, batch='A1')
        self.confirm(self.preview(self.workbook(c=batch, values=[['P', 'A']])))
        self.assertEqual(AttendanceSession.query.count(), 2)
        self.assertEqual(Attendance.query.filter_by(prn='S001').one().total_lectures, 2)

    def test_leave_and_official_duty_retained_and_not_defaulters_without_counted_classes(self):
        self.confirm(self.preview(self.workbook(values=[['L'], ['O']])))
        rows = attendance.report_rows(self.c)
        self.assertEqual([r['total'] for r in rows], [0, 0])
        self.assertEqual(rows[0]['leave'], 1)
        self.assertEqual(rows[1]['official'], 1)
        response = self.client.get('/attendance/defaulters/excel?' + urlencode(self.c))
        wb = load_workbook(io.BytesIO(response.data))
        self.assertEqual(wb.active.max_row, 3)

    def test_calendar_holiday_and_cancellation_exclusions_and_zero_summary_reset(self):
        self.post('/attendance/calendar', dict(semester_id=6, academic_year='2025-26', start_date='2025-09-01', end_date='2025-09-30'))
        self.post('/attendance/holidays', dict(semester_id=6, academic_year='2025-26', holiday_date='2025-09-08', division='ALL', reason='College holiday'))
        plans = attendance.planned(self.c, date(2025, 9, 1), date(2025, 9, 8))
        self.assertEqual(len(plans), 1)
        self.confirm(self.preview(self.workbook(values=[['P'], ['P']])))
        self.post('/attendance/sessions/cancel', dict(self.c, date='2025-09-01', start_time='10:00', end_time='11:00', slot_division='A', slot_batch='ALL', reason='Class cancelled by department'))
        self.assertEqual(AttendanceSession.query.one().status, 'CANCELLED')
        self.assertEqual(AttendanceEntry.query.count(), 2)
        self.assertEqual(Attendance.query.filter_by(prn='S001').one().total_lectures, 0)
        self.assertEqual(attendance.planned(self.c, date(2025, 9, 1), date(2025, 9, 8)), [])
        self.assertIsNone(attendance.report_rows(self.c)[0]['percentage'])

    def test_makeup_class_counts_only_after_confirmation_and_does_not_overlap(self):
        response = self.post('/attendance/sessions/makeup', dict(self.c, date='2025-09-02', start_time='12:00', end_time='13:00', reason='Makeup for cancelled lecture'))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(AttendanceSession.query.one().status, 'EXTRA')
        self.assertEqual(attendance.report_rows(self.c)[0]['total'], 0)
        self.confirm(self.preview(self.workbook(start='2025-09-02', end='2025-09-02', values=[['P'], ['A']])))
        self.assertEqual(attendance.report_rows(self.c)[0]['total'], 1)
        self.post('/attendance/sessions/makeup', dict(self.c, date='2025-09-01', start_time='10:30', end_time='11:30', reason='Should be blocked overlap'))
        self.assertEqual(AttendanceSession.query.count(), 1)

    def test_changed_roster_and_revoked_assignment_block_confirmation_and_other_owner_blocked(self):
        self.fixture.logout()
        self.fixture.login('teacher', 'teacher123')
        response = self.preview(self.workbook(values=[['P'], ['A']]))
        from term_services import placement
        placement('S001', self.c['semester_id'], self.c['academic_year']).batch = 'A9'
        db.session.commit()
        self.assertEqual(self.confirm(response).status_code, 409)
        self.assertEqual(AttendanceEntry.query.count(), 0)
        placement('S001', self.c['semester_id'], self.c['academic_year']).batch = 'A1'
        db.session.commit()
        response = self.preview(self.workbook(values=[['P'], ['A']]))
        TeachingAssignment.query.one().is_active = False
        db.session.commit()
        self.assertEqual(self.post(response.location + '/confirm').status_code, 403)
        self.assertEqual(AttendanceEntry.query.count(), 0)

    def test_stale_calendar_changed_entries_and_expired_preview_cannot_commit(self):
        first = self.preview(self.workbook(values=[['P'], ['A']]))
        second = self.preview(self.workbook(values=[['A'], ['P']]))
        self.confirm(first)
        self.assertEqual(self.confirm(second).status_code, 409)
        self.assertEqual(AttendanceEntry.query.filter_by(prn='S001').one().status, 'P')
        third = self.preview(self.workbook(values=[['P'], ['A']]))
        review = WorkflowReview.query.filter_by(token=third.location.rsplit('/', 1)[-1]).one()
        review.expires_at = now() - timedelta(seconds=1)
        db.session.commit()
        self.assertEqual(self.confirm(third).status_code, 409)

    def test_reports_keep_business_column_order_pending_and_register_and_pdf(self):
        self.confirm(self.preview(self.workbook(values=[['P'], ['A']])))
        c = dict(self.c, start_date='2025-09-01', end_date='2025-09-01')
        response = self.client.get('/attendance/report/excel?' + urlencode(c))
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(io.BytesIO(response.data))
        self.assertEqual(wb.sheetnames, ['Attendance Summary', 'Session Register'])
        self.assertEqual([wb.active.cell(4, i).value for i in range(1, 4)], ['Sr. No.', 'PRN', 'Student Name'])
        self.assertEqual(wb['Session Register'].max_row, 3)
        response = self.client.get('/attendance/report/pdf?' + urlencode(c))
        self.assertEqual(response.status_code, 200)
        import fitz
        pdf = fitz.open(stream=response.data, filetype='pdf')
        text = ''.join(p.get_text() for p in pdf)
        self.assertIn('Asha Patil', text)
        self.assertIn('Bharat More', text)

    def test_division_specific_holiday_does_not_hide_other_division_or_add_absences(self):
        self.slot.division = 'ALL'
        db.session.commit()
        all_scope = dict(self.c, division='ALL')
        self.post('/attendance/holidays', dict(semester_id=6, academic_year='2025-26', holiday_date='2025-09-01', division='A', reason='Division A non-teaching day'))
        content = self.workbook(c=all_scope, values=[['N/A'], ['N/A'], ['P']])
        self.confirm(self.preview(content))
        self.assertEqual(AttendanceEntry.query.count(), 1)
        self.assertEqual(AttendanceEntry.query.one().prn, 'S003')
        self.assertIsNone(attendance.report_rows(all_scope)[0]['percentage'])

    def test_hidden_metadata_tampering_and_other_review_owner_are_denied(self):
        self.fixture.logout()
        self.fixture.login('teacher', 'teacher123')
        content = self.workbook(values=[['P'], ['A']])
        wb = load_workbook(io.BytesIO(content))
        wb['_meta']['B3'] = 'LB601'
        stream = io.BytesIO()
        wb.save(stream)
        self.assertEqual(self.preview(stream.getvalue()).status_code, 403)
        response = self.preview(content)
        from models import User
        other = User(username='otherteacher', role='TEACHER', is_active=True)
        other.set_password('otherteacher123')
        db.session.add(other)
        db.session.flush()
        db.session.add(TeachingAssignment(teacher_id=other.user_id, subject_code='TH601', semester_id=6, academic_year='2025-26'))
        db.session.commit()
        self.fixture.logout()
        self.fixture.login('otherteacher', 'otherteacher123')
        self.assertEqual(self.client.get(response.location).status_code, 403)
        self.assertEqual(self.client.get(response.location + '/source').status_code, 403)
        self.assertEqual(self.post(response.location + '/confirm').status_code, 403)
        self.assertEqual(AttendanceEntry.query.count(), 0)

    def test_calendar_changes_invalidate_preview_and_cannot_hide_recorded_classes(self):
        response = self.preview(self.workbook(values=[['P'], ['A']]))
        self.post('/attendance/holidays', dict(semester_id=6, academic_year='2025-26', holiday_date='2025-09-01', division='ALL', reason='Official holiday'))
        self.assertEqual(self.confirm(response).status_code, 409)
        self.assertEqual(AttendanceEntry.query.count(), 0)
        holiday = AttendanceHoliday.query.one()
        self.post(f'/attendance/holidays/{holiday.id}/remove')
        self.confirm(self.preview(self.workbook(values=[['P'], ['A']])))
        self.post('/attendance/holidays', dict(semester_id=6, academic_year='2025-26', holiday_date='2025-09-01', division='ALL', reason='Should require explicit cancellation'))
        self.assertFalse(AttendanceHoliday.query.one().is_active)
        self.post('/attendance/calendar', dict(semester_id=6, academic_year='2025-26', start_date='2025-09-02', end_date='2025-09-30'))
        self.assertEqual(AttendanceCalendar.query.count(), 0)

    def test_timetable_cannot_turn_credit_course_into_activity_or_double_book(self):
        template = self.client.get('/attendance/download-timetable-template')
        wb = load_workbook(io.BytesIO(template.data))
        ws = wb['Class Timetable']
        ws['B4'], ws['D4'], ws['H4'] = '2025-26', 6, 'A'
        for row in range(7, 13):
            for col in range(2, 10):
                ws.cell(row, col, '')
        ws['B7'] = 'TH601 - Theory Subject [ATTENDANCE_ONLY]'
        stream = io.BytesIO()
        wb.save(stream)
        self.post('/attendance/upload-timetable', dict(timetable_file=(io.BytesIO(stream.getvalue()), 'timetable.xlsx')))
        self.assertEqual(db.session.get(Subject, 'TH601').credits, 4)
        self.assertFalse(db.session.get(Subject, 'TH601').is_attendance_only)
        self.assertEqual(TimetableSlot.query.filter_by(is_active=True).count(), 1)
        ws['B7'] = 'TH601 - Theory Subject\nLB601 - Lab Subject'
        stream = io.BytesIO()
        wb.save(stream)
        self.post('/attendance/upload-timetable', dict(timetable_file=(io.BytesIO(stream.getvalue()), 'timetable.xlsx')))
        self.assertEqual(TimetableSlot.query.filter_by(is_active=True).count(), 1)

    def test_approved_course_cannot_be_edited_or_deleted_through_legacy_subject_form(self):
        v = self.version()
        self.approve(v)
        db.session.add(Semester(semester_id=3, academic_year='2024-25', is_active=False))
        db.session.commit()
        self.client.post('/subjects/save', data=dict(original_subject_code='NEW301', subject_code='NEW301',
            subject_name='Changed', semester_id=3, subject_type='THEORY', credits=9, is_active='on'))
        self.assertEqual(db.session.get(Subject, 'NEW301').credits, 3)
        self.client.post('/subjects/NEW301/delete')
        self.assertIsNotNone(db.session.get(Subject, 'NEW301'))

    def test_curriculum_roster_stale_preview_and_wrong_admission_are_rejected(self):
        v = self.version()
        self.approve(v)
        self.post(f'/curriculum/{v.token}/bind', dict(regular_cohort=2023, semester_id=3, academic_year='2024-25', group='COMMON'))
        binding = CurriculumBinding.query.one()
        db.session.get(Student, 'S003').year_of_admission = 2020
        db.session.commit()
        bad = self.post(f'/curriculum/bindings/{binding.id}/preview', dict(prns=['S003']))
        self.assertNotIn('roster-reviews', bad.location)
        response = self.post(f'/curriculum/bindings/{binding.id}/preview', dict(prns=['S001']))
        db.session.get(Student, 'S001').name = 'Changed name'
        db.session.commit()
        self.assertEqual(self.post(response.location + '/confirm').status_code, 303)
        self.assertEqual(StudentCurriculumTerm.query.count(), 0)
        self.assertEqual(Enrollment.query.filter_by(subject_code='NEW301').count(), 0)

    def test_legacy_duplicate_classes_are_counted_once_and_conflicts_need_review(self):
        for scope, status in [('ALL', 'P'), ('A1', 'P')]:
            session = AttendanceSession(**{**self.c, 'batch': scope}, lecture_date=date(2025, 9, 1), start_time=time(10), end_time=time(11), session_type='THEORY', status='COMPLETED')
            db.session.add(session)
            db.session.flush()
            db.session.add(AttendanceEntry(session_id=session.id, prn='S001', status=status))
        db.session.commit()
        self.assertEqual(attendance.report_rows(self.c)[0]['total'], 1)
        AttendanceEntry.query.order_by(AttendanceEntry.id.desc()).first().status = 'A'
        db.session.commit()
        row = attendance.report_rows(self.c)[0]
        self.assertEqual(row['total'], 0)
        self.assertEqual(row['conflicts'], 1)
        self.confirm(self.preview(self.workbook(values=[['P'], ['P']])))
        row = attendance.report_rows(self.c)[0]
        self.assertEqual(row['total'], 1)
        self.assertEqual(row['conflicts'], 0)
        self.assertEqual(AttendanceEntry.query.count(), 3)  # Old copies retained.

    def test_digit_prefixed_nep_subject_codes_work_in_the_timetable(self):
        content = (Path(__file__).parent / 'fixtures/second_year_syllabus_schemes.pdf').read_bytes()
        self.post('/curriculum/upload', dict(title='NEP second year', branch='ECE', scheme='NEP/MIS', effective_year='2025-26', syllabus=(io.BytesIO(content), 'syllabus.pdf')))
        v = CurriculumVersion.query.one()
        self.approve(v)
        db.session.add(Semester(semester_id=3, academic_year='2025-26', is_active=False))
        db.session.commit()
        template = self.client.get('/attendance/download-timetable-template?semester_id=3&academic_year=2025-26')
        wb = load_workbook(io.BytesIO(template.data))
        self.assertIn('Subject Codes', wb.sheetnames)
        ws = wb['Class Timetable']
        self.assertEqual(ws['D4'].value, 3)
        for row in range(7, 13):
            for col in range(2, 10):
                ws.cell(row, col, '')
        ws['B7'] = '25AF1844PC302 - Analog and Digital Electronics'
        stream = io.BytesIO()
        wb.save(stream)
        self.post('/attendance/upload-timetable', dict(timetable_file=(io.BytesIO(stream.getvalue()), 'timetable.xlsx')))
        self.assertEqual(TimetableSlot.query.filter_by(subject_code='25AF1844PC302', is_active=True).count(), 1)
        from attendance_utils import _session_type_from_text
        subject = db.session.get(Subject, 'TH601')
        self.assertEqual(_session_type_from_text('TH601 - Software Engineering', subject), 'THEORY')


if __name__ == '__main__':
    unittest.main()
