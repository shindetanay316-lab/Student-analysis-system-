from pathlib import Path

from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .analysis_service import build_report_analysis


DATA_HEADER_ROW = 9
DATA_START_ROW = DATA_HEADER_ROW + 1

TRUST_NAME = "Chhatrapati Shahu Maharaj Shikshan Sanstha's"
COLLEGE_NAME = "CSMSS Chh. Shahu College of Engineering"
COLLEGE_ADDRESS = "Kanchanwadi, Chhatrapati Sambhajinagar – 431011"
DEPARTMENT_NAME = "Department of Electronics and Computer Engineering"

NAVY = "17365D"
BLUE = "1F4E78"
TEAL = "0F766E"
GOLD = "D6A84B"
PALE_BLUE = "D9EAF7"
PALE_GREEN = "E2F0D9"
PALE_YELLOW = "FFF4CC"
PALE_RED = "FCE8E6"
WHITE = "FFFFFF"
LIGHT_BORDER = Side(style="thin", color="D9E2F3")

_IMAGES_DIRECTORY = (
    Path(__file__).resolve().parents[2] / "static" / "result_data" / "images"
)
_LEFT_LOGO = _IMAGES_DIRECTORY / "college-left-logo.jpg"
_RIGHT_LOGO = _IMAGES_DIRECTORY / "college-right-logo.jpg"


def add_brand_header(sheet, report, sheet_title: str) -> None:
    """Add the official college identity above a worksheet's content."""
    sheet.sheet_view.showGridLines = False

    for row, height in {1: 22, 2: 28, 3: 20, 4: 22, 5: 22, 6: 27, 7: 22}.items():
        sheet.row_dimensions[row].height = height

    sheet.column_dimensions["A"].width = max(sheet.column_dimensions["A"].width or 0, 16)
    sheet.column_dimensions["I"].width = max(sheet.column_dimensions["I"].width or 0, 16)

    _add_logo(sheet, _LEFT_LOGO, "A1", width=72, height=96)
    _add_logo(sheet, _RIGHT_LOGO, "I1", width=96, height=80)

    for range_address, value, size, bold in (
        ("B1:H1", TRUST_NAME, 12, False),
        ("B2:H2", COLLEGE_NAME, 18, True),
        ("B3:H3", COLLEGE_ADDRESS, 11, False),
        ("B4:H5", DEPARTMENT_NAME, 15, True),
    ):
        sheet.merge_cells(range_address)
        cell = sheet[range_address.split(":", 1)[0]]
        cell.value = value
        cell.font = Font(name="Times New Roman", size=size, bold=bold, color=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in sheet.iter_rows(min_row=1, max_row=5, min_col=1, max_col=9):
        for cell in row:
            cell.border = Border(
                left=LIGHT_BORDER,
                right=LIGHT_BORDER,
                top=LIGHT_BORDER,
                bottom=LIGHT_BORDER,
            )

    sheet.merge_cells("A6:I6")
    title = sheet["A6"]
    title.value = sheet_title.upper()
    title.font = Font(name="Aptos Display", size=14, bold=True, color=WHITE)
    title.fill = PatternFill("solid", fgColor=NAVY)
    title.alignment = Alignment(horizontal="center", vertical="center")
    title.border = Border(bottom=Side(style="medium", color=GOLD))

    sheet.merge_cells("A7:I7")
    subtitle = sheet["A7"]
    subtitle.value = _report_context(report)
    subtitle.font = Font(name="Aptos", size=10, bold=True, color=BLUE)
    subtitle.fill = PatternFill("solid", fgColor=PALE_BLUE)
    subtitle.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def add_analysis_sheet(workbook, job, report) -> None:
    analysis = build_report_analysis(report)
    sheet = workbook.create_sheet("Analysis", 1)
    add_brand_header(sheet, report, "Result Analysis")
    sheet.sheet_view.zoomScale = 75
    sheet.freeze_panes = "A9"

    for column in range(1, 26):
        sheet.column_dimensions[get_column_letter(column)].width = 12
    sheet.column_dimensions["B"].width = 17
    sheet.column_dimensions["C"].width = 17
    sheet.column_dimensions["D"].width = 17
    sheet.column_dimensions["A"].width = 18

    summary = analysis["summary"]
    cards = (
        ("A9:C9", "A10:C11", "Students", summary["students"], BLUE),
        ("D9:F9", "D10:F11", "Pass rate", summary["pass_rate"], TEAL),
        ("G9:I9", "G10:I11", "Average SGPA", summary["average_sgpa"], BLUE),
        ("J9:L9", "J10:L11", "Highest SGPA", summary["highest_sgpa"], GOLD),
        ("M9:O9", "M10:O11", "Completion", summary["completion_rate"], TEAL),
    )
    for label_range, value_range, label, value, color in cards:
        sheet.merge_cells(label_range)
        sheet.merge_cells(value_range)
        label_cell = sheet[label_range.split(":", 1)[0]]
        value_cell = sheet[value_range.split(":", 1)[0]]
        label_cell.value = label
        value_cell.value = value if value is not None else "—"
        label_cell.fill = PatternFill("solid", fgColor=color)
        label_cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        label_cell.alignment = Alignment(horizontal="center", vertical="center")
        value_cell.fill = PatternFill("solid", fgColor="F8FAFC")
        value_cell.font = Font(name="Aptos Display", size=18, bold=True, color=NAVY)
        value_cell.alignment = Alignment(horizontal="center", vertical="center")
        _outline_range(sheet, label_range)
        _outline_range(sheet, value_range)
    sheet["D10"].number_format = "0.0\"%\""
    sheet["M10"].number_format = "0.0\"%\""
    sheet["G10"].number_format = "0.00"
    sheet["J10"].number_format = "0.00"

    sgpa_end = _write_table(
        sheet,
        14,
        1,
        "SGPA Distribution",
        ("Band", "Students"),
        [(row["label"], row["count"]) for row in analysis["sgpa_distribution"]],
    )
    status_end = _write_table(
        sheet,
        14,
        11,
        "Result Status",
        ("Status", "Students"),
        [(row["label"], row["count"]) for row in analysis["result_status"]],
    )
    _add_column_chart(
        sheet,
        min_col=1,
        max_col=2,
        header_row=15,
        data_end_row=sgpa_end,
        title="Students by SGPA Band",
        anchor="D14",
        width=8.5,
        height=7.5,
        y_title="Students",
    )
    _add_doughnut_chart(
        sheet,
        label_col=11,
        value_col=12,
        header_row=15,
        data_end_row=status_end,
        title="Declared Result Status",
        anchor="N14",
    )

    subject_rows = [
        (
            row["code"],
            row["name"],
            row["average_percentage"],
            row["pass_rate"],
            row["students"],
        )
        for row in analysis["subject_performance"]
    ]
    subject_end = _write_table(
        sheet,
        31,
        1,
        "Subject Performance",
        ("Code", "Subject", "Average %", "Pass Rate %", "Students"),
        subject_rows,
    )
    for row in range(33, subject_end + 1):
        sheet.cell(row, 3).number_format = "0.0"
        sheet.cell(row, 4).number_format = "0.0"
    sheet.column_dimensions["B"].width = 34
    _add_bar_chart(
        sheet,
        category_col=1,
        value_cols=(3, 4),
        header_row=32,
        data_end_row=subject_end,
        title="Subject Average and Pass Rate (%)",
        anchor="G31",
        width=10,
        height=8.5,
    )

    grade_end = _write_table(
        sheet,
        31,
        17,
        "Grade Distribution",
        ("Grade", "Count"),
        [(row["grade"], row["count"]) for row in analysis["grade_distribution"]],
    )
    _add_column_chart(
        sheet,
        min_col=17,
        max_col=18,
        header_row=32,
        data_end_row=grade_end,
        title="Grades Awarded",
        anchor="T31",
        width=8,
        height=8.5,
        y_title="Results",
    )

    detail_row = max(subject_end + 3, 52)
    _write_table(
        sheet,
        detail_row,
        1,
        "Top Performers",
        ("Seat Number", "Student", "SGPA", "Total Marks", "Result"),
        [
            (
                row["seat_number"],
                row["name"],
                row["sgpa"],
                row["total_marks"],
                row["result_status"],
            )
            for row in analysis["top_performers"]
        ],
    )
    _write_table(
        sheet,
        detail_row,
        8,
        "Attention Required",
        ("Seat Number", "Student", "SGPA", "Reason", "Validation"),
        [
            (
                row["seat_number"],
                row["name"],
                row["sgpa"],
                row["reason"],
                row["validation_status"],
            )
            for row in analysis["attention_students"]
        ],
    )

    validation_row = detail_row
    _write_table(
        sheet,
        validation_row,
        15,
        "Validation Status",
        ("Status", "Students"),
        [(row["label"], row["count"]) for row in analysis["validation_status"]],
    )

    sheet.oddFooter.center.text = (
        f"Import #{job.id} · Generated by Result Data Management System"
    )
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0


def _add_logo(sheet, path: Path, anchor: str, *, width: int, height: int) -> None:
    if not path.exists():
        return
    logo = Image(str(path))
    logo.width = width
    logo.height = height
    sheet.add_image(logo, anchor)


def _report_context(report) -> str:
    parts = []
    if report.branch:
        parts.append(report.branch)
    elif report.programme:
        parts.append(report.programme)
    if report.semester:
        parts.append(f"Semester {report.semester}")
    if report.examination_session:
        parts.append(report.examination_session)
    if report.examination_type:
        parts.append(report.examination_type.title())
    return "  ·  ".join(parts) or "Validated result dataset"


def _outline_range(sheet, range_address: str) -> None:
    for row in sheet[range_address]:
        for cell in row:
            cell.border = Border(
                left=LIGHT_BORDER,
                right=LIGHT_BORDER,
                top=LIGHT_BORDER,
                bottom=LIGHT_BORDER,
            )


def _write_table(sheet, title_row, start_col, title, headers, rows) -> int:
    end_col = start_col + len(headers) - 1
    sheet.merge_cells(
        start_row=title_row,
        start_column=start_col,
        end_row=title_row,
        end_column=end_col,
    )
    title_cell = sheet.cell(title_row, start_col)
    title_cell.value = title
    title_cell.fill = PatternFill("solid", fgColor=NAVY)
    title_cell.font = Font(name="Aptos", size=11, bold=True, color=WHITE)
    title_cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[title_row].height = 23

    header_row = title_row + 1
    for offset, header in enumerate(headers):
        cell = sheet.cell(header_row, start_col + offset, header)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=LIGHT_BORDER)

    data_rows = rows or [("No data",) + (None,) * (len(headers) - 1)]
    for row_offset, values in enumerate(data_rows, start=1):
        for column_offset, value in enumerate(values):
            cell = sheet.cell(header_row + row_offset, start_col + column_offset, value)
            cell.fill = PatternFill(
                "solid", fgColor="F8FAFC" if row_offset % 2 else "EEF4F8"
            )
            cell.border = Border(bottom=Side(style="hair", color="E2E8F0"))
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if column_offset == 0 and "Number" in headers[0]:
                cell.number_format = "@"
    return header_row + len(data_rows)


def _add_column_chart(
    sheet,
    *,
    min_col,
    max_col,
    header_row,
    data_end_row,
    title,
    anchor,
    width,
    height,
    y_title,
) -> None:
    if data_end_row <= header_row:
        return
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = title
    chart.y_axis.title = y_title
    chart.height = height
    chart.width = width
    chart.legend = None
    chart.add_data(
        Reference(
            sheet,
            min_col=max_col,
            min_row=header_row,
            max_row=data_end_row,
        ),
        titles_from_data=True,
    )
    chart.set_categories(
        Reference(
            sheet,
            min_col=min_col,
            min_row=header_row + 1,
            max_row=data_end_row,
        )
    )
    sheet.add_chart(chart, anchor)


def _add_bar_chart(
    sheet,
    *,
    category_col,
    value_cols,
    header_row,
    data_end_row,
    title,
    anchor,
    width,
    height,
) -> None:
    if data_end_row <= header_row:
        return
    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = title
    chart.x_axis.title = "Percentage"
    chart.height = height
    chart.width = width
    for value_col in value_cols:
        chart.add_data(
            Reference(
                sheet,
                min_col=value_col,
                min_row=header_row,
                max_row=data_end_row,
            ),
            titles_from_data=True,
        )
    chart.set_categories(
        Reference(
            sheet,
            min_col=category_col,
            min_row=header_row + 1,
            max_row=data_end_row,
        )
    )
    sheet.add_chart(chart, anchor)


def _add_doughnut_chart(
    sheet,
    *,
    label_col,
    value_col,
    header_row,
    data_end_row,
    title,
    anchor,
) -> None:
    if data_end_row <= header_row:
        return
    chart = DoughnutChart()
    chart.title = title
    chart.style = 10
    chart.holeSize = 55
    chart.height = 7.5
    chart.width = 7.5
    chart.add_data(
        Reference(
            sheet,
            min_col=value_col,
            min_row=header_row,
            max_row=data_end_row,
        ),
        titles_from_data=True,
    )
    chart.set_categories(
        Reference(
            sheet,
            min_col=label_col,
            min_row=header_row + 1,
            max_row=data_end_row,
        )
    )
    sheet.add_chart(chart, anchor)
