import json
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO

import pandas as pd
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from ..models import ExamReport, ImportJob
from .excel_presentation_service import (
    DATA_HEADER_ROW,
    DATA_START_ROW,
    add_analysis_sheet,
    add_brand_header,
)


EXPORT_SCHEMA_VERSION = "1.1"


@dataclass(frozen=True, slots=True)
class ExportArtifact:
    data: bytes
    filename: str
    mimetype: str


def build_excel_export(job: ImportJob) -> ExportArtifact:
    _require_report(job)
    from .report_excel_service import build_faculty_workbook
    return ExportArtifact(build_faculty_workbook(job), _filename(job, "xlsx"),
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def build_pdf_export(job: ImportJob) -> ExportArtifact:
    _require_report(job)
    from .report_pdf_service import build_faculty_pdf
    return ExportArtifact(build_faculty_pdf(job), _filename(job, "pdf"), "application/pdf")


def build_detailed_excel_export(job: ImportJob) -> ExportArtifact:
    report = _require_report(job)
    frames = _data_frames(report)
    stream = BytesIO()

    with pd.ExcelWriter(stream, engine="openpyxl") as writer:
        for sheet_name, frame in frames.items():
            frame.to_excel(
                writer,
                sheet_name=sheet_name,
                index=False,
                startrow=DATA_HEADER_ROW - 1,
            )
        _format_workbook(writer.book, job, report, frames)

    return ExportArtifact(
        data=stream.getvalue(),
        filename=_filename(job, "xlsx"),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def build_csv_export(job: ImportJob) -> ExportArtifact:
    report = _require_report(job)
    frame = _results_frame(report)
    csv_text = frame.to_csv(index=False, lineterminator="\n")
    return ExportArtifact(
        data=("\ufeff" + csv_text).encode("utf-8"),
        filename=_filename(job, "csv"),
        mimetype="text/csv; charset=utf-8",
    )


def build_json_export(job: ImportJob) -> ExportArtifact:
    report = _require_report(job)
    payload = _json_payload(job, report)
    return ExportArtifact(
        data=json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8"),
        filename=_filename(job, "json"),
        mimetype="application/json",
    )


def _require_report(job: ImportJob) -> ExamReport:
    if job.report is None:
        raise ValueError("This import has not been extracted.")
    if job.report.validation_status == "NOT_VALIDATED":
        raise ValueError("Run validation before exporting this import.")
    return job.report


def _data_frames(report: ExamReport) -> dict[str, pd.DataFrame]:
    frames = {
        "Overview": _overview_frame(report),
        "Students": _students_frame(report),
        "Subjects": _subjects_frame(report),
        "Results": _results_frame(report),
        "Components": _components_frame(report),
        "Validation Issues": _issues_frame(report),
        "Corrections": _corrections_frame(report),
    }
    # Page references remain in PostgreSQL and in the machine-facing CSV/JSON
    # exports, but are intentionally omitted from the presentation workbook.
    for frame in frames.values():
        if "Source Page" in frame.columns:
            frame.drop(columns=["Source Page"], inplace=True)
    return frames


def _overview_frame(report: ExamReport) -> pd.DataFrame:
    pass_count = sum(student.result_status == "PASS" for student in report.students)
    fail_count = sum(student.result_status == "FAIL" for student in report.students)
    cancel_count = sum(
        student.result_status == "CANCEL" for student in report.students
    )
    rows = [
        ("Export schema version", EXPORT_SCHEMA_VERSION),
        ("University", report.university),
        ("Examination", report.examination_title),
        ("Programme", report.programme),
        ("Branch", report.branch),
        ("Semester", report.semester),
        ("Examination session", report.examination_session),
        ("Examination type", report.examination_type),
        ("Institute code", report.institute_code),
        ("Institute name", report.institute_name),
        ("Maximum marks", _number(report.maximum_marks)),
        ("Maximum credits", _number(report.maximum_credits)),
        ("Maximum grade points", _number(report.maximum_grade_points)),
        ("Students", len(report.students)),
        ("Subjects", len(report.subjects)),
        ("Subject results", report.result_count),
        ("Pass", pass_count),
        ("Fail", fail_count),
        ("Cancel", cancel_count),
        ("Validation status", report.validation_status),
        ("Validation errors", report.validation_error_count),
        ("Validation warnings", report.validation_warning_count),
        ("Workflow status", report.import_job.status),
        ("Corrections", len(report.corrections)),
        ("Confirmed at (UTC)", _iso_utc(report.import_job.confirmed_at)),
        (
            "Confirmed by",
            (
                report.import_job.confirmed_by.username
                if report.import_job.confirmed_by
                else None
            ),
        ),
        ("Confirmation notes", report.import_job.confirmation_notes),
        ("Calculation rule", report.calculation_rule_version),
        (
            "Validated at (UTC)",
            _iso_utc(report.validated_at),
        ),
        ("Parser", f"{report.parser_name} {report.parser_version}"),
    ]
    return pd.DataFrame(rows, columns=["Field", "Value"])


def _students_frame(report: ExamReport) -> pd.DataFrame:
    rows = []
    for student in report.students:
        summary = student.summary
        rows.append(
            {
                "Seat Number": student.seat_number,
                "Student Name": student.name,
                "Gender Marker": student.gender_marker,
                "Center Code": student.center_code,
                "Registration Scope": student.registration_scope,
                "Institute Code": student.institute_code,
                "Institute Name": student.institute_name,
                "Source Page": student.source_page,
                "Reported Result Status": student.result_status,
                "Calculated Result Status": (
                    summary.calculated_status if summary else None
                ),
                "Validation Status": student.validation_status,
                "Reported SGPA": _number(summary.reported_sgpa if summary else None),
                "Calculated SGPA": _number(
                    summary.calculated_sgpa if summary else None
                ),
                "Reported Total Grade Points": _number(
                    summary.reported_total_grade_points if summary else None
                ),
                "Calculated Total Grade Points": _number(
                    summary.calculated_total_grade_points if summary else None
                ),
                "Reported Earned Credits": _number(
                    summary.reported_earned_credits if summary else None
                ),
                "Calculated Earned Credits": _number(
                    summary.calculated_earned_credits if summary else None
                ),
                "Reported Total Marks": _number(
                    summary.reported_total_marks if summary else None
                ),
                "Calculated Total Marks": _number(
                    summary.calculated_total_marks if summary else None
                ),
            }
        )
    extra_columns = []
    for row, student in zip(rows, report.students):
        for key, value in report.additional_summary_for(student.seat_number).items():
            column = "Source " + key.replace("_", " ").title()
            if column not in extra_columns:
                extra_columns.append(column)
            row[column] = value.get("raw_text")
    return pd.DataFrame(rows, columns=_student_columns() + extra_columns)


def _subjects_frame(report: ExamReport) -> pd.DataFrame:
    rows = []
    for subject in report.subjects:
        structure = "; ".join(
            f"{assessment.component_type} {assessment.maximum_marks}"
            + (
                f"/{assessment.minimum_marks}"
                if assessment.minimum_marks is not None
                else ""
            )
            for assessment in subject.assessments
        )
        rows.append(
            {
                "Display Order": subject.display_order,
                "Subject Code": subject.subject_code,
                "Subject Name": subject.subject_name,
                "Credits": _number(subject.credits),
                "Audit Subject": subject.is_audit,
                "Total Maximum": _number(subject.total_maximum),
                "Total Minimum": _number(subject.total_minimum),
                "Assessment Count": len(subject.assessments),
                "Assessment Structure": structure or "Grade only",
            }
        )
    return pd.DataFrame(rows, columns=_subject_columns())


def _results_frame(report: ExamReport) -> pd.DataFrame:
    component_types = _component_types(report)
    rows = []
    for student in report.students:
        for result in student.subject_results:
            components = {
                component.component_type: component for component in result.components
            }
            row = {
                "Seat Number": student.seat_number,
                "Student Name": student.name,
                "Reported Result Status": student.result_status,
                "Student Validation Status": student.validation_status,
                "Source Page": result.source_page,
                "Subject Order": result.subject.display_order,
                "Subject Code": result.subject.subject_code,
                "Subject Name": result.subject.subject_name,
                "Credits": _number(result.subject.credits),
                "Audit Subject": result.subject.is_audit,
                "Attempt Session": result.attempt_session,
            }
            for component_type in component_types:
                component = components.get(component_type)
                row[f"{component_type} Raw"] = (
                    component.raw_text if component else None
                )
                row[f"{component_type} Value"] = _number(
                    component.numeric_value if component else None
                )
                row[f"{component_type} Carried Forward"] = (
                    component.is_carried_forward if component else None
                )
                row[f"{component_type} Special Code"] = (
                    component.special_code if component else None
                )
            row.update(
                {
                    "Reported Total Raw": result.total_raw,
                    "Reported Total": _number(result.total_value),
                    "Component Sum": _number(result.component_sum),
                    "Total Carried Forward": result.total_is_carried_forward,
                    "Total Special Code": result.total_special_code,
                    "Reported Grade Raw": result.grade_raw,
                    "Reported Grade": result.grade,
                    "Expected Grade": result.expected_grade,
                    "Reported Grade Point": _number(result.grade_point),
                    "Expected Grade Point": _number(result.expected_grade_point),
                    "Reported Credit Points": _number(result.credit_points),
                    "Expected Credit Points": _number(result.expected_credit_points),
                    "Grace Marks": _number(result.grace_marks),
                    "Audit Grade": result.audit_grade,
                    "Result Validation Status": result.validation_status,
                }
            )
            rows.append(row)
    extra_columns = []
    extras_by_seat = {student.seat_number: report.additional_summary_for(student.seat_number)
                      for student in report.students}
    for row in rows:
        for key, value in extras_by_seat.get(row["Seat Number"], {}).items():
            column = "Source " + key.replace("_", " ").title()
            if column not in extra_columns:
                extra_columns.append(column)
            row[column] = value.get("raw_text")
    return pd.DataFrame(rows, columns=_result_columns(component_types) + extra_columns)


def _components_frame(report: ExamReport) -> pd.DataFrame:
    rows = []
    for student in report.students:
        for result in student.subject_results:
            for component in result.components:
                rows.append(
                    {
                        "Seat Number": student.seat_number,
                        "Student Name": student.name,
                        "Subject Code": result.subject.subject_code,
                        "Subject Name": result.subject.subject_name,
                        "Component Type": component.component_type,
                        "Component Name": component.assessment.display_name,
                        "Raw Value": component.raw_text,
                        "Numeric Value": _number(component.numeric_value),
                        "Carried Forward": component.is_carried_forward,
                        "Special Code": component.special_code,
                        "Maximum Marks": _number(component.assessment.maximum_marks),
                        "Minimum Marks": _number(component.assessment.minimum_marks),
                        "Source Page": result.source_page,
                    }
                )
    return pd.DataFrame(rows, columns=_component_columns())


def _issues_frame(report: ExamReport) -> pd.DataFrame:
    rows = []
    for issue in report.validation_issues:
        rows.append(
            {
                "Severity": issue.severity,
                "Code": issue.code,
                "Scope": issue.scope,
                "Seat Number": issue.student.seat_number if issue.student else None,
                "Student Name": issue.student.name if issue.student else None,
                "Subject Code": (
                    issue.subject_result.subject.subject_code
                    if issue.subject_result
                    else None
                ),
                "Field": issue.field_name,
                "Message": issue.message,
                "Reported Value": issue.reported_value,
                "Calculated Value": issue.calculated_value,
                "Source Page": (
                    issue.subject_result.source_page
                    if issue.subject_result
                    else (issue.student.source_page if issue.student else None)
                ),
            }
        )
    return pd.DataFrame(rows, columns=_issue_columns())


def _corrections_frame(report: ExamReport) -> pd.DataFrame:
    rows = []
    for correction in reversed(report.corrections):
        rows.append(
            {
                "Changed At (UTC)": _excel_utc(correction.created_at),
                "Seat Number": (
                    correction.student.seat_number if correction.student else None
                ),
                "Student Name": correction.student.name if correction.student else None,
                "Entity Type": correction.entity_type,
                "Entity Label": correction.entity_label,
                "Field": correction.field_name,
                "Original Value": correction.original_value,
                "Previous Value": correction.previous_value,
                "New Value": correction.new_value,
                "Reason": correction.reason,
                "Changed By": (
                    correction.changed_by.username
                    if correction.changed_by
                    else "Account removed"
                ),
            }
        )
    return pd.DataFrame(rows, columns=_correction_columns())


def _format_workbook(workbook, job, report, frames) -> None:
    navy = "17365D"
    blue = "1F4E78"
    pale_blue = "D9EAF7"
    pale_red = "FCE8E6"
    pale_yellow = "FFF4CC"
    pale_green = "E2F0D9"
    white = "FFFFFF"
    light_border = Side(style="thin", color="D9E2F3")

    workbook.properties.title = "Result Data Export"
    workbook.properties.subject = report.examination_title
    workbook.properties.creator = "Result Data Management System"
    workbook.properties.description = (
        "Extracted and validated result dataset. Reported and calculated values "
        "are intentionally stored separately."
    )

    add_analysis_sheet(workbook, job, report)

    for sheet_name, frame in frames.items():
        sheet = workbook[sheet_name]
        add_brand_header(sheet, report, sheet_name)
        header_row = DATA_HEADER_ROW
        freeze_panes = {
            "Overview": f"A{DATA_START_ROW}",
            "Students": f"C{DATA_START_ROW}",
            "Subjects": f"C{DATA_START_ROW}",
            "Results": f"C{DATA_START_ROW}",
            "Components": f"C{DATA_START_ROW}",
            "Validation Issues": f"D{DATA_START_ROW}",
            "Corrections": f"D{DATA_START_ROW}",
        }
        sheet.freeze_panes = freeze_panes[sheet_name]

        for cell in sheet[header_row]:
            cell.fill = PatternFill("solid", fgColor=blue)
            cell.font = Font(color=white, bold=True)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(bottom=light_border)
        sheet.row_dimensions[header_row].height = 30

        if sheet_name == "Overview":
            # Overview is a normal cell range, so it needs a worksheet filter.
            # Other populated sheets use Excel tables, which create their own
            # filters. Adding both over the same range makes Excel repair the
            # workbook because overlapping table/worksheet filters are invalid.
            sheet.auto_filter.ref = f"A{DATA_HEADER_ROW}:B{sheet.max_row}"
            for row in range(DATA_START_ROW, sheet.max_row + 1):
                sheet.cell(row, 1).font = Font(bold=True, color=navy)
                if sheet.cell(row, 1).value == "Validation status":
                    status_cell = sheet.cell(row, 2)
                    status_cell.fill = PatternFill(
                        "solid",
                        fgColor=(
                            pale_green
                            if report.validation_status == "VALID"
                            else (
                                pale_red
                                if report.validation_status == "ERROR"
                                else pale_yellow
                            )
                        ),
                    )
            sheet.column_dimensions["A"].width = 30
            sheet.column_dimensions["B"].width = 85
            for row in range(DATA_START_ROW, sheet.max_row + 1):
                if sheet.cell(row, 1).value == "Institute code":
                    code_cell = sheet.cell(row, 2)
                    if code_cell.value is not None:
                        code_cell.value = str(code_cell.value)
                    code_cell.number_format = "@"
                    code_cell.alignment = Alignment(horizontal="left")
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0
            sheet.oddFooter.center.text = (
                f"Import #{job.id} · Generated by Result Data Management System"
            )
            continue

        if len(frame.index) > 0:
            last_column = get_column_letter(len(frame.columns))
            table = Table(
                displayName=re.sub(r"[^A-Za-z0-9]", "", sheet_name) + "Table",
                ref=(
                    f"A{DATA_HEADER_ROW}:"
                    f"{last_column}{len(frame.index) + DATA_HEADER_ROW}"
                ),
            )
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            sheet.add_table(table)

        for column_index, column_name in enumerate(frame.columns, start=1):
            values = [column_name]
            values.extend(frame[column_name].head(200).tolist())
            maximum_length = max(
                len(str(value)) if value is not None else 0 for value in values
            )
            width = min(max(maximum_length + 2, 11), 42)
            if column_name in {
                "Message",
                "Reason",
                "Subject Name",
                "Institute Name",
            }:
                width = min(max(maximum_length + 2, 24), 55)
            sheet.column_dimensions[get_column_letter(column_index)].width = width

            if column_name in {"Seat Number", "Center Code", "Institute Code"}:
                for cell in sheet.iter_cols(
                    min_col=column_index,
                    max_col=column_index,
                    min_row=DATA_START_ROW,
                    max_row=sheet.max_row,
                ):
                    for item in cell:
                        if item.value is not None:
                            item.value = str(item.value)
                        item.number_format = "@"
                        item.alignment = Alignment(horizontal="left")
            if column_name in {"Message", "Reason"}:
                for cell in sheet.iter_cols(
                    min_col=column_index,
                    max_col=column_index,
                    min_row=DATA_START_ROW,
                    max_row=sheet.max_row,
                ):
                    for item in cell:
                        item.alignment = Alignment(wrap_text=True, vertical="top")
                for row_index in range(DATA_START_ROW, sheet.max_row + 1):
                    sheet.row_dimensions[row_index].height = 34
            if column_name.endswith("(UTC)"):
                for cell in sheet.iter_cols(
                    min_col=column_index,
                    max_col=column_index,
                    min_row=DATA_START_ROW,
                    max_row=sheet.max_row,
                ):
                    for item in cell:
                        item.number_format = "yyyy-mm-dd hh:mm:ss"
            if _is_numeric_column(column_name):
                for cell in sheet.iter_cols(
                    min_col=column_index,
                    max_col=column_index,
                    min_row=DATA_START_ROW,
                    max_row=sheet.max_row,
                ):
                    for item in cell:
                        item.number_format = "0.##"

        _add_status_formatting(
            sheet,
            frame,
            pale_red,
            pale_yellow,
            pale_green,
            header_row=DATA_HEADER_ROW,
        )

        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.oddFooter.center.text = (
            f"Import #{job.id} · Generated by Result Data Management System"
        )


def _add_status_formatting(
    sheet, frame, red, yellow, green, *, header_row: int
) -> None:
    if frame.empty:
        return
    first_row = header_row + 1
    last_row = len(frame.index) + header_row
    for index, name in enumerate(frame.columns, start=1):
        if "Status" not in name and name != "Severity":
            continue
        letter = get_column_letter(index)
        target = f"{letter}{first_row}:{letter}{last_row}"
        sheet.conditional_formatting.add(
            target,
            FormulaRule(
                formula=[f'${letter}{first_row}="ERROR"'],
                fill=PatternFill("solid", fgColor=red),
            ),
        )
        sheet.conditional_formatting.add(
            target,
            FormulaRule(
                formula=[
                    f'OR(${letter}{first_row}="WARNING",'
                    f'${letter}{first_row}="FAIL")'
                ],
                fill=PatternFill("solid", fgColor=yellow),
            ),
        )
        sheet.conditional_formatting.add(
            target,
            FormulaRule(
                formula=[
                    f'OR(${letter}{first_row}="VALID",'
                    f'${letter}{first_row}="PASS")'
                ],
                fill=PatternFill("solid", fgColor=green),
            ),
        )


def _json_payload(job: ImportJob, report: ExamReport) -> dict:
    return {
        "schema_version": EXPORT_SCHEMA_VERSION,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source_import": {
            "id": job.id,
            "original_filename": job.original_filename,
            "sha256": job.file_hash,
            "workflow_status": job.status,
            "page_count": job.page_count,
            "parser_name": report.parser_name,
            "parser_version": job.parser_version,
            "confirmed_at": _iso_utc(job.confirmed_at),
            "confirmed_by": job.confirmed_by.username if job.confirmed_by else None,
            "confirmation_notes": job.confirmation_notes,
        },
        "report": {
            "university": report.university,
            "examination_title": report.examination_title,
            "programme": report.programme,
            "branch": report.branch,
            "semester": report.semester,
            "examination_session": report.examination_session,
            "examination_type": report.examination_type,
            "institute_code": report.institute_code,
            "institute_name": report.institute_name,
            "maximum_marks": _number(report.maximum_marks),
            "maximum_credits": _number(report.maximum_credits),
            "maximum_grade_points": _number(report.maximum_grade_points),
            "report_generated_at_raw": report.report_generated_at_raw,
            "calculation_rule_version": report.calculation_rule_version,
        },
        "source_layouts": report.source_data.data if report.source_data else {},
        "subjects": [_json_subject(subject) for subject in report.subjects],
        "students": [_json_student(student) for student in report.students],
        "validation": {
            "status": report.validation_status,
            "error_count": report.validation_error_count,
            "warning_count": report.validation_warning_count,
            "validated_at": (
                _iso_utc(report.validated_at)
            ),
            "issues": [_json_issue(issue) for issue in report.validation_issues],
        },
        "corrections": [_json_correction(item) for item in report.corrections],
    }


def _json_subject(subject) -> dict:
    return {
        "display_order": subject.display_order,
        "subject_code": subject.subject_code,
        "subject_name": subject.subject_name,
        "credits": _number(subject.credits),
        "is_audit": subject.is_audit,
        "total_maximum": _number(subject.total_maximum),
        "total_minimum": _number(subject.total_minimum),
        "assessments": [
            {
                "component_type": assessment.component_type,
                "display_name": assessment.display_name,
                "display_order": assessment.display_order,
                "source_row": assessment.source_row,
                "maximum_marks": _number(assessment.maximum_marks),
                "minimum_marks": _number(assessment.minimum_marks),
            }
            for assessment in subject.assessments
        ],
    }


def _json_student(student) -> dict:
    summary = student.summary
    return {
        "student": {
            "seat_number": student.seat_number,
            "name": student.name,
            "gender_marker": student.gender_marker,
            "center_code": student.center_code,
            "registration_scope": student.registration_scope,
            "institute_code": student.institute_code,
            "institute_name": student.institute_name,
            "source_page": student.source_page,
            "reported_result_status": student.result_status,
            "validation_status": student.validation_status,
        },
        "results": [_json_result(result) for result in student.subject_results],
        "summary": {
            "additional_source_fields": student.report.additional_summary_for(student.seat_number),
            "reported": {
                "sgpa": _number(summary.reported_sgpa if summary else None),
                "total_grade_points": _number(
                    summary.reported_total_grade_points if summary else None
                ),
                "earned_credits": _number(
                    summary.reported_earned_credits if summary else None
                ),
                "total_marks": _number(
                    summary.reported_total_marks if summary else None
                ),
                "result_status": summary.reported_status if summary else None,
            },
            "calculated": {
                "sgpa": _number(summary.calculated_sgpa if summary else None),
                "total_grade_points": _number(
                    summary.calculated_total_grade_points if summary else None
                ),
                "earned_credits": _number(
                    summary.calculated_earned_credits if summary else None
                ),
                "total_marks": _number(
                    summary.calculated_total_marks if summary else None
                ),
                "result_status": summary.calculated_status if summary else None,
            },
        },
    }


def _json_result(result) -> dict:
    return {
        "subject_code": result.subject.subject_code,
        "source_page": result.source_page,
        "attempt_session": result.attempt_session,
        "components": [
            {
                "component_type": component.component_type,
                "raw_text": component.raw_text,
                "numeric_value": _number(component.numeric_value),
                "is_carried_forward": component.is_carried_forward,
                "special_code": component.special_code,
            }
            for component in result.components
        ],
        "reported": {
            "total_raw": result.total_raw,
            "total": _number(result.total_value),
            "total_is_carried_forward": result.total_is_carried_forward,
            "total_special_code": result.total_special_code,
            "grade_raw": result.grade_raw,
            "grade": result.grade,
            "grade_point": _number(result.grade_point),
            "credit_points": _number(result.credit_points),
            "grace_marks": _number(result.grace_marks),
            "audit_grade": result.audit_grade,
        },
        "calculated": {
            "component_sum": _number(result.component_sum),
            "expected_grade": result.expected_grade,
            "expected_grade_point": _number(result.expected_grade_point),
            "expected_credit_points": _number(result.expected_credit_points),
        },
        "validation_status": result.validation_status,
    }


def _json_issue(issue) -> dict:
    return {
        "severity": issue.severity,
        "scope": issue.scope,
        "code": issue.code,
        "seat_number": issue.student.seat_number if issue.student else None,
        "subject_code": (
            issue.subject_result.subject.subject_code
            if issue.subject_result
            else None
        ),
        "field_name": issue.field_name,
        "message": issue.message,
        "reported_value": issue.reported_value,
        "calculated_value": issue.calculated_value,
    }


def _json_correction(correction) -> dict:
    return {
        "changed_at": _iso_utc(correction.created_at),
        "seat_number": (
            correction.student.seat_number if correction.student else None
        ),
        "entity_type": correction.entity_type,
        "entity_label": correction.entity_label,
        "field_name": correction.field_name,
        "original_value": correction.original_value,
        "previous_value": correction.previous_value,
        "new_value": correction.new_value,
        "reason": correction.reason,
        "changed_by": (
            correction.changed_by.username
            if correction.changed_by
            else "Account removed"
        ),
    }


def _component_types(report: ExamReport) -> list[str]:
    component_types = []
    for subject in report.subjects:
        for assessment in subject.assessments:
            if assessment.component_type not in component_types:
                component_types.append(assessment.component_type)
    return component_types


def _student_columns() -> list[str]:
    return [
        "Seat Number",
        "Student Name",
        "Gender Marker",
        "Center Code",
        "Registration Scope",
        "Institute Code",
        "Institute Name",
        "Source Page",
        "Reported Result Status",
        "Calculated Result Status",
        "Validation Status",
        "Reported SGPA",
        "Calculated SGPA",
        "Reported Total Grade Points",
        "Calculated Total Grade Points",
        "Reported Earned Credits",
        "Calculated Earned Credits",
        "Reported Total Marks",
        "Calculated Total Marks",
    ]


def _subject_columns() -> list[str]:
    return [
        "Display Order",
        "Subject Code",
        "Subject Name",
        "Credits",
        "Audit Subject",
        "Total Maximum",
        "Total Minimum",
        "Assessment Count",
        "Assessment Structure",
    ]


def _result_columns(component_types: list[str]) -> list[str]:
    columns = [
        "Seat Number",
        "Student Name",
        "Reported Result Status",
        "Student Validation Status",
        "Source Page",
        "Subject Order",
        "Subject Code",
        "Subject Name",
        "Credits",
        "Audit Subject",
        "Attempt Session",
    ]
    for component_type in component_types:
        columns.extend(
            [
                f"{component_type} Raw",
                f"{component_type} Value",
                f"{component_type} Carried Forward",
                f"{component_type} Special Code",
            ]
        )
    columns.extend(
        [
            "Reported Total Raw",
            "Reported Total",
            "Component Sum",
            "Total Carried Forward",
            "Total Special Code",
            "Reported Grade Raw",
            "Reported Grade",
            "Expected Grade",
            "Reported Grade Point",
            "Expected Grade Point",
            "Reported Credit Points",
            "Expected Credit Points",
            "Grace Marks",
            "Audit Grade",
            "Result Validation Status",
        ]
    )
    return columns


def _component_columns() -> list[str]:
    return [
        "Seat Number",
        "Student Name",
        "Subject Code",
        "Subject Name",
        "Component Type",
        "Component Name",
        "Raw Value",
        "Numeric Value",
        "Carried Forward",
        "Special Code",
        "Maximum Marks",
        "Minimum Marks",
        "Source Page",
    ]


def _issue_columns() -> list[str]:
    return [
        "Severity",
        "Code",
        "Scope",
        "Seat Number",
        "Student Name",
        "Subject Code",
        "Field",
        "Message",
        "Reported Value",
        "Calculated Value",
        "Source Page",
    ]


def _correction_columns() -> list[str]:
    return [
        "Changed At (UTC)",
        "Seat Number",
        "Student Name",
        "Entity Type",
        "Entity Label",
        "Field",
        "Original Value",
        "Previous Value",
        "New Value",
        "Reason",
        "Changed By",
    ]


def _number(value: Decimal | None):
    if value is None:
        return None
    return int(value) if value == value.to_integral_value() else float(value)


def _iso_utc(value: datetime | None) -> str | None:
    if value is None:
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(timezone.utc).isoformat()


def _excel_utc(value: datetime | None) -> datetime | None:
    if value is None:
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(timezone.utc).replace(tzinfo=None)


def _is_numeric_column(column_name: str) -> bool:
    return any(
        token in column_name
        for token in (
            "Marks",
            "Value",
            "Credits",
            "Grade Point",
            "Credit Points",
            "SGPA",
            "Total",
            "Maximum",
            "Minimum",
            "Source Page",
            "Display Order",
            "Subject Order",
        )
    ) and not column_name.endswith("Raw")


def _filename(job: ImportJob, extension: str) -> str:
    report = job.report
    identity = "-".join(
        value
        for value in (
            report.branch if report else None,
            f"semester-{report.semester}" if report and report.semester else None,
            report.examination_session if report else None,
        )
        if value
    )
    stem = re.sub(r"[^A-Za-z0-9_-]+", "-", identity).strip("-").lower()
    stem = stem[:90] or "result-data"
    return f"{stem}-import-{job.id}-validated.{extension}"
