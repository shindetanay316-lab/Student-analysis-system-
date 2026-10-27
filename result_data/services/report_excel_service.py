"""Faculty workbook exporter for the existing Python application.

Uses the application's established openpyxl dependency, so Windows users do
not need a separate authoring runtime to download a workbook.
"""
from io import BytesIO
from math import ceil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

from .faculty_report_service import build_faculty_report
from .excel_presentation_service import TRUST_NAME, COLLEGE_NAME, COLLEGE_ADDRESS, DEPARTMENT_NAME

NAVY='17365D'; BLUE='1F4E78'; RED='FCE8E6'; PALE='EEF3F8'; WHITE='FFFFFF'
LINE=Side(style='thin',color='CBD5E1')
ASSETS=Path(__file__).resolve().parents[2]/'static'/'result_data'/'images'
HEADER=10


def write(sheet,row,column,value):
    cell=sheet.cell(row,column,value)
    # All source text is literal, including PRNs with leading zeros and names
    # that begin with '='. Only explicitly authored formulas are formula cells.
    if isinstance(value,str):cell.data_type='s'
    cell.font=Font(name='Times New Roman',size=10)
    cell.alignment=Alignment(vertical='center',wrap_text=True)
    return cell


def brand(sheet,model,title,last=8):
    last=max(last,8)
    sheet.sheet_view.showGridLines=False
    for col in range(1,last+1):sheet.column_dimensions[get_column_letter(col)].width=17
    for row,height in {1:21,2:30,3:22,4:25,5:8,6:25,7:22,8:8}.items():sheet.row_dimensions[row].height=height
    for row,value,size,bold in [(1,TRUST_NAME,11,False),(2,COLLEGE_NAME.upper(),16,True),
                               (3,COLLEGE_ADDRESS,10,False),(4,DEPARTMENT_NAME,12,True)]:
        sheet.merge_cells(start_row=row,start_column=2,end_row=row,end_column=last-1)
        cell=write(sheet,row,2,value);cell.font=Font(name='Times New Roman',size=size,bold=bold,color=NAVY)
        cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    for filename,anchor,w,h in [('college-left-logo.jpg','A1',54,75),('college-right-logo.jpg',f'{get_column_letter(last)}1',65,62)]:
        if (ASSETS/filename).exists():
            logo=Image(str(ASSETS/filename));logo.width=w;logo.height=h;sheet.add_image(logo,anchor)
    for row,value in [(6,title),(7,f'Semester {model["report"].semester or "-"}   {model["report"].examination_session or ""}   {model["report"].examination_type or ""}   {model["status"]}   Date: {model["options"]["report_date"]}')]:
        sheet.merge_cells(start_row=row,start_column=1,end_row=row,end_column=last)
        cell=write(sheet,row,1,value)
        cell.font=Font(name='Times New Roman',size=13 if row==6 else 10,bold=True,color=WHITE if row==6 else NAVY)
        cell.fill=PatternFill('solid',fgColor=NAVY if row==6 else PALE)
        cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A3 if last>13 else sheet.PAPERSIZE_A4
    sheet.sheet_properties.pageSetUpPr.fitToPage=True
    sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
    sheet.print_options.horizontalCentered=True
    sheet.print_title_rows=f'1:{HEADER}'
    sheet.oddFooter.center.text='Page &P of &N'
    sheet.oddFooter.left.text=COLLEGE_NAME
    sheet.oddFooter.right.text=model['options']['report_date']
    return last


def data_table(sheet,headers,rows,status_column=None,grade_matrix=False):
    last=len(headers)
    for col,label in enumerate(headers,1):
        cell=write(sheet,HEADER,col,label)
        cell.font=Font(name='Times New Roman',size=9,bold=True,color=WHITE)
        cell.fill=PatternFill('solid',fgColor=NAVY)
        cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        cell.border=Border(left=LINE,right=LINE,top=LINE,bottom=LINE)
    height=max(28,max(ceil(len(str(h))/17)+str(h).count('\n') for h in headers)*12+14)
    sheet.row_dimensions[HEADER].height=min(145,height) if grade_matrix else 34
    for row_index,values in enumerate(rows,HEADER+1):
        for col,value in enumerate(values,1):
            cell=write(sheet,row_index,col,value)
            cell.border=Border(left=LINE,right=LINE,top=LINE,bottom=LINE)
            cell.alignment=Alignment(vertical='center',horizontal='left' if col==3 else 'center',wrap_text=True)
            if row_index%2==0:cell.fill=PatternFill('solid',fgColor='F6F8FB')
            if isinstance(value,(float,int)) and col==last:cell.number_format='0.00'
        if last>=3:sheet.cell(row_index,2).number_format='@'
        sheet.row_dimensions[row_index].height=32
    end=max(HEADER+1,HEADER+len(rows))
    sheet.auto_filter.ref=f'A{HEADER}:{get_column_letter(last)}{end}'
    sheet.freeze_panes=f'D{HEADER+1}' if last>=3 else f'A{HEADER+1}'
    sheet.print_area=f'A1:{get_column_letter(max(last,8))}{end}'
    if status_column and rows:
        result_letter=get_column_letter(status_column)
        sheet.conditional_formatting.add(f'A{HEADER+1}:{get_column_letter(last)}{end}',FormulaRule(
            formula=[f'${result_letter}{HEADER+1}="FAIL"'],fill=PatternFill('solid',fgColor=RED),
            font=Font(color='9C0006')))
    return end


def add_signatures(sheet,model,row,last):
    # Four signature blocks span the actual sheet width.
    last=max(8,last)
    for index,(role,name) in enumerate(model['signatures']):
        first=1+index*last//4;end=(index+1)*last//4
        for r,value in [(row,role),(row+1,name)]:
            if end>first:sheet.merge_cells(start_row=r,start_column=first,end_row=r,end_column=end)
            cell=write(sheet,r,first,value);cell.font=Font(name='Times New Roman',size=9,bold=r==row)
            cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        sheet.row_dimensions[row].height=28;sheet.row_dimensions[row+1].height=25
    sheet.print_area=f'A1:{get_column_letter(last)}{row+2}'


def build_faculty_workbook(job):
    model=build_faculty_report(job)
    workbook=Workbook();workbook.remove(workbook.active)
    workbook.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True)
    summary=workbook.create_sheet('Result Summary');brand(summary,model,'Overall University Result')
    counts=model['counts'];stats=model['analysis']['summary']
    data_table(summary,['Class','Total','All Clear','Failed / Backlog','Withheld','Cancelled','Other','Pass %'],
        [[f'Semester {model["report"].semester or "-"}',len(model['students']),counts['PASS'],counts['FAIL'],counts['WITHHELD'],counts['CANCEL'],counts['OTHER'],None]])
    summary['H11']='=IF(C11+D11=0,"-",C11/(C11+D11))';summary['H11'].number_format='0.00%'
    summary.auto_filter.ref=None
    for row,label,value in [(15,'Highest SGPA',stats['highest_sgpa']),(16,'Student name(s)',stats['highest_student_names']),
                           (18,'Average SGPA',stats['average_sgpa']),(20,'Source',job.original_filename),
                           (22,'Status',model['status'])]:
        summary.merge_cells(start_row=row,start_column=1,end_row=row,end_column=2)
        summary.merge_cells(start_row=row,start_column=3,end_row=row,end_column=8)
        write(summary,row,1,label).font=Font(name='Times New Roman',size=11,bold=True,color=NAVY)
        cell=write(summary,row,3,value);cell.number_format='0.00' if row in {15,18} else 'General'
        summary.row_dimensions[row].height=40 if row==16 else 25
    for row,text in [(25,'Pass % = Passed / (Passed + Failed). Withheld and cancelled results are separate.'),
                     (27,'Reviewed university values are used. Missing grades are not inferred; inapplicable electives are blank.'),
                     (29,'Subject outcomes use printed grades. Audit subjects and missing grades have no inferred passing percentage.')]:
        summary.merge_cells(start_row=row,start_column=1,end_row=row,end_column=8)
        write(summary,row,1,text);summary.row_dimensions[row].height=30
    add_signatures(summary,model,33,8)

    sheet=workbook.create_sheet('Student Results')
    headers=['Sr. No.','PRN','Student Name']+[s.subject_name+'\n'+s.subject_code for s in model['subjects']]+['Result','SGPA']
    brand(sheet,model,'All Students Results',len(headers))
    rows=[[r['serial'],r['prn'],r['name']]+r['grades']+[r['result'],float(r['sgpa']) if r['sgpa'] is not None else None] for r in model['students']]
    end=data_table(sheet,headers,rows,status_column=len(headers)-1,grade_matrix=True)
    sheet.merge_cells(start_row=8,start_column=1,end_row=8,end_column=len(headers))
    write(sheet,8,1,'N/A: subject not applicable. NOT PRINTED: grade absent in PDF. MISSING: expected record absent. REVIEW: manually cleared. CANCELLED: no declared result.')
    sheet.row_dimensions[8].height=28

    sheet.column_dimensions['A'].width=9;sheet.column_dimensions['B'].width=21;sheet.column_dimensions['C'].width=37
    sheet.column_dimensions[get_column_letter(len(headers)-1)].width=15
    sheet.column_dimensions[get_column_letter(len(headers))].width=12
    sheet.sheet_view.zoomScale=75
    add_signatures(sheet,model,end+4,len(headers))

    sheet=workbook.create_sheet('Subject Analysis');brand(sheet,model,'Subject-wise Result')
    rows=[[r['serial'],r['name'],r['code'],r['faculty'],r['appeared'],r['passed'],r['failed'],None] for r in model['subject_rows']]
    end=data_table(sheet,['Sr. No.','Subject Name','Subject Code','Faculty Name','Appeared','Passed','Failed','Passing %'],rows)
    for row_index,r in enumerate(model['subject_rows'],HEADER+1):
        if r['passed'] is not None:
            sheet.cell(row_index,8,f'=IF(F{row_index}+G{row_index}=0,"-",F{row_index}/(F{row_index}+G{row_index}))').number_format='0.00%'
        sheet.row_dimensions[row_index].height=55
    sheet.column_dimensions['A'].width=9;sheet.column_dimensions['B'].width=48;sheet.column_dimensions['C'].width=22;sheet.column_dimensions['D'].width=30
    add_signatures(sheet,model,end+4,8)

    sheet=workbook.create_sheet('Top 10');brand(sheet,model,'Top 10 Students')
    rows=[[i,r['seat_number'],r['name'],r['rank'],r['sgpa']] for i,r in enumerate(model['top_students'],1)]
    end=data_table(sheet,['Sr. No.','PRN','Student Name','Rank','SGPA'],rows)
    sheet.column_dimensions['B'].width=22;sheet.column_dimensions['C'].width=40
    add_signatures(sheet,model,end+4,8)

    sheet=workbook.create_sheet('Failed Students');brand(sheet,model,'Failed Students')
    rows=[[i,r['prn'],r['name'],r['failed_subjects'],r['result'],float(r['sgpa']) if r['sgpa'] is not None else None] for i,r in enumerate(model['failed_students'],1)]
    end=data_table(sheet,['Sr. No.','PRN','Student Name','Failed Subjects','Result','SGPA'],rows,status_column=5)
    sheet.column_dimensions['B'].width=22;sheet.column_dimensions['C'].width=35;sheet.column_dimensions['D'].width=65
    for i,r in enumerate(model['failed_students'],HEADER+1):sheet.row_dimensions[i].height=max(45,ceil(len(r['failed_subjects'])/65)*14)
    add_signatures(sheet,model,end+4,8)

    sheet=workbook.create_sheet('Charts');brand(sheet,model,'University Result Charts',12)
    analysis=model['analysis'];block=10
    chart_sets=[('Result Status',['Status','Students'],[[r['label'],r['count']] for r in analysis['result_status']],'doughnut'),
                ('SGPA Distribution',['SGPA Band','Students'],[[r['label'],r['count']] for r in analysis['sgpa_distribution']],'bar'),
                ('Subject Performance',['Subject Name','Average Marks %','Pass Rate %'],[[r['name'],r['average_percentage'],r['pass_rate']] for r in analysis['subject_performance']],'horizontal'),
                ('Grade Distribution',['Grade','Subject Results'],[[r['grade'],r['count']] for r in analysis['grade_distribution']],'bar')]
    for title,labels,data,kind in chart_sets:
        write(sheet,block,1,title).font=Font(name='Times New Roman',bold=True,size=12,color=NAVY)
        for i,label in enumerate(labels,1):write(sheet,block+1,i,label)
        for j,values in enumerate(data,block+2):
            for i,value in enumerate(values,1):write(sheet,j,i,value)
            sheet.row_dimensions[j].height=40 if kind=='horizontal' else 21
        if data:
            chart=DoughnutChart() if kind=='doughnut' else BarChart()
            if kind=='horizontal':chart.type='bar';chart.height=max(12,len(data)*1.05);chart.width=25;chart.x_axis.title='Percentage'
            else:chart.height=10;chart.width=22
            chart.title=title
            chart.add_data(Reference(sheet,min_col=2,max_col=len(labels),min_row=block+1,max_row=block+1+len(data)),titles_from_data=True)
            chart.set_categories(Reference(sheet,min_col=1,min_row=block+2,max_row=block+1+len(data)))
            chart.style=10;sheet.add_chart(chart,f'E{block}')
        block+=max(25,len(data)*2+8) if kind=='horizontal' else 25
    sheet.column_dimensions['A'].width=42;sheet.column_dimensions['B'].width=18;sheet.column_dimensions['C'].width=18
    sheet.freeze_panes='A10';sheet.print_area=f'A1:P{block}'
    stream=BytesIO();workbook.save(stream);return stream.getvalue()
