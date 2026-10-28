"""Portable student history using the project's existing Excel/PDF stack."""
from io import BytesIO
from xml.sax.saxutils import escape
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def display(value, pending='Pending'):
    return value if value is not None else pending


def sections(data):
    student = data['student']
    prefix = lambda i: [i, student.prn, student.name]
    identity = ['Sr. No.', 'PRN', 'Student Name']
    local = [prefix(i) + [r['academic_year'], r['semester_id'], r['subject_name'], r['subject_code'], r['credits'],
        display(r['internal']), display(r['external'], 'Absent' if (r['states'].get('EXTERNAL') == 'ABSENT' or r['states'].get('LAB_EXTERNAL') == 'ABSENT') else 'Pending'),
        display(r['total']), r['grade'] or 'Pending', r['result'], ', '.join(k + ' absent' for k, v in r['states'].items() if v == 'ABSENT')]
        for i, r in enumerate(data['local'], 1)]
    official = []
    for i, item in enumerate(data['official'], 1):
        r = item['record']
        official.append(prefix(i) + [r.academic_year, r.semester_id, r.source['subject_name'], r.source['subject_code'], r.subject_code,
            r.publication.examination_session or '', r.publication.examination_type or '',
            display(r.source['total'], r.source.get('total_raw') or 'Not printed'),
            r.source['grade_display'],
            r.source['student_status'], display(r.source.get('sgpa'), 'Not printed') if not data['scoped'] else 'Withheld in teacher scope',
            'Earlier source version' if item['superseded'] else 'Current source version',
            'Yes' if r.ese_copied else 'No', r.publication.metadata_snapshot['filename'],
            '; '.join(v['type'] + ': ' + str(v['raw']) + (' (carried)' if v['carried'] else '') for v in r.source['components'])])
    attendance = [prefix(i) + [r['academic_year'], r['semester_id'], r['subject_name'], r['subject_code'], r['total'], r['present'],
        r['total'] - r['present'], r['leave'], r['official'], display(r['percentage']), 'Review conflicts' if r['conflicts'] else 'Recorded' if r['total'] else 'Pending']
        for i, r in enumerate(data['attendance'], 1)]
    curriculum = [prefix(i) + [r.academic_year, r.semester_id, r.binding.version.title, r.binding.version.scheme,
        r.binding.regular_cohort, r.binding.group_name] for i, r in enumerate(data['curriculum'], 1)]
    corrections = [prefix(i) + [r.context['academic_year'], r.context['semester_id'], r.context['subject_code'], r.category, r.component,
        display(r.previous['value'], 'Absent'), display(r.proposed['value'], 'Absent'), r.status, r.reason, r.decision_note or '', str(r.created_at)]
        for i, r in enumerate(data['corrections'], 1)]
    tables = [
        ('Local Results', identity + ['Teaching year', 'Semester', 'Subject', 'Code', 'Credits', 'Internal', 'External', 'Total', 'Grade', 'Result', 'Absence details'], local),
        ('University Evidence', identity + ['Teaching year', 'Semester', 'Printed subject', 'Printed code', 'Local code', 'Examination session', 'Examination type',
            'Printed total', 'Printed grade', 'Source status', 'Printed SGPA', 'Source version', 'ESE copied', 'Source PDF', 'Printed components'], official),
        ('Attendance', identity + ['Teaching year', 'Semester', 'Subject', 'Code', 'Counted classes', 'Present', 'Absent', 'Leave', 'Official duty', 'Percentage', 'Status'], attendance),
        ('Curriculum', identity + ['Teaching year', 'Semester', 'Version', 'Scheme', 'Regular cohort', 'Group'], curriculum),
        ('Corrections', identity + ['Teaching year', 'Semester', 'Code', 'Category', 'Component', 'Previous', 'Proposed', 'Status', 'Reason', 'Decision note', 'Requested at'], corrections),
    ]
    if data.get('configured_components'):
        rows=[prefix(i)+[r['academic_year'],r['semester_id'],r['subject_name'],r['subject_code'],r['label'],r['code'],r['maximum'],
            float(r['value']) if r['value'] is not None else 'Absent' if r['state']=='ABSENT' else 'Pending',r['state'],r['revision']]
            for i,r in enumerate(data['configured_components'],1)]
        tables.append(('Configured Components',identity+['Teaching year','Semester','Subject','Subject code','Assessment','Component code','Maximum','Mark','State','Revision'],rows))
    return tables


def generate_history(data, fmt):
    student = data['student']
    buffer = BytesIO()
    safe_prn = re.sub(r'[^A-Za-z0-9_-]', '_', student.prn)
    filename = 'Student_History_' + safe_prn
    if fmt == 'excel':
        book = Workbook()
        profile = book.active
        profile.title = 'Student Profile'
        profile.append(['ACADPULSE - Student history'])
        profile.append(['PRN', student.prn])
        profile.append(['Student Name', student.name])
        profile.append(['Actual admission year', student.year_of_admission])
        profile.append(['Legacy default division / batch', (student.division or 'Not recorded') + ' / ' + (student.batch or 'Not recorded')])
        profile.append(['Scope', 'Assigned subjects only' if data['scoped'] else 'Administrator: complete local and linked university record'])
        profile.append(['Interpretation', 'Local calculations and printed university evidence are separate. Unlinked PDFs are not included. Pending is not zero.'])
        for key, spec in data['portfolio_fields'].items():
            text = data['portfolio'].details.get(key) if data['portfolio'] and data['portfolio'].details.get(key) else 'Not recorded'
            for offset in range(0, len(text), 1000):
                profile.append([spec[0] + (' (continued)' if offset else ''), text[offset:offset + 1000]])
        profile.append(['Portfolio note', 'Recorded profile entries are not verification of a credential.'])
        profile.append(['Locally calculated CGPA', 'Withheld in teacher scope' if data['scoped'] else display(data['cgpa'])])
        profile.append([])
        profile.append(['Local semester summaries', 'Teaching year', 'Semester', 'SGPA', 'Registered credits', 'Earned credits', 'Local status'])
        for row in data['summaries']:
            profile.append(['Local calculation', row['academic_year'], row['semester_id'], display(row['sgpa']), row['credits'], row['earned'], row['result']])
        profile.append([])
        profile.append(['Term placements','Teaching year','Semester','Roll number','Division','Practical batch','Source'])
        for row in data.get('placements',[]):
            profile.append(['Placement',row.academic_year,row.semester_id,row.roll_number,row.division,row.batch,row.origin])
        profile.append([])
        profile.append(['Approved course selections','Subject code','Semester','Selected teaching year','Grade','Grade point','Result','Evidence'])
        for row in data.get('selected_outcomes',[]):
            e=row.evidence
            profile.append(['Approved selection',row.subject_code,row.semester_id,e['teaching_year'],e['grade'],e['grade_point'],e['result'],row.source_key])
        profile.append([])
        profile.append(['Progression approvals','Teaching year','Semester','Decision','Revision','Reason'])
        for row in data.get('progression',[]):
            profile.append(['Progression',row.academic_year,row.semester_id,row.decision,row.revision,row.reason])
        profile.column_dimensions['A'].width = 32
        profile.column_dimensions['B'].width = 85
        for c in range(3, 8):
            profile.column_dimensions[get_column_letter(c)].width = 20
        profile.freeze_panes = 'B2'
        for title, headers, rows in sections(data):
            sheet = book.create_sheet(title)
            sheet.append(headers)
            for row in rows:
                sheet.append(row)
            sheet.freeze_panes = 'D2'
            sheet.auto_filter.ref = sheet.dimensions
            for col, header in enumerate(headers, 1):
                width = 9 if col == 1 else 23 if col == 2 else 34 if col == 3 else 38 if any(w in header.lower() for w in ('subject', 'reason', 'components', 'pdf', 'version', 'note')) else 21
                sheet.column_dimensions[get_column_letter(col)].width = width
            for cell in sheet[1]:
                cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='14283F')
            sheet.row_dimensions[1].height = 34
        for sheet in book:
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'  # Profile/printed text must never execute as an Excel formula.
                    if cell.row != 1:
                        cell.font = Font(name='Calibri', size=11)
                    cell.alignment = Alignment(wrap_text=True, vertical='top')
            for row in range(2, sheet.max_row + 1):
                lines = 1
                for cell in sheet[row]:
                    width = sheet.column_dimensions[cell.column_letter].width or 20
                    lines = max(lines, sum(max(1, (len(part) + int(width) - 1) // max(1, int(width) - 3)) for part in str(cell.value or '').split('\n')))
                sheet.row_dimensions[row].height = min(409, max(32, lines * 15 + 8))
        book.save(buffer)
        suffix, mime = '.xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    else:
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, CondPageBreak
        styles = getSampleStyleSheet()
        small = ParagraphStyle('HistoryCell', fontName='Helvetica', fontSize=8, leading=11, wordWrap='CJK')
        paragraph = lambda value: Paragraph(escape(str(value)).replace('\n', '<br/>'), small)
        story = [Paragraph('ACADPULSE - Student history', styles['Title']),
            Paragraph(escape(student.name) + ' | PRN ' + escape(student.prn), styles['Heading2']),
            paragraph(f'Actual admission: {student.year_of_admission}. Division and batch are recorded per teaching term.'),
            paragraph('Scope: assigned subjects only' if data['scoped'] else 'Scope: complete administrator view'),
            paragraph('Printed university evidence and local calculations are separate. Unlinked PDFs are not included. Pending is not zero.'), Spacer(1, 12)]
        story.append(Paragraph('Student portfolio', styles['Heading2']))
        for key, spec in data['portfolio_fields'].items():
            value = data['portfolio'].details.get(key) if data['portfolio'] else None
            story.extend([Paragraph(escape(spec[0]), styles['Heading3']), paragraph(value or 'Not recorded')])
        if data['summaries']:
            story.append(Paragraph('Local semester summaries', styles['Heading2']))
            story.append(paragraph('Locally calculated CGPA: ' + str(display(data['cgpa']))))
            for row in data['summaries']:
                story.append(paragraph(f"{row['academic_year']} / Semester {row['semester_id']} | SGPA {display(row['sgpa'])} | Credits {row['earned']}/{row['credits']} | {row['result']}"))
        # The PDF is a compact readable summary; Excel retains full component
        # evidence, mapping details and correction notes in separate sheets.
        story.append(paragraph('Excel includes the full printed component evidence and source details.'))
        if data.get('placements'):
            story.append(Paragraph('Teaching-term placements',styles['Heading2']))
            for row in data['placements']:
                story.append(paragraph(f'{row.academic_year} / Semester {row.semester_id} | Roll {row.roll_number or "Unassigned"} | Division {row.division or "Unassigned"} | Batch {row.batch or "Unassigned"}'))
        if data.get('selected_outcomes'):
            story.append(Paragraph('Approved course selections',styles['Heading2']))
            for row in data['selected_outcomes']:
                e=row.evidence
                story.append(paragraph(f"Semester {row.semester_id} / {e['subject_name']} ({row.subject_code}) | {e['grade']} / {e['grade_point']} | {e['result']} | {e['source']} {e['teaching_year']}"))
        if data.get('progression'):
            story.append(Paragraph('Progression approvals',styles['Heading2']))
            for row in data['progression']:
                story.append(paragraph(f'{row.academic_year} / Semester {row.semester_id} | {row.decision} | Revision {row.revision} | {row.reason}'))
        for title, headers, rows in sections(data):
            if rows:
                story.append(CondPageBreak(85))
            story.append(Paragraph(title, styles['Heading2']))
            if not rows:
                story.append(paragraph('No records in this scope.'))
                continue
            identity = ['Sr. No.', 'PRN', 'Student Name']
            if title == 'Local Results':
                columns = identity + ['Teaching year / sem', 'Subject', 'Internal / ESE', 'Total', 'Grade', 'Local result']
                summary_rows = [r[:3] + [f'{r[3]} / S{r[4]}', f'{r[5]} ({r[6]})', f'{r[8]} / {r[9]}', r[10], r[11], r[12]] for r in rows]
            elif title == 'University Evidence':
                columns = identity + ['Teaching year / sem', 'Printed subject / local code', 'Examination / version', 'Total / grade', 'Source status / SGPA', 'Source PDF']
                summary_rows = [r[:3] + [f'{r[3]} / S{r[4]}', f'{r[5]} ({r[6]}) / {r[7]}', f'{r[8]} {r[9]} / {r[14]}', f'{r[10]} / {r[11]}', f'{r[12]} / SGPA {r[13]}', r[16]] for r in rows]
            elif title == 'Attendance':
                columns = identity + ['Teaching year / sem', 'Subject', 'Present / counted', 'Leave / OD', 'Percentage', 'Status']
                summary_rows = [r[:3] + [f'{r[3]} / S{r[4]}', f'{r[5]} ({r[6]})', f'{r[8]} / {r[7]}', f'{r[10]} / {r[11]}', r[12], r[13]] for r in rows]
            elif title == 'Curriculum':
                columns = identity + ['Teaching year / sem', 'Approved version', 'Scheme', 'Regular cohort', 'Group', '']
                summary_rows = [r[:3] + [f'{r[3]} / S{r[4]}', r[5], r[6], r[7], r[8], ''] for r in rows]
            elif title=='Configured Components':
                columns=identity+['Teaching year / sem','Subject','Assessment','Mark / maximum','State','Revision']
                summary_rows=[r[:3]+[f'{r[3]} / S{r[4]}',f'{r[5]} ({r[6]})',f'{r[7]} ({r[8]})',f'{r[10]} / {r[9]}',r[11],r[12]] for r in rows]
            else:
                columns = identity + ['Teaching year / sem', 'Subject / component', 'Previous / proposed', 'Status', 'Reason', 'Decision note']
                summary_rows = [r[:3] + [f'{r[3]} / S{r[4]}', f'{r[5]} / {r[6]} {r[7]}', f'{r[8]} / {r[9]}', r[10], r[11], r[12]] for r in rows]
            body = [[paragraph(v) for v in row] for row in [columns] + summary_rows]
            table = Table(body, colWidths=[32, 85, 105, 82, 170, 78, 60, 64, 89], hAlign='LEFT', repeatRows=1, splitInRow=1)
            table.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e5ebf2')),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
                ('GRID', (0, 0), (-1, -1), .25, colors.HexColor('#cbd5e1')),
                ('LEFTPADDING', (0, 0), (-1, -1), 4), ('RIGHTPADDING', (0, 0), (-1, -1), 4)]))
            story.extend([table, Spacer(1, 10)])
        SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=28, rightMargin=28, topMargin=24, bottomMargin=24).build(story)
        suffix, mime = '.pdf', 'application/pdf'
    buffer.seek(0)
    return buffer, filename + suffix, mime
